# VitalAdvisorTool, 
# DiseaseAdvisorTool, 
# BloodReportAnalyzerTool, 
# BloodReportQueryTool, 
# PanelQueryTool, 
# CalorieCalculatorTool

import base64
import PyPDF2
from config import BASE_DIR, PREDEFINED_BLOOD_REPORT_JSON_PATH, PREDEFINED_PANEL_JSON_PATH
from docx import Document
import csv
from PIL import Image
import pytesseract
import requests
from datetime import datetime
import json
import logging
import os
from dotenv import load_dotenv
import re
import time
from typing import List, Dict, Any, Optional, Tuple
from rapidfuzz import process, fuzz
import traceback
from .tool_core import BaseTool, LLMService, ToolResult, ToolType 
from .base_and_utility import DataService, DatabasePersistenceTool 
from helpers.utils import serialize_data

logger = logging.getLogger(__name__)

load_dotenv()

MISTRAL_API_ENDPOINT = os.getenv("MISTRAL_API_ENDPOINT")
MISTRAL_API_KEY = os.getenv("MISTRAL_API_KEY") 

FULL_PANEL_DATA = {}
def _load_panel_json_data():
    global FULL_PANEL_DATA
    try:
        with open(PREDEFINED_PANEL_JSON_PATH, 'r', encoding='utf-8') as f:
            FULL_PANEL_DATA = json.load(f)
        logger.info(f"Successfully loaded data from {PREDEFINED_PANEL_JSON_PATH}")
    except FileNotFoundError:
        logger.error(f"Error: {PREDEFINED_PANEL_JSON_PATH} not found.")
    except json.JSONDecodeError as e:
        logger.error(f"Error decoding JSON from {PREDEFINED_PANEL_JSON_PATH}: {e}")
    except Exception as e:
        logger.error(f"An unexpected error occurred loading panel JSON: {e}")


_load_panel_json_data()

def create_translation_signature(translation, panels):
    blood_data = translation.get('blood_report_data', {})
   
    panel_tests = {}
    for test_name, test_data in blood_data.items():
        panel = test_data.get('Panel', 'UNKNOWN')
        extracted_value = test_data.get('Extracted Value')
       
        if panel not in panel_tests:
            panel_tests[panel] = []
       
        test_signature = f"{test_name}:{extracted_value}"
        panel_tests[panel].append(test_signature)
   
    signature_parts = []
    for panel in sorted(panel_tests.keys()):
        tests = sorted(panel_tests[panel])
        signature_parts.append(f"{panel}:[{','.join(tests)}]")
   
    return '|'.join(signature_parts)
 
def are_translations_identical(sig1, sig2):
    return sig1 == sig2
 
def parse_date(date_str):
    if not date_str:
        return datetime.min
   
    try:
        if '.' in date_str:
            return datetime.strptime(date_str.split('.')[0], '%Y-%m-%d %H:%M:%S')
        else:
            return datetime.strptime(date_str, '%Y-%m-%d %H:%M:%S')
    except ValueError:
        try:
            return datetime.strptime(date_str[:19], '%Y-%m-%d %H:%M:%S')
        except ValueError:
            return datetime.min
 

def find_abnormal_key(test_name, abnormal_test_ranges):
    if test_name in abnormal_test_ranges:
        return test_name
   
    for key in abnormal_test_ranges.keys():
        if key.lower().replace(',', '').replace(' ', '') == test_name.lower().replace(',', '').replace(' ', ''):
            return key
       
        if key.lower() in test_name.lower() or test_name.lower() in key.lower():
            return key
   
    return None
 
def is_test_abnormal(test_data):
    extracted_value = test_data.get('Extracted Value')
    normal_range = test_data.get('Normal Range', [])
   
    if extracted_value is None or not normal_range:
        return False
   
    try:
        range_str = normal_range[0] if isinstance(normal_range, list) else str(normal_range)
       
        if '<' in range_str:
            upper_match = re.search(r'<\s*(\d+\.?\d*)', range_str)
            if upper_match:
                upper_limit = float(upper_match.group(1))
                return extracted_value >= upper_limit
               
        elif '>' in range_str:
            lower_match = re.search(r'>\s*(\d+\.?\d*)', range_str)
            if lower_match:
                lower_limit = float(lower_match.group(1))
                return extracted_value <= lower_limit
               
        else:
            range_match = re.search(r'(\d+\.?\d*)\s*[–\-]\s*(\d+\.?\d*)', range_str)
            if range_match:
                lower_limit = float(range_match.group(1))
                upper_limit = float(range_match.group(2))
                return extracted_value < lower_limit or extracted_value > upper_limit
   
    except (ValueError, AttributeError, IndexError):
        pass
   
    return False

def process_blood_report_translations(raw_data):
    valid_translations = []
   
    for report in raw_data:
        if not isinstance(report, dict) or not report.get('translation'):
            continue
           
        translation_str = report.get('translation')
        report_date = report.get('created_date')
       
        analysis_from_api = None
        if isinstance(translation_str, str):
            try:
                analysis_from_api = json.loads(translation_str)
                if not isinstance(analysis_from_api, dict):
                    raise ValueError("Translation is not a dict after parsing")
            except (ValueError, SyntaxError):
                continue
        elif isinstance(translation_str, dict):
            analysis_from_api = translation_str
       
        if analysis_from_api and analysis_from_api.get("blood_report_data"):
            filtered_translation = filter_abnormal_tests_only(analysis_from_api)
           
            if filtered_translation and filtered_translation.get("blood_report_data"):
                valid_translations.append({
                    'date': report_date,
                    'translation': filtered_translation,
                    'panels': filtered_translation.get('detected_panels', []),
                    'test_names': list(filtered_translation.get('blood_report_data', {}).keys()),
                    'document_id': report.get('document_id')
                })
   
    valid_translations.sort(key=lambda x: parse_date(x['date']), reverse=True)
   
    final_translations = []
    processed_signatures = {}
   
    for item in valid_translations:
        translation = item['translation']
        date = item['date']
        panels = item['panels']
        document_id = item['document_id']
       
        translation_signature = create_translation_signature(translation, panels)
       
        should_include = False
        is_duplicate = False
       
        for existing_sig, existing_data in processed_signatures.items():
            if are_translations_identical(translation_signature, existing_sig):
                if parse_date(date) > parse_date(existing_data['date']):
                    final_translations = [t for t in final_translations if t.get('document_id') != existing_data['document_id']]
                    should_include = True
                    del processed_signatures[existing_sig]
                else:
                    is_duplicate = True
                break
       
        if not is_duplicate and not should_include:
            should_include = True
       
        if should_include:
            processed_signatures[translation_signature] = {
                'date': date,
                'document_id': document_id,
                'panels': panels
            }
            final_translations.append(translation)
   
    return final_translations
 
def filter_abnormal_tests_only(analysis_from_api):
    blood_report_data = analysis_from_api.get("blood_report_data", {})
    abnormal_test_ranges = analysis_from_api.get("abnormal_test_ranges", {})
   
    filtered_blood_data = {}
    filtered_abnormal_ranges = {}
   
    for test_name, test_data in blood_report_data.items():
        abnormal_key = find_abnormal_key(test_name, abnormal_test_ranges)
       
        if abnormal_key:
            filtered_blood_data[test_name] = test_data
            filtered_abnormal_ranges[test_name] = abnormal_test_ranges[abnormal_key]
        else:
            if is_test_abnormal(test_data):
                filtered_blood_data[test_name] = test_data
                extracted_value = test_data.get('Extracted Value')
                if extracted_value is not None:
                    normal_range = test_data.get('Normal Range', [])
                    status = determine_abnormal_status(extracted_value, normal_range)
                    if status:
                        filtered_abnormal_ranges[test_name] = {
                            "value": extracted_value,
                            "status": status
                        }
   
    if filtered_blood_data:
        filtered_translation = analysis_from_api.copy()
        filtered_translation["blood_report_data"] = filtered_blood_data
        filtered_translation["abnormal_test_ranges"] = filtered_abnormal_ranges
        filtered_translation["matched_test_names"] = list(filtered_blood_data.keys())
        return filtered_translation
   
    return None

def determine_abnormal_status(extracted_value, normal_range):
    if not normal_range:
        return None
   
    try:
        range_str = normal_range[0] if isinstance(normal_range, list) else str(normal_range)
       
        if '<' in range_str:
            upper_match = re.search(r'<\s*(\d+\.?\d*)', range_str)
            if upper_match:
                upper_limit = float(upper_match.group(1))
                return "high" if extracted_value >= upper_limit else None
               
        elif '>' in range_str:
            lower_match = re.search(r'>\s*(\d+\.?\d*)', range_str)
            if lower_match:
                lower_limit = float(lower_match.group(1))
                return "low" if extracted_value <= lower_limit else None
               
        else:
            range_match = re.search(r'(\d+\.?\d*)\s*[–\-]\s*(\d+\.?\d*)', range_str)
            if range_match:
                lower_limit = float(range_match.group(1))
                upper_limit = float(range_match.group(2))
                if extracted_value < lower_limit:
                    return "low"
                elif extracted_value > upper_limit:
                    return "high"
   
    except (ValueError, AttributeError, IndexError):
        pass
   
    return None

def get_age_group(age_val):
    if age_val is None:
        return "Adult"
    if 0 <= age_val <= 1:
        return "Newborn"
    elif 1 < age_val <= 12:
        return "Child"
    elif 13 <= age_val <= 17:
        return "Adolescent"
    elif 18 <= age_val <= 64:
        return "Adult"
    elif age_val >= 65:
        return "Elderly"
    return "Adult"

def is_in_range(value, range_str, vital_name):
    if value is None or range_str == "N/A" or not range_str:
        return False
    range_str_cleaned = range_str.replace(' ', '').replace('–', '-').lower()
    def extract_numeric(s):
        match = re.search(r'[+-]?(\d+(\.\d*)?|\.\d+)', s)
        return float(match.group(0)) if match else None
    if vital_name == "Blood Pressure" and isinstance(value, dict):
        if '/' in range_str_cleaned:
            parts = range_str_cleaned.split('/')
            if len(parts) != 2:
                return False
            sys_range, dia_range = parts[0], parts[1]
            sys_check = False
            if '-' in sys_range:
                s_parts = sys_range.split('-')
                s_min = extract_numeric(s_parts[0])
                s_max = extract_numeric(s_parts[1])
                sys_check = s_min <= value.get('systolic', float('inf')) <= s_max if s_min is not None and s_max is not None else False
            else:
                s_val = extract_numeric(sys_range)
                sys_check = value.get('systolic', float('inf')) == s_val if s_val is not None else False
            dia_check = False
            if '-' in dia_range:
                d_parts = dia_range.split('-')
                d_min = extract_numeric(d_parts[0])
                d_max = extract_numeric(d_parts[1])
                dia_check = d_min <= value.get('diastolic', float('inf')) <= d_max if d_min is not None and d_max is not None else False
            else:
                d_val = extract_numeric(dia_range)
                dia_check = value.get('diastolic', float('inf')) == d_val if d_val is not None else False
            return sys_check and dia_check
        return False
    if not isinstance(value, (int, float)):
        return False
    if '%' in range_str_cleaned:
        range_str_cleaned = range_str_cleaned.replace('%', '')
    if '-' in range_str_cleaned:
        parts = range_str_cleaned.split('-')
        if len(parts) == 2:
            min_val = extract_numeric(parts[0])
            max_val = extract_numeric(parts[1])
            return min_val <= value <= max_val if min_val is not None and max_val is not None else False
    else:
        numeric_val = extract_numeric(range_str_cleaned)
        if numeric_val is not None:
            return value == numeric_val
    return False
TEST_ABBREVIATIONS = {}

def _get_test_name_variations(test_name):
    variations = []    
    common_mappings = {
        "CHOLESTEROL": ["TOTAL CHOLESTEROL", "CHOL"],
        "LOW DENSITY LIPOPROTEIN": ["LDL", "LDL CHOLESTEROL", "LOW DENSITY LIPOPROTEIN"],
        "HIGH DENSITY LIPOPROTEIN": ["HDL", "HDL CHOLESTEROL", "HIGH DENSITY LIPOPROTEIN"],
        "TRIGLYCERIDES": ["TG", "TRIGLYCERIDES"],
        "VLDL": ["VLDL", "VERY LOW DENSITY LIPOPROTEIN"],
        "ALANINE AMINOTRANSFERASE": ["ALT", "SGPT", "ALANINE TRANSAMINASE"],
        "ASPARTATE AMINOTRANSFERASE": ["AST", "SGOT", "ASPARTATE TRANSAMINASE"],
        "FATTY LIVER GRADE": ["FATTY LIVER", "FATTY LIVER GRADE"],
        "THYROID STIMULATING HORMONE": ["TSH", "THYROID STIMULATING HORMONE"],
        "THYROXINE": ["T4", "FREE T4", "TOTAL T4", "THYROXINE", "FT4"],
        "TRIIODOTHYRONINE": ["T3", "FREE T3", "TOTAL T3", "TRIIODOTHYRONINE", "FT3"],
        "REVERSE T3": ["RT3", "REVERSE T3"],
        "ANTI-TPO ANTIBODIES": ["ANTI-TPO", "ANTI-TPO ANTIBODIES", "ANTI THYROID PEROXIDASE"],
        "ANTI-TG ANTIBODIES": ["ANTI-TG", "ANTI-TG ANTIBODIES", "ANTI THYROGLOBULIN"],
        "TRAB / TSI": ["TRAB", "TSI", "THYROTROPIN RECEPTOR ANTIBODY", "THYROID STIMULATING IMMUNOGLOBULINS"],
        "FASTING BLOOD SUGAR": ["FBS", "FASTING BLOOD SUGAR"],
        "POSTPRANDIAL BLOOD SUGAR": ["PPBS", "POSTPRANDIAL BLOOD SUGAR"],
        "RANDOM BLOOD SUGAR": ["RBS", "RANDOM GLUCOSE", "RANDOM BLOOD SUGAR"],
        "HBA1C": ["HBA1C", "GLYCATED HEMOGLOBIN"],
        "GLUCOSE TOLERANCE TEST": ["GTT", "OGTT", "GLUCOSE TOLERANCE TEST"],
        "FRUCTOSAMINE": ["FRUCTOSAMINE"],
        "INSULIN (FASTING)": ["FASTING INSULIN", "INSULIN (FASTING)"],
        "C-PEPTIDE": ["C-PEPTIDE"],
        "HOMA-IR": ["HOMA-IR", "HOMA INDEX", "HOMEOSTATIC MODEL ASSESSMENT"]
    }
    
    test_name_upper = test_name.upper()
    
    for full_name, abbrevs in common_mappings.items():
        if full_name in test_name_upper or test_name_upper in abbrevs:
            variations.extend(abbrevs)
            variations.append(full_name)    
    variations.append(test_name)
    return list(set(variations))  


def _load_blood_report_json_data():
    global TEST_ABBREVIATIONS, FULL_BLOOD_REPORT_DATA
    try:
        with open(PREDEFINED_BLOOD_REPORT_JSON_PATH, 'r', encoding='utf-8') as f:
            FULL_BLOOD_REPORT_DATA = json.load(f)
        
        for panel_name, panel_data in FULL_BLOOD_REPORT_DATA.items():
            tests = panel_data.get("Tests", [])
            if isinstance(tests, list):
                for test_info in tests:
                    if isinstance(test_info, dict) and "Test Name" in test_info:
                        test_abbr = test_info["Test Name"]
                        TEST_ABBREVIATIONS[test_abbr.lower()] = {
                            **test_info,
                            "panel": panel_name  
                        }
                        
                        variations = _get_test_name_variations(test_abbr)
                        for variation in variations:
                            TEST_ABBREVIATIONS[variation.lower()] = {
                                **test_info,
                                "panel": panel_name
                            }

    except FileNotFoundError:
        logger.error(f"Error: {PREDEFINED_BLOOD_REPORT_JSON_PATH} not found.")
    except json.JSONDecodeError as e:
        logger.error(f"Error decoding JSON from {PREDEFINED_BLOOD_REPORT_JSON_PATH}: {e}")
    except Exception as e:
        logger.error(f"An unexpected error occurred loading blood report JSON: {e}")
_load_blood_report_json_data()


def classify_vital_signs(vital_signs_input, vital_signs_data, age=None, gender=None):
    results = {}
    current_age_group = get_age_group(age)
    for vital_name_input, value_input in vital_signs_input.items():
        standardized_vital_name = vital_name_input
        if "Heart Rate" in vital_name_input:
            standardized_vital_name = "Heart Rate"
        elif "Blood Pressure" in vital_name_input:
            standardized_vital_name = "Blood Pressure"
            if isinstance(value_input, str) and '/' in value_input:
                try:
                    systolic_str, diastolic_str = value_input.split('/')
                    systolic = float(systolic_str.strip())
                    diastolic = float(diastolic_str.strip())
                    value_input = {'systolic': systolic, 'diastolic': diastolic}
                except ValueError:
                    logging.warning("Invalid Blood Pressure input format: %s. Skipping.", value_input)
                    continue
            else:
                if isinstance(value_input, (int, float)):
                    value_input = {'systolic': float(value_input), 'diastolic': float(value_input)}
                else:
                    logging.warning("Invalid Blood Pressure input: %s. Skipping.", vital_name_input)
                    continue
        elif "Blood Glucose" in vital_name_input:
            standardized_vital_name = "Blood Glucose"
        elif "Blood Oxygen Saturation" in vital_name_input:
            standardized_vital_name = "Blood Oxygen Saturation"
        elif "Respiration Rate" in vital_name_input:
            standardized_vital_name = "Respiration Rate"
        elif "Body Temperature" in vital_name_input:
            standardized_vital_name = "Body Temperature"
        elif "Blood Ketones" in vital_name_input:
            standardized_vital_name = "Blood Ketones"
        elif "Body Fat" in vital_name_input:
            standardized_vital_name = "Body Fat %"
        vital_result = {}
        found_nutritional_guideline = False
        for vital_guideline in vital_signs_data.get('vital_signs', []):
            if vital_guideline['name'] == standardized_vital_name:
                for condition in vital_guideline['conditions']:
                    condition_range_str = condition.get('range', '').strip()
                    condition_status = condition.get('status', '').strip()
                    age_group_matches = current_age_group in condition_status
                    if age_group_matches:
                        if standardized_vital_name == "Body Fat %":
                            gender_in_status_matched = False
                            if gender:
                                if gender.lower() == 'male' and '(men' in condition_status.lower():
                                    gender_in_status_matched = True
                                elif gender.lower() == 'female' and '(women' in condition_status.lower():
                                    gender_in_status_matched = True
                            if gender_in_status_matched and is_in_range(value_input, condition_range_str, standardized_vital_name):
                                found_nutritional_guideline = True
                        else:
                            if is_in_range(value_input, condition_range_str, standardized_vital_name):
                                found_nutritional_guideline = True
                    if found_nutritional_guideline:
                        vital_result.update({
                            'nutritional_status_type': condition_status,
                            'totalCalories': condition.get('calories', ''),
                            'nutritionalConsiderations': condition.get('nutritional_considerations', ''),
                            'foodRecommendations': condition.get('food_recoms', [])
                        })
                        break
                if found_nutritional_guideline:
                    break
        if not found_nutritional_guideline:
            logging.warning("No matching condition found for %s: %s", vital_name_input, value_input)
            vital_result.update({
                'nutritional_status_type': "Unknown",
                'totalCalories': "Maintain",
                'nutritionalConsiderations': "Consult healthcare provider",
                'foodRecommendations': []
            })
        results[vital_name_input] = vital_result
    return results

def _find_best_test_match(test_name, threshold=85):
    test_name_clean = test_name.lower().strip()
    if test_name_clean in TEST_ABBREVIATIONS:
        return TEST_ABBREVIATIONS[test_name_clean], 100
    
    for key in TEST_ABBREVIATIONS.keys():
        if test_name_clean in key or key in test_name_clean:
            if len(test_name_clean) >= 2 and len(key) >= 2: 
                return TEST_ABBREVIATIONS[key], 85
    return None, 0

class BloodReportQueryTool(BaseTool):
    def __init__(self, llm_service: LLMService, db_tool: DatabasePersistenceTool):
        super().__init__(ToolType.BLOOD_REPORT_QUERY.value, "Answers specific questions about blood report tests, panels, and provides related food recommendations from the loaded blood report data.")
        self.llm_service = llm_service
        self.db_tool = db_tool
 
    async def _generate_dynamic_food_response(self, query: str, entity_name: str, food_rec_type: str,
                                            all_recommended: set, all_avoid: set,
                                            user_abnormal_test_ranges: dict, medical_conditions: list, allergies: list) -> str:
        food_context = {
            "recommended_foods": list(all_recommended),
            "foods_to_avoid": list(all_avoid),
            "allergies": list(allergies)
        }
        has_diabetes = False
        diabetes_instruction = ""
        
        if medical_conditions:
            for condition in medical_conditions:
                if 'diabetes' in condition.lower():
                    has_diabetes = True
                    break
        
        if has_diabetes:
            diabetes_instruction = """
        **SPECIAL DIABETES FOCUS**: Since you have diabetes, please include specific guidance about blood sugar testing:
        - Explain the importance of 2-hour post-meal blood sugar (post prandial blood sugar (PPBS ))testing
        - Guide them on 12-hour fasting blood sugar monitoring
        - Mention prediabetic ranges (RBS 140-199 mg/dL) and diabetic ranges (RBS ≥200 mg/dL)
        - Motivate them about how regular testing helps them understand their body's response to food
        - Keep the tone encouraging and empowering about taking control of their health
        """
        abnormal_tests_context = []
        for test_name, range_status in user_abnormal_test_ranges.items():
            if range_status != "normal":
                abnormal_tests_context.append(f"{test_name}: {range_status}")
        
        condition_prompts = []
        for condition in medical_conditions:
            for key, condition_prompt in MEDICAL_CONDITION_PROMPTS.items():
                if key.lower() in condition.lower() or condition.lower() in key.lower():
                    condition_prompts.append(f"""
        **Condition-Specific Nutrition Rules for {key}:**
        {condition_prompt}
        """)
                    break
        system_prompt = f"""You are Friska, an AI nutrition assistant. The user has asked a specific question about food recommendations based on their blood report.
 
        User's Query: "{query}"
        Entity/Focus: "{entity_name}"
        Recommendation Type: "{food_rec_type}"
 
        User's Abnormal Test Results: {', '.join(abnormal_tests_context) if abnormal_tests_context else 'None - all tests are normal'}

        {diabetes_instruction}

        Available Food Data:
        - Recommended Foods: {', '.join(food_context['recommended_foods']) if food_context['recommended_foods'] else 'None available'}
        - Foods to Avoid: {', '.join(food_context['foods_to_avoid']) if food_context['foods_to_avoid'] else 'None available'}
        - Allergies: {', '.join(food_context['allergies']) if food_context['allergies'] else 'None available'}

        {''.join(condition_prompts) if condition_prompts else ""}
        INSTRUCTIONS:
        1. Answer the user's SPECIFIC question contextually - if they ask for "light snacks", focus on light snack options
        2. If they ask about foods to avoid or allergies, focus on that
        3. Use the available food data but filter and present it according to their specific request
        4. If asking for snacks/specific meal types, suggest 3-5 specific items from the recommended foods that fit their request
        5. Keep responses concise and practical
        6. Always end with the standard medical disclaimer

        **Your Communication Style**:
        - Use warm and considerate language that makes them feel capable
        - Explain *why* behind advice, not just *what*
        - Make them feel like they're taking powerful steps toward better health
        - Be their knowledgeable friend who genuinely cares about their success

        Format your response naturally, don't just list everything. Be helpful and specific to their query."""
        user_prompt = f"""Based on my blood report analysis and the available food recommendations, please answer my question: "{query}"
 
        Available data context:
        - Recommended foods from my blood analysis: {food_context['recommended_foods']}
        - Foods I should avoid: {food_context['foods_to_avoid']}
 
        Please provide a helpful, specific answer to my question."""
 
        try:
            dynamic_response = await self.llm_service.query(
                prompt=user_prompt,
                system_prompt=system_prompt,
                max_tokens=600
            )
            return dynamic_response
        except Exception as e:
            logger.error(f"Error generating dynamic food response: {e}")
            return self._generate_fallback_response(query, all_recommended, all_avoid)
 
    def _generate_fallback_response(self, query: str, all_recommended: set, all_avoid: set) -> str:
        response_parts = []
       
        query_lower = query.lower()
        if any(word in query_lower for word in ["snack", "light", "quick", "small"]):
            snack_foods = [food for food in all_recommended if any(snack_word in food.lower()
                        for snack_word in ["nuts", "berries", "yogurt", "fruit", "seeds", "egg", "cucumber", "carrot"])]
            if snack_foods:
                response_parts.append(f"**Light snack suggestions based on your blood report:**")
                response_parts.append(f"{', '.join(snack_foods[:5])}")
        else:
            if all_recommended:
                response_parts.append(f"**Recommended foods:** {', '.join(list(all_recommended)[:8])}")
            if all_avoid:
                response_parts.append(f"**Foods to avoid:** {', '.join(list(all_avoid)[:8])}")
       
        response_parts.append(f"\nRemember to consult with your healthcare provider for personalized medical advice.")
        return "\n".join(response_parts)
 
    async def execute(self, query: str, constraints: Dict, last_agent_context: Dict, **kwargs) -> ToolResult:
        global FULL_BLOOD_REPORT_DATA
        if not FULL_BLOOD_REPORT_DATA:
            return ToolResult(success=False, data=None, error="I don't have any of your blood report data yet. Please upload or share your report so I can look into it for you.")
        medical_conditions = constraints.get("medical_conditions", [])
        allergies = constraints.get('allergies', [])
        user_blood_report_data = constraints.get("blood_report_analysis_data", {})
        token = kwargs.get("token")
 
        if not user_blood_report_data and token:
            db_result2 = await self.db_tool.execute(operation="get_all_blood_reports", token=token)
            if db_result2.success and db_result2.data and db_result2.data.get('data'):
                raw_data = db_result2.data['data']
                processed_translations = process_blood_report_translations(raw_data)
                combined_blood_report_data = {}
                combined_abnormal_test_ranges = {}
                combined_foods_to_avoid = []
                combined_recommended_foods = []
               
                for translation in processed_translations:
                    blood_data = translation.get("blood_report_data", {})
                    combined_blood_report_data.update(blood_data)
                   
                    abnormal_ranges = translation.get("abnormal_test_ranges", {})
                    combined_abnormal_test_ranges.update(abnormal_ranges)
                   
                    foods_to_avoid = translation.get("foods_to_avoid_from_report", [])
                    for food in foods_to_avoid:
                        if food not in combined_foods_to_avoid:
                            combined_foods_to_avoid.append(food)
                   
                    recommended_foods = translation.get("recommended_foods_from_report", [])
                    for food in recommended_foods:
                        if food not in combined_recommended_foods:
                            combined_recommended_foods.append(food)
               
                if combined_blood_report_data:
                    constraints.update({
                        "blood_report_analysis_data": combined_blood_report_data,
                        "abnormal_test_ranges": combined_abnormal_test_ranges,
                        "foods_to_avoid_from_report": combined_foods_to_avoid,
                        "recommended_foods_from_report": combined_recommended_foods
                    })
                    user_blood_report_data = combined_blood_report_data
                    user_abnormal_test_ranges = combined_abnormal_test_ranges
 
        user_abnormal_test_ranges = constraints.get("abnormal_test_ranges", {})
        if not user_blood_report_data:
            return ToolResult(success=True, data={"answer": "I don't have any of your blood report data saved. Please upload a report first so I can answer your questions about it."})
 
        system_prompt = """You are an expert assistant for blood reports and nutrition. Your task is to analyze the user's query and determine:
        1. If they are asking about a specific blood test (e.g., "What is WBC?", "Tell me about my cholesterol?").
        2. If they are asking about a specific panel (e.g., "What is in the Lipid Panel?", "Tell me about my CBC?").
        3. If they are asking for food recommendations related to blood reports or general health based on blood report data. When this query type is identified, you MUST also provide a list of relevant food items.
        4. If the query is general and doesn't fit the above, identify it as 'general_info'.
 
        Return ONLY a JSON object with the following structure:
        {
            "query_type": "test_info" | "panel_info" | "food_recommendations" | "general_info",
            "entity_name": "Name of the test or panel or food group (e.g., 'WBC', 'Lipid Panel', 'Fruits')",
            "food_recommendation_type": "recommended" | "avoid" | null # Only if query_type is 'food_recommendations'
        }
        Examples:
        - User: "What is my Hemoglobin?" -> {"query_type": "test_info", "entity_name": "Hemoglobin"}
        - User: "Tell me about the Thyroid Panel." -> {"query_type": "panel_info", "entity_name": "Thyroid Panel"}
        - User: "What foods should I eat for high cholesterol?" -> {"query_type": "food_recommendations", "entity_name": "Cholesterol", "food_recommendation_type": "recommended"}
        - User: "What foods should I avoid if my blood sugar is high?" -> {"query_type": "food_recommendations", "entity_name": "Blood Sugar", "food_recommendation_type": "avoid"}
        - User: "What are the common symptoms of low iron?" -> {"query_type": "general_info", "entity_name": "low iron symptoms"}
        - User: "Suggest me a high protein lunch based on my blood report." -> {"query_type": "food_recommendations", "entity_name": "blood report", "food_recommendation_type": "recommended"}
        """
       
        parsed_query_intent = await self.llm_service.query_json(
            prompt=f"Analyze this query: '{query}'",
            system_prompt=system_prompt,
            json_schema={
                "type": "object",
                "properties": {
                    "query_type": {"type": "string", "enum": ["test_info", "panel_info", "food_recommendations", "general_info"]},
                    "entity_name": {"type": "string"},
                    "food_recommendation_type": {"type": "string", "enum": ["recommended", "avoid", "both", "all", None]}
                },
                "required": ["query_type", "entity_name"]
            },
            max_tokens=500,
        )
 
        query_type = parsed_query_intent.get("query_type")
        entity_name = parsed_query_intent.get("entity_name", "").lower()
        food_rec_type = parsed_query_intent.get("food_recommendation_type")
        response_parts = []
       
        if query_type == "test_info":
            matched_test_info, _ = _find_best_test_match(entity_name)
            if matched_test_info:
                response_parts.append(f"Here is information about **{matched_test_info.get('Test Name', 'N/A')}**:")
                response_parts.append(f"- Full Form/Description: {matched_test_info.get('Full Form / Description', 'N/A')}")
                response_parts.append(f"- Normal Range: {', '.join(matched_test_info.get('Normal Range', ['N/A']))}")
                response_parts.append(f"- What It Indicates: {', '.join(matched_test_info.get('What It Indicates', ['N/A']))}")
                if matched_test_info.get('High Value Indicates'):
                    response_parts.append(f"- High Value Indicates: {', '.join(matched_test_info['High Value Indicates'])}")
                if matched_test_info.get('Low Value Indicates'):
                    response_parts.append(f"- Low Value Indicates: {', '.join(matched_test_info['Low Value Indicates'])}")
                if matched_test_info.get('Symptoms if Abnormal'):
                    response_parts.append(f"- Symptoms if Abnormal: {', '.join(matched_test_info['Symptoms if Abnormal'])}")
                if matched_test_info.get('Symptoms / Risks of Imbalance'):
                    response_parts.append(f"- Symptoms/Risks of Imbalance: {', '.join(matched_test_info['Symptoms / Risks of Imbalance'])}")
               
                user_test_status = None
                for test_key, test_data in user_blood_report_data.items():
                    if fuzz.ratio(test_key.lower(), matched_test_info.get('Test Name', '').lower()) > 85:
                        user_test_status = test_data.get('range')
                        break
 
                if user_test_status:
                    if user_test_status == 'normal':
                        response_parts.append(f"\nYour **{matched_test_info.get('Test Name', 'N/A')}** level is currently **normal** - no specific dietary changes needed!")
                    else:
                        response_parts.append(f"\nYour **{matched_test_info.get('Test Name', 'N/A')}** level is **{user_test_status}**.")
                       
                        diet_keywords = ["diet", "food", "eat", "nutrition", "should i eat", "avoid", "recommend"]
                        is_asking_for_diet = any(keyword in query.lower() for keyword in diet_keywords)
                       
                        if is_asking_for_diet:
                            response_parts.append("Here are some dietary recommendations based on your test results:")
                            self._add_food_recommendations_for_entity(entity_name, food_rec_type, response_parts)
                        else:
                            response_parts.append("Would you like specific dietary recommendations based on this test result?")
                else:
                    response_parts.append("\nI don't have your specific test results to provide personalized advice for this test. Please ensure your blood report has been analyzed by uploading it.")
 
            else:
                response_parts.append(f"I couldn't find detailed information for the test '{entity_name}'. Please ensure the test name is correct.")
 
        elif query_type == "panel_info":
            matched_panel_data = None
            panel_key_found = None
            highest_panel_score = 0
 
            for panel_key, panel_data in FULL_BLOOD_REPORT_DATA.items():
                current_panel_name = panel_data.get("Full Name", panel_key.replace('_', ' ')).lower()
               
                score_key = fuzz.ratio(entity_name, panel_key.lower())
                score_full_name = fuzz.ratio(entity_name, current_panel_name)
               
                current_score = max(score_key, score_full_name)
 
                if current_score > highest_panel_score and current_score >= 80:
                    highest_panel_score = current_score
                    matched_panel_data = panel_data
                    panel_key_found = panel_key
 
            if matched_panel_data:
                response_parts.append(f"Here is information about the **{matched_panel_data.get('Full Name', entity_name.title())}** panel:")
                tests_in_panel = [test.get("Test Name", "Unknown Test") for test in matched_panel_data.get("Tests", [])]
                if tests_in_panel:
                    response_parts.append(f"- Includes tests: {', '.join(tests_in_panel)}")
                else:
                    response_parts.append("- No specific tests listed for this panel.")
               
                is_panel_abnormal = False
                abnormal_tests_in_panel = []
                for test_name_in_panel in tests_in_panel:
                    if test_name_in_panel in user_abnormal_test_ranges:
                        is_panel_abnormal = True
                        abnormal_tests_in_panel.append(test_name_in_panel)
                    for user_test_key, test_data in user_blood_report_data.items():
                        if fuzz.ratio(user_test_key.lower(), test_name_in_panel.lower()) > 85 and test_data.get('range') != 'normal':
                            is_panel_abnormal = True
                            abnormal_tests_in_panel.append(user_test_key)
               
                if is_panel_abnormal:
                    response_parts.append(f"\nSome tests in your **{matched_panel_data.get('Full Name', entity_name.title())}** panel show abnormal values: {', '.join(set(abnormal_tests_in_panel))}")
                   
                    diet_keywords = ["diet", "food", "eat", "nutrition", "should i eat", "avoid", "recommend"]
                    is_asking_for_diet = any(keyword in query.lower() for keyword in diet_keywords)
                   
                    if is_asking_for_diet:
                        response_parts.append("Here are nutritional recommendations based on your panel results:")
                        nutritional_recs = matched_panel_data.get("NUTRITIONAL_RECOMMENDATION", [])
                        if nutritional_recs:
                            for rec_group in nutritional_recs:
                                if rec_group.get("Recommended Foods"):
                                    response_parts.append(f"- **Recommended {rec_group.get('Food Group / Topic', 'Foods')}:** {', '.join(rec_group['Recommended Foods'])}")
                                if rec_group.get("Foods to Avoid"):
                                    response_parts.append(f"- **Foods to Avoid ({rec_group.get('Food Group / Topic', 'Foods')}):** {', '.join(rec_group['Foods to Avoid'])}")
                       
                    else:
                        response_parts.append("Would you like specific dietary recommendations based on these abnormal results?")
                else:
                    response_parts.append(f"\nAll your test results within the **{matched_panel_data.get('Full Name', entity_name.title())}** panel are currently **normal** - great job!")
               
            else:
                response_parts.append(f"I couldn't find information for the panel '{entity_name}'. Please ensure the panel name is correct.")
       
        elif query_type == "food_recommendations":
            blood_report_keywords = [
                "blood report", "report", "my blood report", "test results", "blood test",
                "my report", "my test", "blood work", "lab report", "lab results",
                "my blood test", "my blood work", "my lab results", "my tests",
                "blood results", "test report", "my blood", "blood panel",
                "my blood panel", "blood values", "my values", "results", "my results"
            ]
           
            is_blood_report_query = False
            entity_lower = entity_name.lower().strip()
           
            for keyword in blood_report_keywords:
                if entity_lower == keyword or fuzz.ratio(entity_lower, keyword) > 85:
                    is_blood_report_query = True
                    break
           
            if not is_blood_report_query and entity_lower.startswith("my "):
                health_terms = ["cholesterol", "sugar", "hemoglobin", "vitamin", "thyroid", "liver", "kidney"]
                for term in health_terms:
                    if term in entity_lower:
                        is_blood_report_query = True
                        break
           
            if is_blood_report_query:
                all_recommended = set()
                all_avoid = set()
                processed_panels = set()
 
                for test_name, test_range in user_abnormal_test_ranges.items():
                    if test_range != "normal":
                        test_panel = None
                        for panel_key, panel_data in FULL_BLOOD_REPORT_DATA.items():
                            tests_in_panel = panel_data.get("Tests", [])
                            for test in tests_in_panel:
                                if fuzz.ratio(test_name.lower(), test.get("Test Name", "").lower()) > 85:
                                    test_panel = panel_key
                                    break
                            if test_panel:
                                break
                       
                        if test_panel and test_panel not in processed_panels:
                            processed_panels.add(test_panel)
                            panel_data = FULL_BLOOD_REPORT_DATA[test_panel]
                           
                            nutritional_recs = panel_data.get("NUTRITIONAL_RECOMMENDATION", [])
                            for rec_group in nutritional_recs:
                                if rec_group.get("Recommended Foods"):
                                    all_recommended.update(rec_group["Recommended Foods"])
                                if rec_group.get("Foods to Avoid"):
                                    all_avoid.update(rec_group["Foods to Avoid"])
                                if rec_group.get("Foods and Beverages to Avoid"):
                                    all_avoid.update(rec_group["Foods and Beverages to Avoid"])
                           
                if all_recommended or all_avoid:
                    dynamic_response = await self._generate_dynamic_food_response(
                        query, entity_name, food_rec_type,
                        all_recommended, all_avoid,
                        user_abnormal_test_ranges, 
                        medical_conditions = medical_conditions, allergies = allergies
                    )
                    response_parts.append(dynamic_response)
                else:
                    response_parts.append(
                        "Based on your blood report, your results are mostly normal! No major dietary changes are required. "
                        "Continue maintaining a balanced diet rich in fruits, vegetables, lean proteins, and whole grains."
                    )
 
            else:
                food_recs = self._add_food_recommendations_for_entity(entity_name, food_rec_type, None)
                if food_recs:
                    contextual_response = await self._generate_dynamic_food_response(
                        query, entity_name, food_rec_type,
                        set(), set(),  
                        {}, medical_conditions = medical_conditions, 
                        allergies = allergies
                    )
                    response_parts.append(contextual_response)
                else:
                    response_parts.append(f"I couldn't find specific food recommendations for '{entity_name}'.")
       
        elif query_type == "general_info":
            user_context = {
                "blood_report_data": user_blood_report_data,
                "abnormal_test_ranges": user_abnormal_test_ranges,
                "foods_to_avoid": constraints.get("foods_to_avoid_from_report", []),
                "recommended_foods": constraints.get("recommended_foods_from_report", [])
            }
           
            abnormal_summary = []
            normal_summary = []
            for test_name, test_data in user_blood_report_data.items():
                test_range = test_data.get('Normal Range', 'unknown')
                test_value = test_data.get('Extracted Value', 'N/A')
               
                if test_range != 'normal':
                    abnormal_summary.append(f"{test_name}: {test_value} ({test_range})")
                else:
                    normal_summary.append(f"{test_name}: {test_value} (normal)")
            user_blood_context = ""
            if abnormal_summary:
                user_blood_context += f"**User's Abnormal Test Results:**\n" + "\n".join([f"- {result}" for result in abnormal_summary]) + "\n\n"
           
            if normal_summary:
                user_blood_context += f"**User's Normal Test Results:**\n" + "\n".join([f"- {result}" for result in normal_summary[:10]]) + "\n\n"  
           
            if not user_blood_context:
                user_blood_context = "**User's Blood Report:** No specific test results available.\n\n"
           
            llm_system_prompt = f"""You are Friska, an AI nutrition assistant with access to comprehensive blood report data and the user's personal blood test results.
 
            IMPORTANT INSTRUCTIONS:
            1. Answer the user's general question about blood reports or health based on ONLY their personal test results shown below.
            2. Make your response HIGHLY PERSONALIZED by referencing the user's specific test values and their status (high/low/normal).
            3. If the question relates to any of their abnormal test results, provide specific guidance and explain what those values mean.
            4. If their relevant tests are normal, reassure them and provide maintenance advice.
            5. Focus on explaining their specific test results in simple terms and what they should know about their health status.
            6. Only include food/diet recommendations if the user explicitly asks about diet, food, or nutrition.
            7. Always maintain a supportive and informative tone.
            8. Connect their actual test values to the symptoms, risks, or health implications mentioned in their query.
            9. If the question is not related to any of their test results, politely indicate that and provide general guidance.
 
            USER'S PERSONAL BLOOD TEST RESULTS:
            {user_blood_context}
           
            Remember: Base your response primarily on their actual test results and values. Explain what their specific numbers mean for their health."""
           
            llm_user_prompt = f"""Based on my personal blood test results shown above, please answer this question: '{query}'
           
            Please make your response specific to my test results where relevant, and explain any connections between my values and the topic I'm asking about."""
           
            answer = await self.llm_service.query(llm_user_prompt, llm_system_prompt, max_tokens=800)
           
            if abnormal_summary:
                answer += "\n\n**Note:** Since you have some abnormal test results, please consult with your healthcare provider for personalized medical advice and treatment recommendations."
           
            response_parts.append(answer)        
       
        else:
            response_parts.append("I couldn't understand your query regarding the blood report. Could you please rephrase it?")
 
        return ToolResult(success=True, data={"answer": "\n".join(response_parts)})
 
    def _add_food_recommendations_for_entity(self, entity_name: str, food_rec_type: Optional[str] = None, response_parts: Optional[List[str]] = None) -> str:
        all_recommended = set()
        all_avoid = set()
 
        panel_matched = False
        for panel_key, panel_data in FULL_BLOOD_REPORT_DATA.items():
            if fuzz.ratio(entity_name, panel_key.lower()) > 85 or fuzz.ratio(entity_name, panel_data.get("Full Name", "").lower()) > 85:
                panel_matched = True
                nutritional_recs = panel_data.get("NUTRITIONAL_RECOMMENDATION", [])
                for rec_group in nutritional_recs:
                    if food_rec_type in ["recommended", "both", "all", None] and rec_group.get("Recommended Foods"):
                        all_recommended.update(rec_group["Recommended Foods"])
                    if food_rec_type in ["avoid", "both", "all", None]:
                        if rec_group.get("Foods to Avoid"):
                            all_avoid.update(rec_group["Foods to Avoid"])
                        if rec_group.get("Foods and Beverages to Avoid"):
                            all_avoid.update(rec_group["Foods and Beverages to Avoid"])
 
        if not panel_matched:
            matched_panel_data = None
            matched_panel_key = None
           
            for panel_key, panel_data in FULL_BLOOD_REPORT_DATA.items():
                tests_in_panel = panel_data.get("Tests", [])
               
                for test in tests_in_panel:
                    test_name = test.get("Test Name", "").lower()
                    test_score = fuzz.ratio(entity_name, test_name)
                   
                    if test_score >= 85:
                        matched_panel_data = panel_data
                        matched_panel_key = panel_key
                        logger.info(f"Found test '{test.get('Test Name')}' in panel '{panel_data.get('Full Name', panel_key)}' with score {test_score}")
                        break
               
                if matched_panel_data:
                    break
           
            if not matched_panel_data:
                for panel_key, panel_data in FULL_BLOOD_REPORT_DATA.items():
                    tests_in_panel = panel_data.get("Tests", [])
                   
                    for test in tests_in_panel:
                        test_name = test.get("Test Name", "").lower()
                        test_score = fuzz.partial_ratio(entity_name, test_name)
                       
                        if test_score >= 90:
                            matched_panel_data = panel_data
                            matched_panel_key = panel_key
                            logger.info(f"Found test '{test.get('Test Name')}' in panel '{panel_data.get('Full Name', panel_key)}' with partial score {test_score}")
                            break
                   
                    if matched_panel_data:
                        break
           
            if matched_panel_data:
                panel_matched = True
                nutritional_recs = matched_panel_data.get("NUTRITIONAL_RECOMMENDATION", [])
                for rec_group in nutritional_recs:
                    if food_rec_type in ["recommended", "both", "all", None] and rec_group.get("Recommended Foods"):
                        all_recommended.update(rec_group["Recommended Foods"])
                    if food_rec_type in ["avoid", "both", "all", None]:
                        if rec_group.get("Foods to Avoid"):
                            all_avoid.update(rec_group["Foods to Avoid"])
                        if rec_group.get("Foods and Beverages to Avoid"):
                            all_avoid.update(rec_group["Foods and Beverages to Avoid"])
 
        for panel_key, panel_data in FULL_BLOOD_REPORT_DATA.items():
            nutritional_recs = panel_data.get("NUTRITIONAL_RECOMMENDATION", [])
            for rec_group in nutritional_recs:
                if fuzz.ratio(entity_name, rec_group.get("Food Group / Topic", "").lower()) > 85:
                    if food_rec_type in ["recommended", "both", "all", None] and rec_group.get("Recommended Foods"):
                        all_recommended.update(rec_group["Recommended Foods"])
                    if food_rec_type in ["avoid", "both", "all", None]:
                        if rec_group.get("Foods to Avoid"):
                            all_avoid.update(rec_group["Foods to Avoid"])
                        if rec_group.get("Foods and Beverages to Avoid"):
                            all_avoid.update(rec_group["Foods and Beverages to Avoid"])
 
        matched_test_info, _ = _find_best_test_match(entity_name)
        if matched_test_info:
            if food_rec_type in ["recommended", "both", "all", None]:
                if matched_test_info.get("Recommended Foods"):
                    all_recommended.update(matched_test_info["Recommended Foods"])
                if matched_test_info.get("Recommended Food"):
                    all_recommended.update(matched_test_info["Recommended Food"])
           
            if food_rec_type in ["avoid", "both", "all", None]:
                if matched_test_info.get("Foods to Avoid"):
                    all_avoid.update(matched_test_info["Foods to Avoid"])
                if matched_test_info.get("Foods and Beverages to Avoid"):
                    all_avoid.update(matched_test_info["Foods and Beverages to Avoid"])
           
            if matched_test_info.get("panel"):
                panel_name = matched_test_info["panel"]
                panel_data = FULL_BLOOD_REPORT_DATA.get(panel_name)
               
                if panel_data:
                    logger.info(f"Found panel {panel_name} for test {entity_name}")
                    nutritional_recs = panel_data.get("NUTRITIONAL_RECOMMENDATION", [])
                    for rec_group in nutritional_recs:
                        if food_rec_type in ["recommended", "both", "all", None] and rec_group.get("Recommended Foods"):
                            all_recommended.update(rec_group["Recommended Foods"])
                        if food_rec_type in ["avoid", "both", "all", None]:
                            if rec_group.get("Foods to Avoid"):
                                all_avoid.update(rec_group["Foods to Avoid"])
                            if rec_group.get("Foods and Beverages to Avoid"):
                                all_avoid.update(rec_group["Foods and Beverages to Avoid"])
 
        if response_parts is not None:
            if all_recommended or all_avoid:
                if all_recommended and food_rec_type in ["recommended", "both", "all", None]:
                    response_parts.append(f"\n**Recommended Foods:**")
                    response_parts.append(", ".join(sorted(list(all_recommended))[:10]))
               
                if all_avoid and food_rec_type in ["avoid", "both", "all", None]:
                    response_parts.append(f"\n**Foods to Avoid:**")
                    response_parts.append(", ".join(sorted(list(all_avoid))[:10]))
               
                response_parts.append(f"\nRemember to consult with your healthcare provider for personalized medical advice.")
            else:
                response_parts.append(f"I couldn't find specific food recommendations for '{entity_name}'.")
        else:
            formatted_response = []
            if all_recommended or all_avoid:
                formatted_response.append(f"**Dietary recommendations for {entity_name.title()}:**")
               
                if all_recommended and food_rec_type in ["recommended", "both", "all", None]:
                    formatted_response.append(f"\n**Recommended Foods:**")
                    formatted_response.append(", ".join(sorted(list(all_recommended))[:10]))
               
                if all_avoid and food_rec_type in ["avoid", "both", "all", None]:
                    formatted_response.append(f"\n**Foods to Avoid:**")
                    formatted_response.append(", ".join(sorted(list(all_avoid))[:10]))
               
                formatted_response.append(f"\nRemember to consult with your healthcare provider for personalized medical advice.")
                return "\n".join(formatted_response)
            else:
                return f"I couldn't find specific food recommendations for '{entity_name}'."


class CalorieCalculatorTool(BaseTool):
    def __init__(self, data_service: DataService):
        super().__init__(ToolType.CALORIE_CALCULATOR.value, "Calculates calorie targets, BMI, TDEE, and meal totals")
        self.data_service = data_service

    def _extract_calorie_target(self, query: str) -> ToolResult:
        patterns = [r'(\d{3,5})\s*(?:kcal|calories|calorie|cal)\b', r'target\s*(?:of|is)?\s*(\d{3,5})\b', r'\b(?:around|about|approx|approximately)\s*(\d{3,5})\b', r'(\d{3,5})\s*daily\b', r'budget\s*(?:of|is)?\s*(\d{3,5})\b']
        for pattern in patterns:
            match = re.search(pattern, query.lower())
            if match:
                try: calorie_val_str = next(g for g in match.groups() if g and g.isdigit())
                except StopIteration: continue
                if calorie_val_str: return ToolResult(success=True, data=int(calorie_val_str))
        return ToolResult(success=True, data=None, error="No calorie target found in query.")

    def _calculate_bmi(self, weight_kg: float, height_cm: float) -> Optional[float]:
        if not weight_kg or not height_cm or weight_kg <= 0 or height_cm <= 0: return None
        return weight_kg / ((height_cm / 100) ** 2)

    def _calculate_tdee(self, gender: str, weight_kg: float, height_cm: float, age: float, activity_level: str) -> Optional[float]:
        if not all([gender, weight_kg, height_cm, age]) or any(v <= 0 for v in [weight_kg, height_cm, age]): return None
        bmr: float
        if gender.lower() == "male":
            bmr = (10 * weight_kg) + (6.25 * height_cm) - (5 * age) + 5
        elif gender.lower() == "female":
            bmr = (10 * weight_kg) + (6.25 * height_cm) - (5 * age) - 161
        else:
            bmr_male=(10*weight_kg)+(6.25*height_cm)-(5*age)+5
            bmr_female=(10*weight_kg)+(6.25*height_cm)-(5*age)-161
            bmr=(bmr_male+bmr_female)/2
        activity_factors = {"sedentary":1.2, "somewhat active":1.375, "moderately active":1.55, "very active":1.725}
        return bmr * activity_factors.get(activity_level.lower(), 1.2)

    async def execute(self, operation: str, **kwargs) -> ToolResult:
        if operation == "extract_target": return self._extract_calorie_target(kwargs.get('query', ''))
        elif operation == "calculate_profile_metrics":
            profile_data = kwargs.get('constraints', {}); weight, height, age, gender, activity = (profile_data.get(k) for k in ['weight_kg', 'height_cm', 'age', 'gender', 'activity_level'])
            metrics = {}
            if weight and height: metrics['bmi'] = self._calculate_bmi(weight, height)
            if gender and weight and height and age and activity: metrics['tdee'] = self._calculate_tdee(gender, weight, height, age, activity)
            if not metrics: return ToolResult(False, {}, error="Insufficient data for BMI/TDEE. Need age, gender, weight, height, activity.")
            return ToolResult(True, metrics)
        return ToolResult(False, error="Unknown calorie calculation operation")
    



def extract_text_from_pdf(pdf_path):
    text = ""
    try:
        with open(pdf_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            for page_num in range(len(reader.pages)):
                text += reader.pages[page_num].extract_text()
    except Exception as e:
        logger.error(f"Error extracting text from PDF {pdf_path}: {e}")
    return text

def extract_text_from_docx(docx_path):
    text_parts = []
    
    try:
        doc = Document(docx_path)
        
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())       
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_text = ""
                    for paragraph in cell.paragraphs:
                        if paragraph.text.strip():
                            cell_text += paragraph.text.strip() + " "
                    if cell_text.strip():
                        row_text.append(cell_text.strip())
                
                if row_text:
                    text_parts.append(" | ".join(row_text))
        
        for section in doc.sections:
            if section.header:
                for para in section.header.paragraphs:
                    if para.text.strip():
                        text_parts.append(f"[HEADER] {para.text.strip()}")
            if section.footer:
                for para in section.footer.paragraphs:
                    if para.text.strip():
                        text_parts.append(f"[FOOTER] {para.text.strip()}")
        
        logging.info(f"Extracted text from DOCX: {docx_path}")
        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"Error extracting text from DOCX {docx_path}: {e}")
        return ""

def extract_text_from_csv(csv_path):
    text = ""
    try:
        encodings = ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']
        
        for encoding in encodings:
            try:
                with open(csv_path, 'r', encoding=encoding) as file:
                    reader = csv.reader(file)
                    for row_num, row in enumerate(reader, 1):
                        row_text = " | ".join(str(cell).strip() for cell in row if str(cell).strip())
                        if row_text:
                            text += f"Row {row_num}: {row_text}\n"
                break
            except UnicodeDecodeError:
                continue
        
        logging.info(f"Extracted text from CSV: {csv_path}")
    except Exception as e:
        logger.error(f"Error extracting text from CSV {csv_path}: {e}")
    return text

def extract_text_from_txt(txt_path):
    text = ""
    try:
        encodings = ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']
        
        for encoding in encodings:
            try:
                with open(txt_path, 'r', encoding=encoding) as file:
                    text = file.read()
                break
            except UnicodeDecodeError:
                continue
        
        logging.info(f"Extracted text from TXT: {txt_path}")
    except Exception as e:
        logger.error(f"Error extracting text from TXT {txt_path}: {e}")
    return text

def extract_text_from_image_pytesseract(image_path):
    text = ""
    try:
        image = Image.open(image_path)
        text = pytesseract.image_to_string(image)
    except Exception as e:
        logger.error(f"Error extracting text from image {image_path}: {e}")
    return text

def extract_text_from_image(image_path):
    try:
        with open(image_path, "rb") as image_file:
            image_data = base64.b64encode(image_file.read()).decode("utf-8")

        headers = {
            "Authorization": f"Bearer {MISTRAL_API_KEY}",
            "Content-Type": "application/json"
        }

        payload = {
            "messages": [
                {"role": "system", "content": "You are a helpful assistant that extracts text and interprets information from images. Only output the extracted text."},
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:image/jpeg;base64,{image_data}"
                            }
                        },
                        {
                            "type": "text",
                            "text": "Extract all text from this image what it contains."
                        }
                    ]
                }
            ],
            "temperature": 0,
            "max_tokens": 2048
        }

        logger.debug(f"Sending request to mistral endpoint: {MISTRAL_API_ENDPOINT}")
        response = requests.post(MISTRAL_API_ENDPOINT, headers=headers, json=payload, timeout=25)
        response.raise_for_status()
        result = response.json()
        reply = result['choices'][0]['message']['content']
        return reply

    except requests.exceptions.RequestException as e:
        logger.error(f"Error during image text extraction (API request failed): {e}")
        if hasattr(e, 'response') and e.response is not None:
            logger.error(f"Response content: {e.response.text}")
        return None
    except Exception as e:
        logger.error(f"Unexpected error during image text extraction: {e}")
        return None

def _fallback_test_extraction(text_content):
    logger.info("Using fallback regex extraction method...") 
    extracted_values = []    
    patterns = [
        r'([A-Za-z\s]+):\s*([0-9.]+)\s*[A-Za-z/]*\s*\([0-9.\-\s]*\)',
        r'([A-Za-z\s]+)\s+([0-9.]+)\s+[A-Za-z/]*\s+[0-9.\-\s]*',
        r'([A-Za-z\s]+)\s*-\s*([0-9.]+)',
        r'([A-Za-z\s]+)\s+([0-9.]+)(?!\s*[A-Za-z])',
    ]    
    for pattern in patterns:
        matches = re.findall(pattern, text_content, re.IGNORECASE)
        for match in matches:
            if len(match) >= 2:
                test_name = match[0].strip()
                value = match[1].strip()
                
                if (len(test_name) > 2 and 
                    not any(word in test_name.lower() for word in ['page', 'report', 'lab', 'date', 'time', 'patient']) and
                    test_name.replace(' ', '').isalpha()):
                    
                    try:
                        numeric_value = float(value)
                        extracted_values.append({
                            "test_name_abbr": test_name,
                            "value": numeric_value
                        })
                    except ValueError:
                        extracted_values.append({
                            "test_name_abbr": test_name,
                            "value": value
                        })
    
    seen = set()
    unique_values = []
    for item in extracted_values:
        test_name = item["test_name_abbr"].lower()
        if test_name not in seen:
            seen.add(test_name)
            unique_values.append(item)
    
    logger.info(f"Fallback extraction found {len(unique_values)} test values")
    return unique_values
class BloodReportAnalyzerTool(BaseTool):
    def __init__(self, llm_service: LLMService):
        super().__init__(ToolType.BLOOD_REPORT_ANALYZER.value, "Analyzes uploaded blood reports to extract test results and nutritional recommendations.")
        self.llm_service = llm_service

    async def analyze_blood_report(self, file_path: str, current_symptoms: Optional[str] = None) -> dict:
        file_extension = os.path.splitext(file_path)[1].lower()
        raw_content_for_llm = ""
        if file_extension == ".pdf":
            raw_content_for_llm = extract_text_from_pdf(file_path)
        elif file_extension == ".docx":
            raw_content_for_llm = extract_text_from_docx(file_path)
        elif file_extension == ".csv":
            raw_content_for_llm = extract_text_from_csv(file_path)
        elif file_extension == ".txt":
            raw_content_for_llm = extract_text_from_txt(file_path)
        elif file_extension in [".jpg", ".jpeg", ".png"]:
            try:
                raw_content_for_llm = extract_text_from_image(file_path)
            except:
                raw_content_for_llm = extract_text_from_image_pytesseract(file_path)
        else:
            return {"error": f"Unsupported file type: {file_extension}"}

        if not raw_content_for_llm:
            return {"error": "Could not extract text from the uploaded file."}

        system_prompt = """You are an expert in parsing blood lab reports. Your task is to extract all blood test names and their corresponding values from the provided content along with the flag that whether value is normal, low or high.

        IMPORTANT INSTRUCTIONS:
        1. Extract the EXACT test names as they appear in the document (don't abbreviate or expand)
        2. For each test, provide the numerical or string value
        3. If a test has multiple values or ranges, extract the primary value
        4. Be very careful with test name spelling and capitalization
        5. Include common variations (e.g., "Hemoglobin", "Hb", "HGB" are all valid)
        6. Use only low, normal, or high for range. Strictly follow the document's reference range—mark any deviation, even slight, as low or high.

        Common test names to look for:
        - Total Cholesterol, Cholesterol
        - LDL Cholesterol, LDL, Low Density Lipoprotein
        - HDL Cholesterol, HDL, High Density Lipoprotein
        - Triglycerides, TG
        - VLDL
        - ALT, SGPT, Alanine Aminotransferase
        - AST, SGOT, Aspartate Aminotransferase
        - Fatty Liver Grade (Ultrasound based)
        - Thyroid Stimulating Hormone, TSH
        - Free T4, T4, Thyroxine
        - Free T3, T3, Triiodothyronine
        - Total T4
        - Total T3
        - Reverse T3 (rT3)
        - Anti-TPO Antibodies
        - Anti-TG Antibodies
        - TRAb / TSI
        - Fasting Blood Sugar (FBS)
        - Postprandial Blood Sugar (PPBS)
        - Random Blood Sugar (RBS)
        - HbA1c (Glycated Hemoglobin)
        - Glucose Tolerance Test (GTT / OGTT)
        - Fructosamine
        - Insulin (Fasting)
        - C-Peptide
        - HOMA-IR (Calculated)

        Return ONLY a valid JSON array with this exact format:
        [
            {"test_name_abbr": "EXACT_TEST_NAME_FROM_DOCUMENT", "value": numeric_value_or_string, "range": "low/normal/high"},
            {"test_name_abbr": "ANOTHER_TEST_NAME", "value": numeric_value_or_string, "range": "low/normal/high"}
        ]

        Example:
        [
            {"test_name_abbr": "Insulin (Fasting)", "value": 22.5, "range": "normal"},
            {"test_name_abbr": "Random Blood Sugar (RBS)", "value": 210, "range": "high"},
            {"test_name_abbr": "TSH", "value": 0.1, "range": "low"}
        ]"""

        user_prompt = f"""Blood Report Content:
        {raw_content_for_llm}"""

        extracted_values_json_str = await self.llm_service.query(
            prompt=user_prompt,
            system_prompt=system_prompt,
            max_tokens=2500
        )
        
        extracted_values = []
        if extracted_values_json_str:
            try:
                json_pattern = r'```(?:json)?\s*\n?(.*?)\n?```'
                match = re.search(json_pattern, extracted_values_json_str, re.DOTALL)
                
                if match:
                    json_content = match.group(1)
                else:
                    json_content = extracted_values_json_str

                extracted_values = json.loads(json_content)
            except json.JSONDecodeError as e:
                logger.error(f"Failed to parse extracted values JSON: {e}")
                logger.error(f"Raw response: {extracted_values_json_str}")
                
                logger.info("Attempting fallback extraction using regex...")
                extracted_values = _fallback_test_extraction(raw_content_for_llm)
        
        if not extracted_values:
            logger.warning("No test values extracted. Attempting fallback extraction...")
            extracted_values = _fallback_test_extraction(raw_content_for_llm)
            
        
        if not extracted_values:
            return {"error": "Could not extract test values from the blood report. Please ensure the content is clear or try a different file."}

        analyzed_blood_report_data = {}
        detected_panels = set()
        matched_test_names = []
        match_details = []
        abnormal_test_ranges = {}
        logging.info(f"extracted values ---- {extracted_values}")
        for extracted_test in extracted_values:
            test_name_from_report = str(extracted_test.get("test_name_abbr", "")).strip()
            test_value = extracted_test.get("value")
            test_range = extracted_test.get("range")
            if test_range in ['low', 'high']:
                abnormal_test_ranges[test_name_from_report] = {
                    "value": test_value,
                    "status": test_range
                }
            
            if not test_name_from_report:
                continue
                
            matched_test_info, match_score = _find_best_test_match(test_name_from_report)
            if matched_test_info:
                full_test_name = matched_test_info.get("Test Name")
                panel_name = matched_test_info.get("panel")
                
                match_details.append({
                    "original": test_name_from_report,
                    "matched": full_test_name,
                    "panel": panel_name,
                    "score": match_score
                })
                
                if panel_name:
                    detected_panels.add(panel_name)
                    matched_test_names.append(full_test_name)                  
                if full_test_name not in analyzed_blood_report_data:
                    analyzed_blood_report_data[full_test_name] = {
                        "Test Name": full_test_name,
                        "Full Form / Description": matched_test_info.get("Full Form / Description"),
                        "Normal Range": matched_test_info.get("Normal Range"),
                        "Extracted Value": test_value,
                        "What It Indicates": matched_test_info.get("What It Indicates"),
                        "High Value Indicates": matched_test_info.get("High Value Indicates"),
                        "Low Value Indicates": matched_test_info.get("Low Value Indicates"),
                        "Symptoms if Abnormal": matched_test_info.get("Symptoms if Abnormal"),
                        "Symptoms / Risks of Imbalance": matched_test_info.get("Symptoms / Risks of Imbalance"),
                        "Panel": panel_name
                    }
            else:
                logger.warning(f"Could not match test name: '{test_name_from_report}'")

        all_foods_to_avoid = set()
        all_recommended_foods = set()

        for panel_name in detected_panels:
            panel_data = FULL_BLOOD_REPORT_DATA.get(panel_name, {})
            nutritional_entries = panel_data.get("NUTRITIONAL_RECOMMENDATION", [])
            
            if isinstance(nutritional_entries, list):
                for entry in nutritional_entries:
                    if isinstance(entry, dict):
                        foods_to_avoid = entry.get("Foods to Avoid", [])
                        if isinstance(foods_to_avoid, list):
                            all_foods_to_avoid.update(foods_to_avoid)
                        
                        recommended_foods = entry.get("Recommended Foods", [])
                        if isinstance(recommended_foods, list):
                            all_recommended_foods.update(recommended_foods)
        
        final_analysis_result = {
            "blood_report_data": analyzed_blood_report_data,
            "detected_panels": list(detected_panels),
            "matched_test_names": matched_test_names,
            "relevant_panels_used": list(detected_panels),  
            "foods_to_avoid_from_report": list(all_foods_to_avoid),
            "recommended_foods_from_report": list(all_recommended_foods),
            "abnormal_test_ranges": abnormal_test_ranges,
            "match_details": match_details,
            "extraction_debug": {
                "total_extracted": len(extracted_values),
                "matched_tests": len(analyzed_blood_report_data),
                "detected_panels_count": len(detected_panels),
                "nutritional_sections_used": len(detected_panels)
            }
        }
        logger.info(f"Blood report analysis complete. Detected panels: {list(detected_panels)}")
        print("\n--- FINAL BLOOD REPORT ANALYSIS JSON ---")
        print(json.dumps(final_analysis_result, indent=2))
        print("-------------------------------------\n")
        
        return final_analysis_result

    async def execute(self, file_path: str, current_symptoms: Optional[str] = None, **kwargs) -> ToolResult:
        if not file_path:
            return ToolResult(success=False, data=None, error="File path for blood report is missing.")
        
        token = kwargs.get("token")
        user_id = kwargs.get("user_id") 

        analysis_result = await self.analyze_blood_report(file_path, current_symptoms)

        if "error" in analysis_result:
            return ToolResult(success=False, data=None, error=analysis_result["error"])
        logger.info("Blood report analysis complete. Passing payload back to the main process.")
        updated_constraints_for_session = {
            "blood_report_analysis_data": analysis_result.get("blood_report_data", {}),
            "abnormal_test_ranges": analysis_result.get("abnormal_test_ranges", {}),
            "foods_to_avoid_from_report": analysis_result.get("foods_to_avoid_from_report", []),
            "recommended_foods_from_report": analysis_result.get("recommended_foods_from_report", [])
        }

        summary_parts = []
        detected_panels = analysis_result.get("detected_panels", [])
        abnormal_tests = analysis_result.get("abnormal_test_ranges", {})
        
        if detected_panels:
            summary_parts.append(f"I've analyzed your report and identified the following panels: {', '.join(detected_panels)}.")
        if abnormal_tests:
            summary_parts.append(f"I've noted abnormal ranges for: {', '.join(abnormal_tests.keys())}.")
        else:
            summary_parts.append("All extracted test values appear to be within normal ranges.")
        
        summary_parts.append("\nI now have this context for our conversation. How can I help you with these results?")
        
        final_answer = " ".join(summary_parts)

        return ToolResult(
            success=True,
            data={
                "answer": final_answer,
                "updated_constraints": updated_constraints_for_session
            },
            metadata={
                "blood_report_analysis_payload": analysis_result,
                "blood_report_analysis_data": analysis_result.get("blood_report_data", {}),
                "abnormal_test_ranges": analysis_result.get("abnormal_test_ranges", {}),
                "foods_to_avoid_from_report": analysis_result.get("foods_to_avoid_from_report", []),
                "recommended_foods_from_report": analysis_result.get("recommended_foods_from_report", [])
            }
        )
class PanelQueryTool(BaseTool):
    def __init__(self, llm_service: 'LLMService'):
        super().__init__(
            ToolType.PANEL_QUERY.value,
            "Answers specific queries about general health panels, nutritional recommendations, "
            "and recipes from the panel.json file, not specific to an uploaded blood report."
        )
        self.llm_service = llm_service
 
    async def execute(self, query: str, chat_history: List[Dict], constraints: Dict, **kwargs) -> ToolResult:
        """
        Executes the panel query by leveraging an LLM to provide answers
        based on the loaded panel.json data.
 
        Args:
            query (str): The user's specific question about the panel data.
            chat_history (List[Dict]): The current chat history for context.
            constraints (Dict): User profile constraints (optional, for personalization).
 
        Returns:
            ToolResult: A result object containing the LLM's answer or an error.
        """
        if not FULL_PANEL_DATA:
            return ToolResult(
                success=False,
                data=None,
                error="Panel data is not available. Please ensure the panel.json file is loaded correctly."
            )
 
        panel_data_json_str = json.dumps(FULL_PANEL_DATA, indent=2)
 
        system_prompt = f"""You are Friska, a highly knowledgeable and empathetic AI health assistant specializing in providing information from health panel data, nutritional recommendations, and recipes. Your goal is to provide clear, concise, and actionable answers based on the provided panel data and your own general health knowledge.
 
        **User Profile for Context:**
        {json.dumps(constraints, indent=2, default=serialize_data)}
 
        **Available Panel Data (for your reference):**
        {panel_data_json_str}
 
        **Your Task:**
        Based on the user's query and the comprehensive panel data provided above, answer the question.
        - If the user asks about a specific test (e.g., "What is Total Cholesterol?"), provide its full description, normal range, what it indicates, and symptoms/risks of imbalance.
        - If the user asks about a panel (e.g., "Tell me about the Lipid Panel?"), summarize the tests within that panel.
        - If the user asks about nutritional recommendations (e.g., "What foods are recommended for thyroid health?"), provide relevant food groups, recommended foods, and why they help.
        - If the user asks for recipes (e.g., "Do you have any sugar-friendly recipes?"), provide the recipe name, ingredients, quick method, and why it's good for the condition.
        - **CRITICAL: If the user mentions a specific test value being high or low (e.g., "My cholesterol is 300 high", "My TSH is low"), follow this structured response:**
            1.  ** Begin with a calm and supportive acknowledgment. **Do not use phrases like "I'm sorry to hear that."**
            - **Example:** "I see you're asking about your high cholesterol reading. This is a great starting point for making some positive changes, and nutrition can be a powerful tool here."
            2.  ** State which test the user is referring to and the panel it belongs to.
            3.  **You can avoid these foods items:** List 3-5 specific food items to limit, explaining briefly why.
            4.  **I recommend these foods items:** List 3-5 specific food items to include, explaining briefly why.
            5.  **You can eat these food items:** Provide a numbered list of **10 distinct food items** to add to their meals to help manage the condition, with a brief benefit for each.
            6.  ** Conclude your response with the exact question: "How can I assist you further?"
        - Maintain a helpful and encouraging tone.
        - Ensure the response is comprehensive and directly answers the user's query using the provided data and your knowledge.
        """
 
        user_prompt = f"User's query: '{query}'"
 
        try:
            llm_answer = await self.llm_service.query(
                prompt=user_prompt,
                system_prompt=system_prompt,
                chat_history=chat_history,
                max_tokens=2000, 
                temperature=0.4
            )
 
            if not llm_answer:
                return ToolResult(success=False, data=None, error="I couldn't generate a response for your panel query at this time.")
            awaiting_food_add_confirmation = "Should I add these food items to your meal plan?" in llm_answer
 
            return ToolResult(
                success=True,
                data={"answer": llm_answer},
                metadata={"awaiting_food_add_confirmation": awaiting_food_add_confirmation}
            )
 
        except Exception as e:
            import traceback
            traceback.print_exc()
            return ToolResult(success=False, data=None, error=f"An error occurred while processing your panel query: {e}")

class DiseaseAdvisorTool(BaseTool):
    def __init__(self, data_service: DataService, llm_service: LLMService):
        super().__init__(
            ToolType.DISEASE_ADVISOR.value,
            "Provides disease-specific dietary recommendations using an LLM."
        )
        self.data_service = data_service
        self.llm_service = llm_service

    def _get_known_medical_conditions(self) -> List[str]:
        return [
            'Acanthosis Nigricans', 'Acne', 'Addison’s Disease', 'ADHD', 'Alcohol Use', 'Alzheimer’s Disease',
    'Ankylosing Spondylitis', 'Anorexia, Bulimia', 'Anxiety', 'Anxiety Disorders', 'Appetite & Nutritional Needs',
    'Arrhythmia', 'Arrhythmias', 'Arthritis (Osteoarthritis, Rheumatoid Arthritis)', 'Asthma',
    'Autism Spectrum Nutrition Support', 'Bipolar Disorder', 'Bladder Cancer', 'Brain Tumors', 'Breast Cancer',
    'Bronchitis', 'Cervical Cancer', 'Cervical Spondylosis', 'Chickenpox', 'Childhood Obesity',
    'Chikungunya', 'Chronic Kidney Disease (CKD)', 'Chronic Obstructive Pulmonary Disease (COPD)', 'Cirrhosis',
    'Cognitive Decline / Dementia (Alzheimers)', 'Colorectal Cancer', 'Congestive Heart Failure', 'Constipation',
    'Coronary Artery Disease', 'Coronary Artery Disease (CAD)', 'COVID-19', 'COVID-19 Recovery', 'Cushing’s Syndrome',
    'Deep Vein Thrombosis (DVT)', 'Dementia', 'Dengue', 'Dengue / Typhoid Recovery', 'Depression',
    'Dermatological Infections (Fungal, Bacterial)', 'Developmental Disorders', 'Diabetes (Type 1, Type 2, Gestational)',
    'Diabetic Ketoacidosis (DKA) Recovery', 'Dialysis Support', 'Diarrheal Diseases', 'Disc Herniation',
    'During Chemotherapy', 'Eating Disorders (Anorexia, Bulimia)', 'Eating Disorders Recovery', 'Eczema',
    'Encephalitis', 'Epilepsy', 'Epilepsy (Ketogenic Diet)', 'Failure to Thrive', 'Fatty Liver Disease (NAFLD/NASH)',
    'Fibromyalgia', 'Fractures', 'G6PD Deficiency', 'Gallbladder Disease', 'Gallstones',
    'Gastroesophageal Reflux Disease (GERD)', 'Gastritis', 'Glomerulonephritis', 'Gout',
    'Heart Failure', 'Hepatitis', 'Hepatitis (A, B, C)', 'Hepatitis E', 'High Cholesterol (Hyperlipidemia)', 'HIV/AIDS',
    'Hypertension', 'Hypertension (High Blood Pressure)', 'Hypoglycemia', 'Hypothyroidism / Hyperthyroidism',
    'Inflammatory Bowel Disease (Flare-up)', 'Inflammatory Bowel Disease (IBD)', 'Inflammatory Bowel Disease (Remission)',
    'Influenza', 'Insomnia', 'Insomnia / Sleep Apnea', 'Insulin Resistance', 'Interstitial Lung Disease',
    'Irritable Bowel Syndrome (IBS)', 'Kidney Stones', 'Lactation Support', 'Leukemia',
    'Low Back Pain', 'Lung Cancer', 'Lupus', 'Lymphoma', 'Malaria', 'Measles', 'Meningitis', 'Menopause',
    'Metabolic Syndrome', 'Migraine', 'Migraines', 'Multiple Sclerosis', 'Myocardial Infarction (Heart Attack) Recovery',
    'Neuropathy', 'Obesity', 'Obsessive Compulsive Disorder (OCD)', 'Osteoarthritis', 'Osteoporosis',
    'Ovarian Cancer', 'PCOS', 'Pancreatic Cancer', 'Parkinson’s Disease', 'Peptic Ulcer Disease', 'Perimenopause',
    'Peripheral Artery Disease', 'PMS / PMDD', 'Pneumonia', 'Polycystic Ovary Syndrome (PCOS)',
    'Post-Traumatic Stress Disorder (PTSD)', 'Postpartum Recovery', 'Post-treatment Recovery',
    'Pregnancy (Trimester Support)', 'Prenatal & Perinatal Conditions', 'Prostate Cancer', 'Prostate Enlargement (BPH)',
    'Psoriasis', 'Pulmonary Embolism', 'Pulmonary Hypertension', 'Pyelonephritis (Kidney Infection)',
    'Rheumatic Heart Disease', 'Rheumatoid Arthritis', 'Schizophrenia', 'Sedentary Lifestyle',
    'Sexually Transmitted Infections (STIs)', 'Sickle Cell Disease', 'Sleep Apnea',
    'Smoking / Alcohol Use', 'Stevens-Johnson Syndrome (SJS) Recovery', 'Stomach Cancer',
    'Stress / Burnout', 'Stroke Recovery', 'Stroke Recovery Support', 'Substance Use Disorder Recovery',
    'Substance Use Disorders', 'Thalassemia', 'Tuberculosis (TB)', 'Tuberculosis (Nutrition Support)',
    'Type 1 Diabetes Mellitus', 'Type 2 Diabetes Mellitus', 'Typhoid Fever', 'Undernutrition / Malnutrition',
    'Urinary Tract Infection (UTI)', 'Vaccine-preventable Diseases (supportive care)', 'Vitiligo',
    'Work-Related Health Issues'
        ]

    async def execute(self, diseases: List[str], constraints: Dict, **kwargs) -> ToolResult:
        """
        Generates a dynamic, LLM-powered response for disease-related or symptom-related nutritional queries.

        Args:
            diseases (List[str]): A list of diseases or symptoms from the user's profile or query.
            constraints (Dict): The user's current profile for context.
            **kwargs: Additional keyword arguments.

        Returns:
            ToolResult: A result object containing the LLM-generated answer or an error message.
        """
        if not diseases:
            return ToolResult(
                success=True,
                data={"answer": "To give you the best advice, I need to know what medical condition or symptom you're asking about. For general wellness, a balanced diet rich in fruits, vegetables, and whole grains is a great start!"}
            )

        disease_knowledge_base = json.dumps(self.data_service.disease_data, indent=2)
        user_profile_context = json.dumps(constraints, indent=2, default=serialize_data)
        user_query_context = ", ".join(diseases)

        system_prompt = f"""
        You are Friska, an expert and empathetic AI nutrition assistant. Your primary goal is to provide safe, targeted advice.

        **TONE AND PERSONA (CRITICAL)**
        Your tone must be calm, reassuring, and empowering. Acknowledge the user's condition without expressing pity.
        - **DO NOT USE** phrases like "I'm sorry to hear that," or "That sounds difficult."
        - **INSTEAD, USE** proactive and supportive language. Start your response by acknowledging their situation and immediately offering gentle encouragement.
        - **Example Tone:** If a user mentions thyroid issues, a good opening is: "I noticed you are experiencing thyroid health issues. Please don’t worry — with the right care and a few nourishing habits, your thyroid can feel more supported. A few gentle food and lifestyle habits can really support your thyroid."
        - Maintain this supportive and positive tone throughout your response.

        **USER PROFILE (for context and personalization):**
        {user_profile_context}

        **NUTRITIONAL KNOWLEDGE BASE (for reference):**
        {disease_knowledge_base}

        **MANDATORY RESPONSE INSTRUCTIONS:**
        1.  **Analyze the Query:** The user is asking for dietary advice related to: "{user_query_context}".
        2.  **Acknowledge and Reassure First:** Begin your response using the empowering tone described above.
        3.  **Prioritize the Immediate Symptom:** Your primary focus MUST be to provide advice for the user's stated symptom first.
        4.  **Cross-Reference with Profile:** After formulating advice for the symptom, you MUST review the **USER PROFILE**. Ensure your recommendations are safe and appropriate given their existing medical conditions and allergies. Acknowledge this cross-referencing naturally.
        5.  **Generate a Response:** Create a CONCISE response. STRICT LIMIT: 120-150 words maximum. Be direct and actionable.
        6.  **Content Requirements:**
            - Provide 3-5 specific dietary recommendations for the symptom, modified by their profile.
            - Use short bullet points. Briefly explain *why* in 1 sentence per point.
            - Do NOT give medical advice, prescribe treatments, or suggest medication.
            - Do NOT repeat generic diet advice. Focus on condition-specific foods.
        7.  **CRITICAL SAFETY DISCLAIMER (MANDATORY):** You MUST conclude your entire response with the following exact text:
             "This information is for general nutritional guidance and is not a substitute for a medical diagnosis. For a comprehensive treatment plan and to address any critical health concerns, it is essential that you contact a doctor or another qualified healthcare physician."
             """

        # ✨ SMART PERSONA INJECTION: Append domain-specific context if provided
        persona_context = kwargs.get('persona_context')
        if persona_context:
            system_prompt += f"\n\n🎯 **ADAPTIVE PERSONA GUIDANCE:**\n{persona_context}"
            logger.info("✨ Applied smart persona injection to DiseaseAdvisorTool system prompt")

        user_prompt = f"Please provide personalized dietary advice for my current symptom: {user_query_context}."

        try:
            llm_answer = await self.llm_service.query(
                prompt=user_prompt,
                system_prompt=system_prompt,
                max_tokens=500,  
                temperature=0.5
            )

            if not llm_answer:
                return ToolResult(
                    success=False,
                    data=None,
                    error="I was unable to generate advice for that condition at this time. Please try again later."
                )

            return ToolResult(success=True, data={"answer": llm_answer})

        except Exception as e:
            logger.error(f"Error in DiseaseAdvisorTool LLM call: {e}")
            return ToolResult(
                success=False,
                data=None,
                error="I encountered an unexpected issue while preparing your advice. Please try again."
            )

class VitalAdvisorTool(BaseTool):
    def __init__(self, data_service: DataService, llm_service: LLMService, vital_units: Dict):
        super().__init__(ToolType.VITAL_ADVISOR.value, "Provides nutritional advice based on vital signs.")
        self.data_service = data_service
        self.llm_service = llm_service
        self.vital_units = vital_units

    async def execute(self, query: str, chat_history: List[Dict], constraints: Dict, **kwargs) -> ToolResult:
        age = constraints.get('age')
        gender = constraints.get('gender')

        extract_schema = {
            "type": "object",
            "properties": {
                "vital_name": {"type": "string", "nullable": True},
                "vital_value": {"type": ["number", "string"], "nullable": True}
            }
        }
        extract_system_prompt = """Extract a single vital sign name and its numerical value from the user's query.
        If a vital is explicitly mentioned with a value, provide it.
        For "Blood Pressure", provide the value as a string "Systolic/Diastolic" (e.g., "120/80").
        For "Body Fat %", ensure the '%' is not part of the number.
        If no vital and value are explicitly mentioned, return null for both.
        Example outputs:
        - {"vital_name": "Heart Rate", "vital_value": 150}
        - {"vital_name": "Blood Pressure", "vital_value": "130/85"}
        - {"vital_name": "Body Temperature", "vital_value": 98.6}
        - {"vital_name": null, "vital_value": null}
        """

        extracted_vital_data = await self.llm_service.query_json(
            prompt=f"Extract vital sign and value from: '{query}'",
            system_prompt=extract_system_prompt,
            json_schema=extract_schema,
            chat_history=chat_history
        )

        specific_vital_name = extracted_vital_data.get('vital_name')
        specific_vital_value = extracted_vital_data.get('vital_value')

        if specific_vital_name and specific_vital_value is not None:
            vitals_to_classify = {specific_vital_name: specific_vital_value}
        else:
            vitals_to_classify = constraints.get('vitals_numeric', {})

        if not vitals_to_classify:
            return ToolResult(
                success=False,
                error="I don't have your vital signs data to provide advice. Please update your profile with your vital signs or ask about a specific reading."
            )

        classification_results = classify_vital_signs(
            vital_signs_input=vitals_to_classify,
            vital_signs_data=self.data_service.vitals_data,
            age=age,
            gender=gender
        )

        response_parts = []
        llm_response_prompt_parts = []

        for vital_name, result in classification_results.items():
            status = result.get('nutritional_status_type', 'Unknown')
            calories_advice = result.get('totalCalories', 'No changes')
            nutritional_considerations = result.get('nutritionalConsiderations', 'Consult healthcare provider')
            food_recs = result.get('foodRecommendations', [])

            vital_value_display = vitals_to_classify.get(vital_name)
            unit = self.vital_units.get(vital_name, "")
            if vital_name == "Blood Pressure" and isinstance(vital_value_display, dict):
                value_with_unit = f"{vital_value_display.get('systolic')}/{vital_value_display.get('diastolic')} {unit}"
            else:
                value_with_unit = f"{vital_value_display} {unit}".strip()

            llm_response_prompt_parts.append(f"Vital: {vital_name}")
            llm_response_prompt_parts.append(f"Current Value: {value_with_unit}")
            llm_response_prompt_parts.append(f"Status: {status}")
            llm_response_prompt_parts.append(f"Nutritional Considerations: {nutritional_considerations}")
            if calories_advice and calories_advice != "No changes in Total Daily Energy Expenditure (TDEE)":
                llm_response_prompt_parts.append(f"Calorie Advice: {calories_advice}")
            if food_recs:
                llm_response_prompt_parts.append("Food Recommendations:")
                for rec in food_recs:
                    llm_response_prompt_parts.append(f"- {rec.get('name')} ({rec.get('quantity')}): {rec.get('how_it_helps')}")
            llm_response_prompt_parts.append("\n---")

        llm_system_prompt = f"""You are Friska, a caring and knowledgeable AI nutritionist. Your goal is to provide immediate, clear, and actionable dietary advice based on the user's vital signs and your existing knowledge. Your response MUST be in English.

**Tone and Persona (CRITICAL):**
- Acknowledge the user's reading in a supportive, non-alarming way.
- **DO NOT USE** phrases like "I'm sorry about that reading" or "That's a high number."
- **DO NOT USE** phrases like "Thank you for sharing that."
- **INSTEAD, USE** encouraging and proactive language that focuses on positive action.
- **Example:** If blood pressure is elevated, a good opening is: "I see your blood pressure reading is elevated. This is valuable information, and we can explore some nutritional steps that are known to support healthy blood pressure levels."
- **BE CONCISE:** Do not be verbose. Get straight to the point with clear, actionable advice.
Based on the provided vital sign classifications and food recommendations:
1.  **Start with a direct, conversational assessment** of the vital sign(s).
2.  **Explain the status** (e.g., "Your Heart Rate is elevated...").
3.  **Provide the nutritional considerations.**
4.  **Immediately list the recommended food items** for that condition. Explain briefly *why* each food is helpful (from the provided 'how_it_helps' information).
5.  If a vital is normal, state that clearly and positively.
6.  Always end with a gentle reminder to consult a healthcare professional for medical advice.
7.  **If multiple vitals are provided, address each one in turn** following the same structure.
8.  **DO NOT** generate a full daily meal plan (Breakfast, Lunch, Dinner, Snacks) or suggest meal plan adjustments. Focus ONLY on specific food items and dietary tips relevant to the vital sign.
User Profile: {json.dumps(constraints, indent=2, default=serialize_data)}
"""
        
        # ✨ SMART PERSONA INJECTION: Append domain-specific context if provided
        persona_context = kwargs.get('persona_context')
        if persona_context:
            llm_system_prompt += f"\n\n🎯 **ADAPTIVE PERSONA GUIDANCE:**\n{persona_context}"
            logger.info("✨ Applied smart persona injection to VitalAdvisorTool system prompt")
        
        llm_user_prompt = f"Please provide immediate nutritional advice for the following vital sign(s):\n" + "\n".join(llm_response_prompt_parts)

        try:
            llm_final_response = await self.llm_service.query(
                prompt=llm_user_prompt,
                system_prompt=llm_system_prompt,
                chat_history=chat_history,
                max_tokens=3500
            )
            return ToolResult(
                success=True,
                data={"answer": llm_final_response},
                metadata={"awaiting_meal_plan_confirmation_for_vitals": True}
            )
        except Exception as e:
            print(f"Error during LLM vital advice generation: {e}")
            return ToolResult(
                success=False,
                error="I encountered an issue while generating advice for your vital signs. Please try again."
            )        

with open(os.path.join(BASE_DIR, "datasets", "Medical_Condition_Food_Prompt.json"), "r") as f:
    MEDICAL_CONDITION_PROMPTS = json.load(f)
