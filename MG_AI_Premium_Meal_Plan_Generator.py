import streamlit as st
import requests
import json
from docx import Document
from io import BytesIO
import re
from typing import Optional, Dict
import os
import tempfile
import base64
import binascii
from datetime import datetime, timedelta
from audio_recorder_streamlit import audio_recorder
from helpers.meal_plan_component import display_meal_plan, MealPlanRenderer
# BACKEND_API_URL = "http://127.0.0.1:8001/process-query"
BACKEND_API_URL = "http://127.0.0.1:7000/chat"
# Generates a meal plan directly from the frontend-collected profile, bypassing DB profile lookup/auth.
MEAL_PLAN_FROM_PROFILE_API_URL = "http://127.0.0.1:7000/generate-meal-plan-from-profile"

FITNESS_GOAL_OPTIONS = ["Weight Loss", "Muscle Gain", "Increase Overall Strength", "Improve Cardiovascular Fitness", "Improve Flexibility & Mobility", "Rehabilitation & Injury Prevention", "Improve Posture and Balance", "General Fitness", "Weight Maintenance"]
FITNESS_LEVEL_OPTIONS = ["Beginner (0–6 months)", "Intermediate (6 months–2 years)", "Advanced (2+ years)"]
FITNESS_MEDICAL_CONDITIONS_OPTIONS = ["None", "Hypertension (High Blood Pressure)", "Type 2 Diabetes", "Osteoarthritis", "Chronic Lower Back Pain", "Other"]
FITNESS_DAYS_OPTIONS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
FITNESS_DURATION_OPTIONS = ["15-20 minutes", "20-30 minutes", "30-45 minutes", "45-60 minutes"]
FITNESS_LOCATION_OPTIONS = ["Home", "Gym", "Outdoor", "Any"]
FITNESS_EQUIPMENT_OPTIONS = ["Bodyweight Only", "Dumbbells", "Resistance Bands", "Kettlebells", "Barbell", "Bench", "Pull-up Bar", "Yoga Mat", "Machines"]

DIETARY_PREFERENCE_OPTIONS = [
    "Vegetarian", "Vegan", "Non Vegetarian", "No Pork", "No Beef", "No Red Meat"
]
DIETARY_RESTRICTION_OPTIONS = ["Dairy-Free", "Lactose Intolerance", "Gluten-Free"]
MEDICAL_CONDITIONS_NUTRITION_OPTIONS = [
    "Alzheimer’s Disease", "Anemia", "Cardiovascular Disease", "Constipation",
    "Coronary Artery Disease", "Crohn’s Disease", "Depression", "Heart Failure",
    "High Cholesterol", "Hypertension", "Hyperthyroidism", "Hypothyroidism",
    "Liver Disease", "Metabolic Syndrome", "Non-Alcoholic Fatty Liver Disease (NAFLD)",
    "Obesity", "Polycystic Ovary Syndrome (PCOS)", "Type 1 Diabetes", "Type 2 Diabetes"
]
DIGESTIVE_ISSUES_OPTIONS = [
    "Occasional bloating", "Constipation", "Diarrhea", "Stomach pain or cramping",
    "Nausea or vomiting", "Excessive gas"
]
SYMPTOM_AGGRAVATING_FOODS_OPTIONS = [
    "Fried or greasy foods", "Spicy foods", "Caffeinated beverages",
    "Carbonated drinks", "Alcohol", "High-fiber foods", "Raw vegetables",
    "Chocolate", "Garlic and onions", "Processed or packaged foods"
]
CUISINE_OPTIONS = [
    "North American", "African", "Central American", "Chinese", "European", "Indian",
    "Indian South", "International", "Italian", "Japanese", "Korean", "Latin American",
    "Middle Eastern", "North Indian", "South American", "Spanish", "Thai", "Turkish",
    "Vietnamese"
]
FOOD_ALLERGIES_OPTIONS = [
    "Tree Nuts (e.g. Almonds, Walnuts, Cashews)", "Shellfish", "Fish", "Eggs",
    "Milk (Dairy Allergy)", "Wheat", "Soy", "Sesame", "Corn",
    "Legumes (e.g. Lentils, Chickpeas)", "Peanuts"
]

PROFILE_CONSTRAINT_FIELDS = {
    "name", "age", "gender", "height_cm", "weight_kg", "activity_level",
    "waist_circumference_cm", "cuisine", "symptom_aggravating_foods",
    "dietary_preference", "restrictions", "allergies", "medical_conditions",
    "digestive_issues", "vitals_numeric"
}

EXPERT_NAMES = ["Dr. Anjali Singh", "Dr. Amit Patel"]
SESSION_TYPES = ["Yoga", "Stress Relief", "Nutritional Advice"]
APPOINTMENT_TYPES = ["Video Call", "Audio Call"]

st.set_page_config(page_title="MG AI — Personalised Nutrition", page_icon="✦", layout="wide", initial_sidebar_state="collapsed")

st.markdown("""
<style>
:root {
    --mg-ink:#151515; --mg-muted:#6f6f6b; --mg-soft:#999993; --mg-paper:#f7f7f5;
    --mg-card:#ffffff; --mg-line:rgba(21,21,21,.09); --mg-sage:#6e8b74;
    --mg-sage-deep:#536c59; --mg-gold:#c9a86a;
    --mg-shadow:0 18px 45px rgba(21,21,21,.06);
    --mg-shadow-hover:0 26px 58px rgba(21,21,21,.10);
}
html,body,[data-testid="stAppViewContainer"]{background:var(--mg-paper);color:var(--mg-ink);}
[data-testid="stHeader"]{background:rgba(247,247,245,.82);}
[data-testid="stMainBlockContainer"]{max-width:1440px;padding-top:1.25rem;padding-bottom:4rem;}
[data-testid="stSidebar"]{background:#151515;color:#f7f7f5;}
[data-testid="stSidebar"] *{color:#f7f7f5!important;}
.mg-brand-row{display:flex;justify-content:space-between;align-items:center;margin:.25rem 0 2rem;}
.mg-brand{display:flex;align-items:center;gap:12px;letter-spacing:.18em;font-size:.80rem;font-weight:800;}
.mg-brand-mark{width:34px;height:34px;display:grid;place-items:center;border-radius:50%;background:#151515;color:#f7f7f5;box-shadow:0 8px 24px rgba(21,21,21,.15);font-family:Georgia,serif;}
.mg-brand-caption{color:#8c8c86;font-size:.68rem;letter-spacing:.12em;font-weight:700;text-transform:uppercase;}
.mg-eyebrow,.mg-section-kicker{text-transform:uppercase;letter-spacing:.17em;font-size:.70rem;font-weight:800;color:var(--mg-sage-deep);}
.mg-hero-title,.mg-section-title{font-family:Georgia,"Times New Roman",serif;letter-spacing:-.04em;}
.mg-hero-title{font-size:clamp(2.8rem,5.8vw,6.2rem);line-height:.98;margin:0;max-width:980px;}
.mg-section-title{font-size:clamp(1.9rem,3vw,3rem);margin:.25rem 0 .55rem;line-height:1.05;}
.mg-hero-subtitle,.mg-section-copy{color:var(--mg-muted);line-height:1.72;}
.mg-hero-subtitle{font-size:1.02rem;max-width:720px;margin-top:1.1rem;}
.mg-section-copy{margin-bottom:1.35rem;}
.mg-hero{background:radial-gradient(circle at 82% 18%,rgba(201,168,106,.20),transparent 28%),radial-gradient(circle at 92% 80%,rgba(110,139,116,.13),transparent 28%),linear-gradient(135deg,#fff 0%,#f5f4ef 100%);border:1px solid var(--mg-line);border-radius:34px;padding:clamp(2rem,5vw,4.8rem);box-shadow:var(--mg-shadow);overflow:hidden;position:relative;}
.mg-hero:after{content:"";position:absolute;width:260px;height:260px;right:-90px;top:-80px;border-radius:50%;border:1px solid rgba(201,168,106,.28);}
.mg-metric{background:rgba(255,255,255,.78);border:1px solid var(--mg-line);border-radius:20px;padding:1rem 1.1rem;min-height:94px;}
.mg-metric-label{color:var(--mg-soft);font-size:.70rem;text-transform:uppercase;letter-spacing:.10em;font-weight:800;}
.mg-metric-value{font-family:Georgia,"Times New Roman",serif;font-size:1.8rem;letter-spacing:-.03em;margin-top:.18rem;}
.mg-metric-unit{font-size:.80rem;color:var(--mg-muted);}
.mg-profile-card,.mg-nutri-card{background:#fff;border:1px solid var(--mg-line);border-radius:26px;padding:1.35rem;box-shadow:0 10px 34px rgba(21,21,21,.035);}
.mg-chip{display:inline-flex;align-items:center;padding:.48rem .70rem;background:#f2f2ee;border:1px solid rgba(21,21,21,.06);color:#3c3c38;border-radius:999px;font-size:.76rem;margin:.25rem .25rem 0 0;}
.mg-summary{background:#151515;color:#fff;border-radius:28px;padding:1.6rem;box-shadow:0 22px 52px rgba(21,21,21,.13);}
.mg-summary-muted{color:rgba(255,255,255,.62);font-size:.76rem;letter-spacing:.10em;text-transform:uppercase;font-weight:800;}
.mg-summary-big{font-family:Georgia,"Times New Roman",serif;font-size:clamp(2.2rem,4.5vw,4.5rem);line-height:1;margin-top:.35rem;}
.mg-summary-sub{color:rgba(255,255,255,.72);margin-top:.4rem;}
.mg-summary-stat{border-left:1px solid rgba(255,255,255,.13);padding-left:1rem;}
.mg-summary-stat .label{color:rgba(255,255,255,.52);font-size:.68rem;text-transform:uppercase;letter-spacing:.09em;}
.mg-summary-stat .value{font-size:1.25rem;margin-top:.22rem;font-weight:700;}
.mg-meal-card{background:#fff;border:1px solid var(--mg-line);border-radius:28px;overflow:hidden;box-shadow:var(--mg-shadow);transition:transform .22s ease,box-shadow .22s ease;}
.mg-meal-card:hover{transform:translateY(-4px);box-shadow:var(--mg-shadow-hover);}
.mg-meal-image{height:235px;background-size:cover;background-position:center;position:relative;}
.mg-meal-image:after{content:"";position:absolute;inset:0;background:linear-gradient(to top,rgba(0,0,0,.35),rgba(0,0,0,.02) 60%);}
.mg-meal-badge{position:absolute;left:18px;top:16px;z-index:2;background:rgba(255,255,255,.90);color:#151515;border-radius:999px;padding:.42rem .62rem;font-size:.68rem;text-transform:uppercase;letter-spacing:.10em;font-weight:800;}
.mg-meal-body{padding:1.25rem 1.25rem 1rem;}
.mg-meal-type{font-size:.68rem;text-transform:uppercase;letter-spacing:.13em;color:var(--mg-sage-deep);font-weight:800;}
.mg-meal-name{font-family:Georgia,"Times New Roman",serif;font-size:1.55rem;letter-spacing:-.02em;margin:.30rem 0 .8rem;line-height:1.15;}
.mg-meal-lines{color:#4f4f49;font-size:.83rem;line-height:1.65;min-height:55px;}
.mg-meal-cal{font-size:.95rem;font-weight:800;margin-top:.85rem;}
.mg-meal-macros{display:flex;gap:14px;flex-wrap:wrap;margin-top:.45rem;}
.mg-mini{color:#77776f;font-size:.76rem;}.mg-mini b{color:#272723;}
.mg-ai-card{background:radial-gradient(circle at 90% 15%,rgba(201,168,106,.16),transparent 25%),#f0eee8;border:1px solid rgba(201,168,106,.32);border-radius:28px;padding:1.55rem;}
.mg-ai-label{display:inline-block;background:#151515;color:#fff;border-radius:999px;padding:.35rem .58rem;font-size:.65rem;letter-spacing:.12em;font-weight:800;}
.mg-ai-copy{font-family:Georgia,"Times New Roman",serif;font-size:1.42rem;line-height:1.35;margin-top:.7rem;}
.mg-ai-footnote{color:#77736b;font-size:.75rem;line-height:1.5;margin-top:.65rem;}
.mg-progress{height:9px;background:#efefe9;border-radius:999px;overflow:hidden;}
.mg-progress>div{height:100%;border-radius:999px;}
.mg-progress-protein>div{background:#6e8b74;}.mg-progress-carbs>div{background:#c9a86a;}.mg-progress-fat>div{background:#92978d;}
.mg-nutrient-row{margin:.75rem 0 1rem;}.mg-nutrient-header{display:flex;justify-content:space-between;margin-bottom:.35rem;font-size:.8rem;color:#54544e;}
.mg-detail-grid{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:10px;margin-top:.8rem;}
.mg-detail-cell{background:#f5f5f1;border-radius:16px;padding:.8rem;}
.mg-detail-cell .k{color:#888880;font-size:.66rem;text-transform:uppercase;letter-spacing:.08em;}
.mg-detail-cell .v{font-weight:800;margin-top:.24rem;}
.mg-form-hero{background:radial-gradient(circle at 85% 15%,rgba(201,168,106,.18),transparent 26%),#151515;color:#fff;padding:clamp(1.7rem,4vw,3.4rem);border-radius:32px;margin-bottom:1.5rem;}
.mg-form-hero h1{font-family:Georgia,"Times New Roman",serif;font-size:clamp(2.1rem,4.5vw,4.4rem);margin:.2rem 0 .6rem;letter-spacing:-.04em;}
.mg-form-hero p{color:rgba(255,255,255,.68);max-width:720px;line-height:1.7;}
div[data-testid="stForm"]{border:1px solid var(--mg-line);border-radius:26px;padding:1.4rem;background:#fff;box-shadow:var(--mg-shadow);}
div[data-testid="stExpander"]{border:1px solid var(--mg-line);border-radius:22px;background:rgba(255,255,255,.78);}
.stButton>button,.stDownloadButton>button{border-radius:999px!important;min-height:44px!important;font-weight:800!important;letter-spacing:.01em!important;border:1px solid rgba(21,21,21,.10)!important;box-shadow:none!important;transition:all .18s ease!important;}
.stButton>button:hover,.stDownloadButton>button:hover{transform:translateY(-1px)!important;box-shadow:0 10px 24px rgba(21,21,21,.08)!important;}
.stButton>button[kind="primary"]{background:#151515!important;color:#fff!important;border-color:#151515!important;}
.stButton>button[kind="secondary"]{background:#fff!important;color:#151515!important;}
.stTabs [data-baseweb="tab-list"]{gap:24px;border-bottom:1px solid var(--mg-line);}
.stTabs [data-baseweb="tab"]{padding:.85rem 0;font-weight:800;}
[data-testid="stMetric"]{background:#fff;border:1px solid var(--mg-line);border-radius:18px;padding:.9rem;}
input,textarea{border-radius:14px!important;}
@media(max-width:900px){.mg-detail-grid{grid-template-columns:repeat(2,minmax(0,1fr));}.mg-meal-image{height:210px;}}
@media(max-width:640px){[data-testid="stMainBlockContainer"]{padding-left:.75rem;padding-right:.75rem;}.mg-hero{border-radius:24px;padding:1.55rem;}.mg-hero-title{font-size:3rem;}.mg-meal-image{height:190px;}.mg-detail-grid{grid-template-columns:1fr 1fr;}}
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="mg-brand-row">
    <div class="mg-brand"><span class="mg-brand-mark">M</span><span>MG AI</span></div>
    <div class="mg-brand-caption">Personal Nutrition Intelligence</div>
</div>
""", unsafe_allow_html=True)



if "messages" not in st.session_state:
    st.session_state.messages = []
if "profile_complete" not in st.session_state:
    st.session_state.profile_complete = False
if "current_constraints" not in st.session_state:
    st.session_state.current_constraints = {}
if "last_agent_context" not in st.session_state:
    st.session_state.last_agent_context = {}
if "pending_images" not in st.session_state:
    st.session_state.pending_images = []
if "auth_token" not in st.session_state:
    # Use the token from your provided code
    st.session_state.auth_token = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiIrOTE5NzE3MjY1NjExIiwianRpIjoiNDAwMjI5YjAtNzg0OS00NjYwLTg1Y2QtMjkzM2RjMmU3OWRhIiwiVXNlcklkIjoiZWI0NmVkYTItODUzMy00MDNkLWE4ZjktMWQ3NjlkMjk4Nzc1IiwiSWQiOiJlYjQ2ZWRhMi04NTMzLTQwM2QtYThmOS0xZDc2OWQyOTg3NzUiLCJFbWFpbCI6ImF5dXNoLnJhanB1dEBub3VyaXEuYWkiLCJUZW5hbnRJZCI6IjQiLCJEYXRhYmFzZU5hbWUiOiJGcmlza2FBaUNDTV9IRldMIiwiRGF0YUJhc2VOYW1lIjoiRnJpc2thQWlDQ01fSEZXTCIsIlJvbGVOYW1lIjoiUGF0aWVudCIsIldlbGxuZXNzU3RhdHVzIjoiRW5yb2xsZWQiLCJXZWxsbmVzc1N0YXR1c0NvZGUiOiJFTlJPTExFRCIsIlBhdGllbnRGaXJzdE5hbWUiOiJBeXVzaCIsIklzSW50ZXJuYWxVc2VyIjoiRmFsc2UiLCJQcm9qZWN0VHlwZSI6IkNDTVRlc3RpbmciLCJleHAiOjE3Nzk0NzMyMjcsImlzcyI6IkhGV0wiLCJhdWQiOiJIRldMIn0.9BM8MP2EUv1dhs7kaPH4EkdzePlRufWd4xUSue65d_c"
if "blood_report_temp_path" not in st.session_state:
    st.session_state.blood_report_temp_path = None
if "fitness_profile_data" not in st.session_state:
    st.session_state.fitness_profile_data = {}
    
# --- CRITICAL FIX 1: DYNAMIC MIC KEY INITIALIZATION ---
if "mic_key" not in st.session_state:
    st.session_state.mic_key = "initial_mic_key"


def call_backend_api(payload: dict, token: str, files: list = None):
    try:
        headers = {
            "Authorization": f"Bearer {token}"
            # Do NOT manually set Content-Type here; 'requests' does it for multipart
        }
        
        # We need to send 'query' and 'last_agent_context' as Form fields
        form_data = {
            "query": payload.get("query", ""),
            "last_agent_context": json.dumps(payload.get("last_agent_context", {}))
        }

        # Handle Image files for the 'image' parameter in your backend
        file_payload = None
        if files and len(files) > 0:
            # Taking the first image as your backend 'image' parameter expects one file
            file_payload = {"image": (files[0].name, files[0].getvalue(), files[0].type)}

        # Use 'data=' for form-data, 'files=' for the image
        response = requests.post(
            BACKEND_API_URL, 
            data=form_data, 
            files=file_payload, 
            headers=headers, 
            timeout=1200
        )
        
        response.raise_for_status()
        return response.json()
    except requests.exceptions.RequestException as e:
        st.error(f"Backend Error: {e}")
        return None

def call_meal_plan_api(profile: dict, query: str, token: str, force_tool_type: str = "meal_plan_generator", start_date: Optional[str] = None, end_date: Optional[str] = None):
    """Sends the frontend-collected profile directly to the backend so it can be
    used as the meal-plan constraints, bypassing the DB profile lookup/auth."""
    try:
        headers = {"Authorization": f"Bearer {token}"}
        body = {
            "profile": profile,
            "query": query,
            "last_agent_context": st.session_state.last_agent_context,
            "force_tool_type": force_tool_type,
            "start_date": start_date,
            "end_date": end_date
        }
        response = requests.post(
            MEAL_PLAN_FROM_PROFILE_API_URL,
            json=body,
            headers=headers,
            timeout=1200
        )
        response.raise_for_status()
        return response.json()
    except requests.exceptions.RequestException as e:
        st.error(f"Backend Error: {e}")
        return None

def convert_images_to_base64_urls(files: list):
    try:
        base64_urls = []
        for uploaded_file in files:
            image_bytes = uploaded_file.getvalue()
            base64_encoded = base64.b64encode(image_bytes).decode('utf-8')
            
            mime_type = uploaded_file.type if uploaded_file.type else "image/jpeg"
            data_url = f"data:{mime_type};base64,{base64_encoded}"
            base64_urls.append(data_url)
        
        return base64_urls
    except Exception as e:
        st.error(f"Failed to process images: {e}")
        return None

# def send_message_to_backend(user_input: str, force_tool: Optional[str] = None, images: list = None):
#     with st.spinner("MG AI is thinking..."):
#         api_payload = {
#             "query": user_input,
#             "chat_history": st.session_state.messages, 
#             "current_constraints": st.session_state.current_constraints,
#             "last_agent_context": st.session_state.last_agent_context,
#             "force_tool_type": force_tool 
#         }
        
#         response_data = call_backend_api(api_payload, st.session_state.auth_token, files = images) 

#         if response_data:
#             st.session_state.messages.append({
#                 "role": "assistant", 
#                 "content": response_data.get("answer"), 
#                 "audio": response_data.get("audio"), 
#                 "context": response_data.get("context")
#             })
                
#             st.session_state.current_constraints = response_data.get("updated_constraints", {})
#             st.session_state.last_agent_context = response_data.get("updated_last_agent_context", {})
    
#             st.rerun()
            
def send_message_to_backend(user_input: str, force_tool: Optional[str] = None, images: list = None):
    with st.spinner("MG AI is thinking..."):
        api_payload = {
            "query": user_input,
            "constraints": st.session_state.current_constraints,
            "last_agent_context": st.session_state.last_agent_context,
            "force_tool_type": force_tool 
        }
        
        full_response = call_backend_api(api_payload, st.session_state.auth_token, files=images) 

        if full_response and "data" in full_response:
            # Extract 'ai_data' from the 'data' key returned by your FastAPI chat endpoint
            ai_data = full_response["data"]
            
            st.session_state.messages.append({
                "role": "assistant", 
                "content": ai_data.get("answer"), # Backend uses 'answer'
                "audio": ai_data.get("audio_base64"), # Your backend uses 'audio_base64'
                "context": ai_data.get("context")
            })
                
            # Update history and context
            st.session_state.last_agent_context = ai_data.get("updated_last_agent_context", {})
    
            st.rerun()

def save_uploaded_file_temp(uploaded_file):
    try:
        temp_dir = "temp_uploads"
        os.makedirs(temp_dir, exist_ok=True)
        
        temp_file_path = os.path.join(temp_dir, uploaded_file.name)
        counter = 1
        original_temp_file_path = temp_file_path
        while os.path.exists(temp_file_path):
            name, ext = os.path.splitext(original_temp_file_path)
            temp_file_path = f"{name}_{counter}{ext}"
            counter += 1

        with open(temp_file_path, "wb") as f:
            f.write(uploaded_file.getvalue())
        return temp_file_path
    except Exception as e:
        st.error(f"Error saving file: {e}")
        return None


# =========================
# MG AI PREMIUM UI HELPERS
# =========================

MEAL_TYPES = ["Breakfast", "Morning Snack", "Lunch", "Evening Snack", "Dinner"]

MG_AI_MEAL_IMAGES = {
    "Breakfast": "https://images.unsplash.com/photo-1490645935967-10de6ba17061?auto=format&fit=crop&w=1400&q=88",
    "Morning Snack": "https://images.unsplash.com/photo-1546554137-f86b9593a222?auto=format&fit=crop&w=1400&q=88",
    "Lunch": "https://images.unsplash.com/photo-1547592180-85f173990554?auto=format&fit=crop&w=1400&q=88",
    "Evening Snack": "https://images.unsplash.com/photo-1512621776951-a57141f2eefd?auto=format&fit=crop&w=1400&q=88",
    "Dinner": "https://images.unsplash.com/photo-1515003197210-e0cd71810b5f?auto=format&fit=crop&w=1400&q=88",
}

def _clean_ui_text(value):
    value = "" if value is None else str(value)
    value = re.sub(r"<[^>]+>", "", value)
    value = re.sub(r"[*_`#]+", "", value)
    return re.sub(r"\s+", " ", value).strip()

def _format_number(value, suffix=""):
    if value is None:
        return "—"
    if abs(float(value) - round(float(value))) < 0.05:
        return f"{int(round(float(value)))}{suffix}"
    return f"{float(value):.1f}{suffix}"

def _parse_date_from_title(title):
    title = _clean_ui_text(title)
    for fmt in ("%A, %B %d, %Y", "%B %d, %Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(title, fmt).date()
        except ValueError:
            continue
    return None

def _macro_dict_from_text(text):
    flat = re.sub(r"\s+", " ", text)
    out = {}
    patterns = {
        "protein": r"Protein\s+(-?\d+(?:\.\d+)?)\s*g\s*\((-?\d+(?:\.\d+)?)\s*%\)",
        "carbs": r"(?:Carbs|Carbohydrates)\s+(-?\d+(?:\.\d+)?)\s*g\s*\((-?\d+(?:\.\d+)?)\s*%\)",
        "fat": r"Fat\s+(-?\d+(?:\.\d+)?)\s*g\s*\((-?\d+(?:\.\d+)?)\s*%\)",
    }
    for key, pattern in patterns.items():
        m = re.search(pattern, flat, flags=re.I)
        if m:
            out[key] = {"grams": float(m.group(1)), "percent": float(m.group(2))}
    return out

def _extract_item_blocks(meal_text):
    lines = [re.sub(r"\s+", " ", x).strip() for x in meal_text.splitlines()]
    lines = [x for x in lines if x]
    starts = []
    for idx, line in enumerate(lines):
        line_clean = re.sub(r"^[-*]\s*", "", line)
        if re.search(r"\|\s*Household Measure\s*:", line_clean, flags=re.I):
            starts.append((idx, line_clean))
    blocks = []
    for pos, (start_idx, _) in enumerate(starts):
        end_idx = starts[pos + 1][0] if pos + 1 < len(starts) else len(lines)
        blocks.append(" ".join(lines[start_idx:end_idx]))
    return blocks

def _parse_meal_item(block):
    block = _clean_ui_text(block)
    name_match = re.match(r"(.+?)\s*\|\s*Household Measure\s*:", block, flags=re.I)
    if not name_match:
        return None

    item = {"name": _clean_ui_text(name_match.group(1))}
    fields = {
        "household_measure": r"Household Measure\s*:\s*(.*?)(?=\s*\|\s*Portion Weight\s*:)",
        "portion_weight": r"Portion Weight\s*:\s*(.*?)(?=\s*\|\s*Protein\s*:)",
        "protein": r"Protein\s*:\s*(-?\d+(?:\.\d+)?)\s*g",
        "carbs": r"Carbs\s*:\s*(-?\d+(?:\.\d+)?)\s*g",
        "fat": r"Fat\s*:\s*(-?\d+(?:\.\d+)?)\s*g",
        "fiber": r"Fiber\s*:\s*(-?\d+(?:\.\d+)?)\s*g",
        "sodium": r"Sodium\s*:\s*(-?\d+(?:\.\d+)?)\s*mg",
        "iodine": r"Iodine\s*:\s*(-?\d+(?:\.\d+)?)\s*mcg",
        "sugar": r"Sugar\s*:\s*(-?\d+(?:\.\d+)?)\s*g",
        "cholesterol": r"Cholesterol\s*:\s*(-?\d+(?:\.\d+)?)\s*mg",
        "calories": r"Calories\s*:\s*(-?\d+(?:\.\d+)?)\s*kcal",
    }
    for key, pattern in fields.items():
        m = re.search(pattern, block, flags=re.I)
        if m:
            raw = m.group(1).strip()
            item[key] = float(raw) if key not in {"household_measure", "portion_weight"} else raw

    ingredients = re.search(r"Ingredients\s*:\s*(.*?)(?=\s*Recipe\s*:)", block, flags=re.I)
    recipe = re.search(r"Recipe\s*:\s*(.*)$", block, flags=re.I)
    item["ingredients"] = _clean_ui_text(ingredients.group(1)) if ingredients else ""
    item["recipe"] = _clean_ui_text(recipe.group(1)) if recipe else ""

    if item.get("ingredients", "").upper() == "N/A":
        item["ingredients"] = ""
    if item.get("recipe", "").upper() == "N/A":
        item["recipe"] = ""
    return item

def parse_weekly_meal_plan(raw_plan):
    result = []
    raw_plan = raw_plan if isinstance(raw_plan, str) else str(raw_plan or "")
    if not raw_plan.strip():
        return result

    normalized = raw_plan.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    day_matches = list(re.finditer(r"(?im)^\s*(?:#+\s*)?Day\s+(\d+)\s*[:\-–]\s*(.+?)\s*$", normalized))

    if not day_matches:
        day_matches = [None]

    first_day_prefix = normalized[:day_matches[0].start()] if day_matches and day_matches[0] is not None else ""
    for idx, match in enumerate(day_matches):
        if match is None:
            day_num = 1
            title = datetime.now().strftime("%A, %B %d, %Y")
            section = normalized
        else:
            day_num = int(match.group(1))
            title = _clean_ui_text(match.group(2))
            start = match.end()
            end = day_matches[idx + 1].start() if idx + 1 < len(day_matches) else len(normalized)
            section = normalized[start:end]
            # Some report renderers place the first page's meal headings before the
            # "Day 1" heading. Retain that real meal content instead of dropping it.
            if idx == 0 and re.search(
                r"(?im)^\s*(?:#+\s*)?(Breakfast|Morning Snack|Lunch|Evening Snack|Dinner)\s*$",
                first_day_prefix,
            ):
                section = first_day_prefix + "\n" + section

        section = re.sub(r"(?im)^\s*Here is your meal plan.*?$", "", section)
        meal_matches = list(re.finditer(
            r"(?im)^\s*(?:#+\s*)?(Breakfast|Morning Snack|Lunch|Evening Snack|Dinner)\s*$",
            section
        ))
        if not meal_matches:
            meal_matches = list(re.finditer(
                r"(?i)(?:^|\n)\s*(Breakfast|Morning Snack|Lunch|Evening Snack|Dinner)\s*(?=\n|$)",
                section
            ))

        day = {
            "day": day_num,
            "title": title,
            "date": _parse_date_from_title(title),
            "meals": {},
            "daily_calories": None,
            "overall_macros": {},
        }

        for meal_idx, meal_match in enumerate(meal_matches):
            meal_type = meal_match.group(1)
            meal_start = meal_match.end()
            meal_end = meal_matches[meal_idx + 1].start() if meal_idx + 1 < len(meal_matches) else len(section)
            meal_text = section[meal_start:meal_end]
            flat_meal = re.sub(r"\s+", " ", meal_text).strip()

            meal = {"type": meal_type, "items": [], "total_calories": None, "macros": {}}

            total_match = re.search(
                r"Total calories for\s+" + re.escape(meal_type) + r"\s*:\s*(-?\d+(?:\.\d+)?)\s*kcal",
                flat_meal,
                flags=re.I
            )
            if total_match:
                meal["total_calories"] = float(total_match.group(1))

            macro_match = re.search(r"Macronutrient Ratio\s*:\s*(.+)", flat_meal, flags=re.I)
            if macro_match:
                meal["macros"] = _macro_dict_from_text(macro_match.group(0))

            for block in _extract_item_blocks(meal_text):
                parsed_item = _parse_meal_item(block)
                if parsed_item:
                    meal["items"].append(parsed_item)

            if meal["total_calories"] is None and meal["items"]:
                vals = [i.get("calories") for i in meal["items"] if i.get("calories") is not None]
                if vals:
                    meal["total_calories"] = round(sum(vals), 1)

            if not meal["macros"] and meal["items"]:
                totals = {}
                for macro in ("protein", "carbs", "fat"):
                    vals = [i.get(macro) for i in meal["items"] if i.get(macro) is not None]
                    if vals:
                        totals[macro] = {"grams": round(sum(vals), 1), "percent": 0.0}
                gram_total = sum(v["grams"] for v in totals.values())
                for data in totals.values():
                    data["percent"] = round((data["grams"] / gram_total) * 100, 1) if gram_total else 0.0
                meal["macros"] = totals

            if meal["items"] or meal["total_calories"] is not None:
                day["meals"][meal_type] = meal

        flat_section = re.sub(r"\s+", " ", section)
        daily_match = re.search(r"Daily Total Calories\s*:\s*(-?\d+(?:\.\d+)?)\s*kcal", flat_section, flags=re.I)
        if daily_match:
            day["daily_calories"] = float(daily_match.group(1))
        else:
            totals = [m["total_calories"] for m in day["meals"].values() if m.get("total_calories") is not None]
            if totals:
                day["daily_calories"] = round(sum(totals), 1)

        overall_match = re.search(r"Overall Macronutrient Distribution\s*:\s*(.+)", flat_section, flags=re.I)
        if overall_match:
            day["overall_macros"] = _macro_dict_from_text(overall_match.group(0))

        if day["meals"]:
            result.append(day)

    return sorted(result, key=lambda d: d.get("day", 999))

def _profile_value(profile, key, default=None):
    value = profile.get(key, default)
    if value in ("", [], None):
        return default
    if isinstance(value, list):
        return ", ".join(str(v) for v in value) if value else default
    return str(value)

def _profile_coverage(profile):
    fields = [
        "name", "age", "gender", "height_cm", "weight_kg",
        "activity_level", "dietary_preference", "restrictions",
        "allergies", "medical_conditions", "cuisine", "digestive_issues"
    ]
    return round(sum(1 for f in fields if profile.get(f) not in (None, "", [])) / len(fields) * 100)

def _meal_display_names(meal):
    names = [_clean_ui_text(i.get("name")) for i in meal.get("items", []) if i.get("name")]
    if not names:
        return "Personalised meal"
    return " · ".join(names[:2]) + (f" · +{len(names)-2} more" if len(names) > 2 else "")

def _render_profile_snapshot(profile):
    name = _profile_value(profile, "name", "there")
    dietary = _profile_value(profile, "dietary_preference", "Not set")
    cuisine = _profile_value(profile, "cuisine", "Flexible")
    restrictions = _profile_value(profile, "restrictions", "None")
    allergies = _profile_value(profile, "allergies", "None")
    medical = _profile_value(profile, "medical_conditions", "None")
    digestive = _profile_value(profile, "digestive_issues", "None")
    activity = _profile_value(profile, "activity_level", "Not set")

    st.markdown("""
    <div class="mg-section-kicker">Your profile</div>
    <div class="mg-section-title">Built around you.</div>
    <div class="mg-section-copy">Your preferences remain editable; the main experience only surfaces the signals that matter at a glance.</div>
    """, unsafe_allow_html=True)

    st.markdown(f"""
    <div class="mg-profile-card">
        <div style="font-family:Georgia,serif;font-size:1.55rem;">{_clean_ui_text(name)}</div>
        <div style="margin-top:0.75rem;">
            <span class="mg-chip">{_format_number(profile.get("age"))} years</span>
            <span class="mg-chip">{_format_number(profile.get("weight_kg"), " kg")}</span>
            <span class="mg-chip">{_clean_ui_text(activity)}</span>
        </div>
        <div style="margin-top:0.9rem;">
            <span class="mg-chip">Diet · {_clean_ui_text(dietary)}</span>
            <span class="mg-chip">Cuisine · {_clean_ui_text(cuisine)}</span>
            <span class="mg-chip">Restrictions · {_clean_ui_text(restrictions)}</span>
        </div>
        <div style="margin-top:0.9rem;color:#6f6f6b;font-size:0.76rem;line-height:1.65;">
            <b>Allergies:</b> {_clean_ui_text(allergies)}
            &nbsp;&nbsp;·&nbsp;&nbsp;
            <b>Medical:</b> {_clean_ui_text(medical)}
            &nbsp;&nbsp;·&nbsp;&nbsp;
            <b>Digestive:</b> {_clean_ui_text(digestive)}
        </div>
    </div>
    """, unsafe_allow_html=True)

def _render_daily_summary(day):
    macros = day.get("overall_macros", {})
    nice_date = day["date"].strftime("%A, %B %d") if day.get("date") else day.get("title", "Selected day")
    st.markdown(f"""
    <div class="mg-summary">
        <div class="mg-summary-muted">Day {day.get('day', '')} · {nice_date}</div>
        <div class="mg-summary-big">{_format_number(day.get('daily_calories'), ' kcal')}</div>
        <div class="mg-summary-sub">Planned for your selected day</div>
        <div style="display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:18px;margin-top:1.35rem;">
            <div class="mg-summary-stat"><div class="label">Protein</div><div class="value">{_format_number(macros.get('protein', {}).get('grams'), ' g')}</div></div>
            <div class="mg-summary-stat"><div class="label">Carbs</div><div class="value">{_format_number(macros.get('carbs', {}).get('grams'), ' g')}</div></div>
            <div class="mg-summary-stat"><div class="label">Fat</div><div class="value">{_format_number(macros.get('fat', {}).get('grams'), ' g')}</div></div>
        </div>
    </div>
    """, unsafe_allow_html=True)

def _meal_image_url(meal_type):
    return MG_AI_MEAL_IMAGES.get(meal_type, MG_AI_MEAL_IMAGES["Lunch"])

def _render_meal_card(meal, idx):
    meal_type = meal.get("type", "Meal")
    items = meal.get("items", [])
    names = [_clean_ui_text(i.get("name")) for i in items if i.get("name")]
    main_name = " · ".join(names[:2]) if names else "Personalised meal"
    if len(names) > 2:
        main_name += f" · +{len(names)-2} more"

    macros = meal.get("macros", {})
    image_url = _meal_image_url(meal_type)
    html = f"""
    <div class="mg-meal-card">
        <div class="mg-meal-image" style="background-image:url('{image_url}');">
            <span class="mg-meal-badge">{_clean_ui_text(meal_type)}</span>
        </div>
        <div class="mg-meal-body">
            <div class="mg-meal-type">{len(items)} component{"s" if len(items) != 1 else ""}</div>
            <div class="mg-meal-name">{main_name}</div>
            <div class="mg-meal-lines">Personalised portioning, ingredients and recipe details are available below without crowding the main card.</div>
            <div class="mg-meal-cal">{_format_number(meal.get("total_calories"), " kcal")}</div>
            <div class="mg-meal-macros">
                <span class="mg-mini"><b>{_format_number(macros.get('protein', {}).get('grams'), 'g')}</b> protein</span>
                <span class="mg-mini"><b>{_format_number(macros.get('carbs', {}).get('grams'), 'g')}</b> carbs</span>
                <span class="mg-mini"><b>{_format_number(macros.get('fat', {}).get('grams'), 'g')}</b> fat</span>
            </div>
        </div>
    </div>
    """
    st.markdown(html, unsafe_allow_html=True)
    return st.button("View recipe →", key=f"mg_recipe_{idx}_{meal_type}", use_container_width=True)

def _render_recipe_details(meal, day_label):
    st.markdown(f"""
    <div class="mg-ai-card">
        <span class="mg-ai-label">{_clean_ui_text(meal.get('type', 'Meal'))}</span>
        <div class="mg-ai-copy">{_meal_display_names(meal)}</div>
        <div class="mg-ai-footnote">{day_label}</div>
    </div>
    """, unsafe_allow_html=True)

    for item in meal.get("items", []):
        name = _clean_ui_text(item.get("name", "Meal component"))
        st.markdown(f"### {name}")
        st.caption(
            f"{_format_number(item.get('calories'), ' kcal')} · "
            f"{_format_number(item.get('protein'), 'g')} protein · "
            f"{_format_number(item.get('carbs'), 'g')} carbs · "
            f"{_format_number(item.get('fat'), 'g')} fat"
        )
        if item.get("ingredients"):
            with st.expander("Ingredients", expanded=False):
                st.write(item["ingredients"])
        if item.get("recipe"):
            with st.expander("How to prepare", expanded=True):
                st.write(item["recipe"])

        rows = [
            ("Fiber", item.get("fiber"), "g"),
            ("Sodium", item.get("sodium"), "mg"),
            ("Iodine", item.get("iodine"), "mcg"),
            ("Sugar", item.get("sugar"), "g"),
            ("Cholesterol", item.get("cholesterol"), "mg"),
            ("Portion", item.get("portion_weight"), ""),
        ]
        rows = [(k, v, u) for k, v, u in rows if v not in (None, "")]
        if rows:
            cells = "".join(
                f'<div class="mg-detail-cell"><div class="k">{_clean_ui_text(k)}</div><div class="v">{_clean_ui_text(v)}{u}</div></div>'
                for k, v, u in rows
            )
            st.markdown(f'<div class="mg-detail-grid">{cells}</div>', unsafe_allow_html=True)
        st.divider()

def _render_nutrition_intelligence(day):
    macros = day.get("overall_macros", {})
    st.markdown("""
    <div class="mg-section-kicker">Nutrition intelligence</div>
    <div class="mg-section-title">See the shape of your day.</div>
    <div class="mg-section-copy">Detailed nutrient information remains available, while the first view focuses on the signals people can understand quickly.</div>
    """, unsafe_allow_html=True)

    left, right = st.columns([1.25, 0.95], gap="large")
    with left:
        bars = [
            ("Protein", macros.get("protein", {}), "mg-progress-protein"),
            ("Carbohydrates", macros.get("carbs", {}), "mg-progress-carbs"),
            ("Fat", macros.get("fat", {}), "mg-progress-fat"),
        ]
        for label, data, cls in bars:
            pct = float(data.get("percent", 0) or 0)
            st.markdown(f"""
            <div class="mg-nutrient-row">
                <div class="mg-nutrient-header"><span>{label}</span><span><b>{_format_number(data.get('grams'), ' g')}</b> · {pct:.1f}%</span></div>
                <div class="mg-progress {cls}"><div style="width:{max(0,min(pct,100)):.1f}%;"></div></div>
            </div>
            """, unsafe_allow_html=True)

        with st.expander("View detailed nutrition ↓", expanded=False):
            advanced = {
                "Fiber": ("g", 0.0),
                "Sodium": ("mg", 0.0),
                "Iodine": ("mcg", 0.0),
                "Sugar": ("g", 0.0),
                "Cholesterol": ("mg", 0.0),
            }
            seen = {k: False for k in advanced}
            source_keys = {"Fiber":"fiber", "Sodium":"sodium", "Iodine":"iodine", "Sugar":"sugar", "Cholesterol":"cholesterol"}
            for meal in day.get("meals", {}).values():
                for item in meal.get("items", []):
                    for label, key in source_keys.items():
                        if item.get(key) is not None:
                            advanced[label] = (advanced[label][0], advanced[label][1] + float(item[key]))
                            seen[label] = True
            rows = [(label, total, unit) for label, (unit, total) in advanced.items() if seen[label]]
            if rows:
                cells = "".join(
                    f'<div class="mg-detail-cell"><div class="k">{label}</div><div class="v">{total:.1f}{unit}</div></div>'
                    for label, total, unit in rows
                )
                st.markdown(f'<div class="mg-detail-grid">{cells}</div>', unsafe_allow_html=True)
            else:
                st.caption("Advanced nutrient details were not provided in this meal-plan response.")

    with right:
        coverage = _profile_coverage(st.session_state.get("current_constraints", {}))
        st.markdown(f"""
        <div class="mg-nutri-card">
            <div class="mg-section-kicker">MG AI profile coverage</div>
            <div style="font-family:Georgia,serif;font-size:3rem;line-height:1;margin-top:.35rem;">{coverage}%</div>
            <div style="color:#6f6f6b;font-size:.78rem;margin-top:.5rem;line-height:1.55;">Based on completed profile inputs available to MG AI. This is a product metric, not a clinical score.</div>
        </div>
        """, unsafe_allow_html=True)

def _build_ai_insight(day, profile):
    meal_count = sum(len(meal.get("items", [])) for meal in day.get("meals", {}).values())
    meal_types = len(day.get("meals", {}))
    dietary = _profile_value(profile, "dietary_preference", "your dietary preference")
    cuisine = _profile_value(profile, "cuisine", "your selected cuisine")
    calories = _format_number(day.get("daily_calories"), " kcal")
    return (
        f"Your selected day includes {meal_types} meal moments and {meal_count} meal components, "
        f"with a planned intake of {calories}. The plan is presented around "
        f"{_clean_ui_text(dietary)} preferences and {_clean_ui_text(cuisine)} cuisine where provided."
    )

def _render_ai_insight(day, profile):
    st.markdown("""
    <div class="mg-section-kicker">MG AI insight</div>
    <div class="mg-section-title">Personalised context, without the clutter.</div>
    """, unsafe_allow_html=True)
    st.markdown(f"""
    <div class="mg-ai-card">
        <span class="mg-ai-label">MG AI</span>
        <div class="mg-ai-copy">{_clean_ui_text(_build_ai_insight(day, profile))}</div>
        <div class="mg-ai-footnote">Generated from the profile and meal-plan data currently available in this session. It is not a medical diagnosis or treatment claim.</div>
    </div>
    """, unsafe_allow_html=True)

def _render_edit_profile_form():
    constraints = st.session_state.current_constraints
    existing_diet = constraints.get("dietary_preference", [])
    if isinstance(existing_diet, str):
        existing_diet = [existing_diet]

    activity_options = ["Sedentary", "Lightly active", "Moderately active", "Very active", "Extra active"]
    current_activity = str(constraints.get("activity_level", "moderately active")).strip().lower()
    activity_index = next((i for i, val in enumerate(activity_options) if val.lower() == current_activity), 2)

    with st.form("edit_meal_plan_profile_form_premium"):
        edit_col1, edit_col2 = st.columns(2)
        with edit_col1:
            edit_name = st.text_input("First Name", value=constraints.get("name", ""), key="edit_plan_name_premium")
            edit_age = st.number_input("Age", min_value=18, max_value=100, value=int(constraints.get("age", 30)), key="edit_plan_age_premium")
            edit_weight = st.number_input("Weight (kg)", min_value=30.0, max_value=300.0, value=float(constraints.get("weight_kg", 75.0)), key="edit_plan_weight_premium")
            gender_options = ["Male", "Female"]
            edit_gender = st.selectbox("Gender", gender_options, index=0 if constraints.get("gender", "Male") == "Male" else 1, key="edit_plan_gender_premium")
            edit_activity = st.selectbox("Activity Level", activity_options, index=activity_index, key="edit_plan_activity_premium")
        with edit_col2:
            edit_diet = st.multiselect("Dietary Preference", DIETARY_PREFERENCE_OPTIONS, default=[x for x in existing_diet if x in DIETARY_PREFERENCE_OPTIONS], key="edit_plan_diet_premium")
            edit_restrictions = st.multiselect("Dietary Restrictions", DIETARY_RESTRICTION_OPTIONS, default=[x for x in constraints.get("restrictions", []) if x in DIETARY_RESTRICTION_OPTIONS], key="edit_plan_restrictions_premium")
            edit_conditions = st.multiselect("Medical Conditions", MEDICAL_CONDITIONS_NUTRITION_OPTIONS, default=[x for x in constraints.get("medical_conditions", []) if x in MEDICAL_CONDITIONS_NUTRITION_OPTIONS], key="edit_plan_conditions_premium")
            edit_allergies = st.multiselect("Food Allergies", FOOD_ALLERGIES_OPTIONS, default=[x for x in constraints.get("allergies", []) if x in FOOD_ALLERGIES_OPTIONS], key="edit_plan_allergies_premium")
            edit_cuisine = st.multiselect("Preferred Cuisine", CUISINE_OPTIONS, default=[x for x in constraints.get("cuisine", []) if x in CUISINE_OPTIONS], key="edit_plan_cuisine_premium")
            edit_digestive = st.multiselect("Digestive Issues", DIGESTIVE_ISSUES_OPTIONS, default=[x for x in constraints.get("digestive_issues", []) if x in DIGESTIVE_ISSUES_OPTIONS], key="edit_plan_digestive_premium")
            edit_symptoms = st.multiselect("Symptom-Aggravating Foods", SYMPTOM_AGGRAVATING_FOODS_OPTIONS, default=[x for x in constraints.get("symptom_aggravating_foods", []) if x in SYMPTOM_AGGRAVATING_FOODS_OPTIONS], key="edit_plan_symptoms_premium")
        edit_profile_submitted = st.form_submit_button("Save profile & regenerate ✦", type="primary", use_container_width=True)

    if edit_profile_submitted:
        updated_profile = {
            **constraints,
            "name": edit_name,
            "age": edit_age,
            "weight_kg": edit_weight,
            "gender": edit_gender,
            "activity_level": edit_activity.lower(),
            "dietary_preference": edit_diet,
            "restrictions": edit_restrictions,
            "medical_conditions": edit_conditions,
            "allergies": edit_allergies,
            "cuisine": edit_cuisine,
            "digestive_issues": edit_digestive,
            "symptom_aggravating_foods": edit_symptoms,
        }
        edit_start_date = datetime.now().date()
        edit_end_date = edit_start_date + timedelta(days=6)
        with st.spinner("Rebuilding your week around the updated profile..."):
            edit_response = call_meal_plan_api(
                profile=updated_profile,
                query="Generate a complete 7-day meal plan based on my updated profile.",
                token=st.session_state.auth_token,
                force_tool_type="weekly_meal_plan_generator",
                start_date=edit_start_date.isoformat(),
                end_date=edit_end_date.isoformat()
            )
        if edit_response and isinstance(edit_response.get("data"), dict):
            edit_data = edit_response["data"]
            st.session_state.current_constraints = edit_data.get("updated_constraints", updated_profile)
            st.session_state.weekly_meal_plan = edit_data.get("answer", "")
            st.session_state.last_agent_context = edit_data.get("updated_last_agent_context", st.session_state.last_agent_context)
            st.session_state.mg_ai_open_recipe = None
            st.session_state.mg_ai_selected_day = None
            st.rerun()
        else:
            st.error("The updated seven-day meal plan could not be generated.")

def render_premium_meal_plan_dashboard():
    profile = st.session_state.get("current_constraints", {})
    weekly_meal_plan = st.session_state.get("weekly_meal_plan", "")
    days = parse_weekly_meal_plan(weekly_meal_plan)

    if not days:
        st.markdown("""
        <section class="mg-form-hero">
            <div class="mg-eyebrow" style="color:#c9a86a;">MG AI · Your nutrition workspace</div>
            <h1>Your personalised week is waiting.</h1>
            <p>Generate a seven-day plan and MG AI will turn the detailed response into an editorial, easy-to-explore experience.</p>
        </section>
        """, unsafe_allow_html=True)
        if st.button("Create My 7-Day Plan →", type="primary", use_container_width=True, key="mg_empty_generate"):
            start_date = datetime.now().date()
            end_date = start_date + timedelta(days=6)
            with st.spinner("Designing your week..."):
                response = call_meal_plan_api(
                    profile=profile,
                    query="Generate a complete 7-day meal plan based on my profile.",
                    token=st.session_state.auth_token,
                    force_tool_type="weekly_meal_plan_generator",
                    start_date=start_date.isoformat(),
                    end_date=end_date.isoformat(),
                )
            if response and isinstance(response.get("data"), dict):
                data = response["data"]
                st.session_state.weekly_meal_plan = data.get("answer", "")
                st.session_state.current_constraints = data.get("updated_constraints", profile)
                st.session_state.last_agent_context = data.get("updated_last_agent_context", st.session_state.last_agent_context)
                st.rerun()
            else:
                st.error("MG AI could not generate the seven-day meal plan. Your existing profile is preserved.")
        return

    selected = st.session_state.get("mg_ai_selected_day")
    available_days = [d["day"] for d in days]
    if selected not in available_days:
        selected = available_days[0]
        st.session_state.mg_ai_selected_day = selected
    selected_day = next(d for d in days if d["day"] == selected)

    date_range = [d["date"] for d in days if d.get("date")]
    range_text = f"{min(date_range).strftime('%b %d')} — {max(date_range).strftime('%b %d, %Y')}" if date_range else "7-day plan"
    calories = [d["daily_calories"] for d in days if d.get("daily_calories") is not None]
    avg_daily_cal = sum(calories) / len(calories) if calories else None
    name = _clean_ui_text(profile.get("name", "there"))

    st.markdown(f"""
    <section class="mg-hero">
        <div class="mg-eyebrow">MG AI · Your personal nutrition plan</div>
        <div class="mg-hero-title">Your week, intelligently designed.</div>
        <div class="mg-hero-subtitle">A premium seven-day nutrition experience built around {name}'s profile, preferences and the generated meal-plan data currently in your session.</div>
        <div style="display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px;margin-top:1.9rem;">
            <div class="mg-metric"><div class="mg-metric-label">Average day</div><div class="mg-metric-value">{_format_number(avg_daily_cal, " kcal")}</div><div class="mg-metric-unit">{range_text}</div></div>
            <div class="mg-metric"><div class="mg-metric-label">Days</div><div class="mg-metric-value">{len(days)}</div><div class="mg-metric-unit">personalised</div></div>
            <div class="mg-metric"><div class="mg-metric-label">Meal moments</div><div class="mg-metric-value">{len(selected_day.get("meals", {}))}</div><div class="mg-metric-unit">selected day</div></div>
            <div class="mg-metric"><div class="mg-metric-label">Components</div><div class="mg-metric-value">{sum(len(m.get("items", [])) for m in selected_day.get("meals", {}).values())}</div><div class="mg-metric-unit">selected day</div></div>
        </div>
    </section>
    """, unsafe_allow_html=True)

    a, b = st.columns([1, 1], gap="small")
    with a:
        if st.button("Explore My Week →", type="primary", use_container_width=True, key="mg_explore_week"):
            st.session_state.mg_ai_selected_day = days[0]["day"]
            st.rerun()
    with b:
        with st.expander("Edit profile", expanded=False):
            _render_edit_profile_form()

    st.markdown("<div style='height:20px'></div>", unsafe_allow_html=True)
    _render_profile_snapshot(profile)

    st.markdown("<div style='height:20px'></div>", unsafe_allow_html=True)
    st.markdown("""
    <div class="mg-section-kicker">Week at a glance</div>
    <div class="mg-section-title">Choose your day.</div>
    """, unsafe_allow_html=True)

    selector_cols = st.columns(len(days))
    for col, day in zip(selector_cols, days):
        with col:
            date_obj = day.get("date")
            label = f"{date_obj.strftime('%a').upper()}  {date_obj.day}" if date_obj else f"DAY {day.get('day')}"
            if st.button(label, key=f"mg_day_select_{day.get('day')}", type="primary" if day.get("day") == selected else "secondary", use_container_width=True):
                st.session_state.mg_ai_selected_day = day["day"]
                st.session_state.mg_ai_open_recipe = None
                st.rerun()

    st.markdown("<div style='height:16px'></div>", unsafe_allow_html=True)
    _render_daily_summary(selected_day)

    st.markdown("<div style='height:22px'></div>", unsafe_allow_html=True)
    st.markdown("""
    <div class="mg-section-kicker">Your day</div>
    <div class="mg-section-title">Five moments. Thoughtfully composed.</div>
    <div class="mg-section-copy">Start with the visual summary. Open any meal for the complete ingredients, recipe and advanced nutrition information.</div>
    """, unsafe_allow_html=True)

    meals = [selected_day["meals"][m] for m in MEAL_TYPES if m in selected_day.get("meals", {})]
    open_recipe = st.session_state.get("mg_ai_open_recipe")
    for row_start in range(0, len(meals), 2):
        row = meals[row_start:row_start + 2]
        cols = st.columns(len(row), gap="large")
        for local_idx, (col, meal) in enumerate(zip(cols, row)):
            with col:
                if _render_meal_card(meal, row_start + local_idx):
                    st.session_state.mg_ai_open_recipe = (selected_day["day"], meal["type"])
                    st.rerun()
        if row_start + 2 < len(meals):
            st.markdown("<div style='height:16px'></div>", unsafe_allow_html=True)

    open_recipe = st.session_state.get("mg_ai_open_recipe")
    if open_recipe:
        open_day = next((d for d in days if d["day"] == open_recipe[0]), None)
        open_meal = open_day.get("meals", {}).get(open_recipe[1]) if open_day else None
        if open_meal:
            st.markdown("<div style='height:26px'></div>", unsafe_allow_html=True)
            st.markdown("""
            <div class="mg-section-kicker">Recipe studio</div>
            <div class="mg-section-title">Details, when you want them.</div>
            """, unsafe_allow_html=True)
            detail, close = st.columns([5, 1])
            with close:
                if st.button("Close", key="mg_close_recipe", use_container_width=True):
                    st.session_state.mg_ai_open_recipe = None
                    st.rerun()
            with detail:
                day_date = open_day.get("date")
                day_label = day_date.strftime("%A, %B %d, %Y") if day_date else open_day.get("title", "")
                _render_recipe_details(open_meal, day_label)

    st.markdown("<div style='height:32px'></div>", unsafe_allow_html=True)
    _render_nutrition_intelligence(selected_day)

    st.markdown("<div style='height:26px'></div>", unsafe_allow_html=True)
    _render_ai_insight(selected_day, profile)

    st.markdown("<div style='height:34px'></div>", unsafe_allow_html=True)
    st.markdown("""
    <div class="mg-section-kicker">Create another direction</div>
    <div class="mg-section-title">Regenerate your week ✦</div>
    <div class="mg-section-copy">Keep the same profile and ask MG AI to explore a different style of week.</div>
    """, unsafe_allow_html=True)

    regen_choice = st.selectbox(
        "Optional direction",
        [
            "Keep my current style", "More variety", "Higher protein", "Lower calories",
            "Quick & easy meals", "More Indian cuisine", "Mediterranean-inspired",
            "Comfort food", "Surprise me ✦"
        ],
        key="mg_regen_choice",
        label_visibility="collapsed"
    )
    regen_notes = st.text_area(
        "Additional instructions",
        value="",
        height=90,
        placeholder="Example: keep the same calories, change dinner options, reduce repetition...",
        key="mg_regen_notes",
        label_visibility="collapsed"
    )

    if st.button("Regenerate My Week ✦", type="primary", use_container_width=True, key="mg_regenerate_premium"):
        direction = "" if regen_choice == "Keep my current style" else f"Style direction: {regen_choice}."
        regeneration_query = (
            "Generate a fresh complete 7-day meal plan based on my profile. "
            "Apply the user's requested direction and notes where applicable.\n"
            f"{direction}\nAdditional notes:\n{regen_notes}"
        )
        start_date = datetime.now().date()
        end_date = start_date + timedelta(days=6)
        with st.spinner("Designing a new week..."):
            response = call_meal_plan_api(
                profile=st.session_state.current_constraints,
                query=regeneration_query,
                token=st.session_state.auth_token,
                force_tool_type="weekly_meal_plan_generator",
                start_date=start_date.isoformat(),
                end_date=end_date.isoformat()
            )
        if response and isinstance(response.get("data"), dict):
            data = response["data"]
            st.session_state.weekly_meal_plan = data.get("answer", "")
            st.session_state.current_constraints = data.get("updated_constraints", st.session_state.current_constraints)
            st.session_state.last_agent_context = data.get("updated_last_agent_context", st.session_state.last_agent_context)
            st.session_state.mg_ai_open_recipe = None
            st.session_state.mg_ai_selected_day = None
            st.rerun()
        else:
            st.error("The new seven-day plan could not be generated. Your current plan is still available.")

    st.markdown("""
    <div style="margin-top:3rem;padding-top:1.4rem;border-top:1px solid rgba(21,21,21,0.08);text-align:center;color:#8c8c86;">
        <div style="font-family:Georgia,serif;font-size:1.35rem;color:#151515;">MG AI</div>
        <div style="font-size:0.68rem;letter-spacing:0.14em;text-transform:uppercase;margin-top:.3rem;">Personal nutrition intelligence</div>
    </div>
    """, unsafe_allow_html=True)


def render_unified_profile_form():
    st.markdown("""
    <section class="mg-form-hero">
        <div class="mg-eyebrow" style="color:#c9a86a;">MG AI · Personal nutrition intelligence</div>
        <h1>Build your nutrition profile.</h1>
        <p>Tell MG AI what matters to you. Your profile becomes the foundation for a personalised seven-day meal experience.</p>
    </section>
    """, unsafe_allow_html=True)
    st.markdown("### Your profile")
    st.caption("Complete the fields that are relevant to you. Detailed nutrition inputs remain available without crowding the main experience.")
    
    with st.form(key="unified_profile_form"):

        st.markdown("---")
        st.subheader("🥝 Nutrition Assistant Profile")
        
        col1, col2 = st.columns(2)

        with col1:
            st.markdown("##### Personal Details")
            nut_user_name_input = st.text_input("First Name", value="Jansen", key="nut_name")
            nut_age = st.number_input("Age", min_value=18, max_value=100, value=30, step=1, key="nut_age")
            nut_weight = st.number_input("Weight (kg)", min_value=30.0, max_value=300.0, value=75.0, step=0.1, format="%.1f", key="nut_weight")
            
            h_col1, h_col2 = st.columns(2)
            with h_col1:
                nut_height_feet = st.number_input("Height (feet)", min_value=4, max_value=8, value=5, step=1, key="nut_height_ft")
            with h_col2:
                nut_height_inches = st.number_input("Height (inches)", min_value=0, max_value=11, value=9, step=1, key="nut_height_in")
            
            nut_waist_circumference = st.number_input("Waist Circumference (inches)", min_value=20.0, max_value=80.0, value=35.0, step=0.1, format="%.1f", key="nut_waist")

        with col2:
            st.markdown("##### Dietary & Lifestyle")
            nut_gender = st.selectbox("Gender", options=["Male", "Female"], index=0, key="nut_gender")
            nut_activity_level = st.selectbox("Activity Level", options=["Sedentary", "Lightly active", "Moderately active", "Very active", "Extra active"], index=2, key="nut_activity")
            nut_dietary_preference = st.selectbox("Dietary Preference", options=DIETARY_PREFERENCE_OPTIONS, key="nut_diet_pref")
            nut_dietary_restrictions = st.multiselect("Dietary Restrictions", options=DIETARY_RESTRICTION_OPTIONS, key="nut_diet_rest")
            
            nut_medical_conditions = st.multiselect("Medical Conditions", options=MEDICAL_CONDITIONS_NUTRITION_OPTIONS, key="nut_med_cond")
            nut_digestive_issues = st.multiselect("Digestive Issues", options=DIGESTIVE_ISSUES_OPTIONS, key="nut_dig_issues")
            nut_symptom_aggravating_foods = st.multiselect("Symptom-Aggravating Foods", options=SYMPTOM_AGGRAVATING_FOODS_OPTIONS, key="nut_symptom_foods")
            nut_cuisine = st.multiselect("Preferred Cuisine", options=CUISINE_OPTIONS, key="nut_cuisine")
            nut_food_allergies = st.multiselect("Food Allergies", options=FOOD_ALLERGIES_OPTIONS, key="nut_allergies")

        st.markdown("##### Upload Blood Report (Optional)")
        nut_uploaded_blood_report = st.file_uploader("Upload a blood report (PDF, JPG, JPEG, PNG, DOCX, CSV, TXT)", type=["pdf", "jpg", "jpeg", "png", "docx", "csv", "txt"], key="nut_uploader")

        st.markdown("##### Vital Signs (Optional, numeric values only)")
        
        v_col1_r1, v_col2_r1 = st.columns(2)
        with v_col1_r1:
            nut_hr = st.number_input("Heart Rate (beats per minute)", value=80, key="nut_hr")
        with v_col2_r1:
            nut_bp_sys = st.number_input("Blood Pressure - Systolic (mmHg)", value=120, key="nut_bps")

        v_col1_r2, v_col2_r2 = st.columns(2)
        with v_col1_r2:
            nut_glucose = st.number_input("Blood Glucose (mg/dL)", value=90, key="nut_glucose")
        with v_col2_r2:
            nut_bp_dia = st.number_input("Blood Pressure - Diastolic (mmHg)", value=80, key="nut_bpd")

        v_col1_r3, v_col2_r3 = st.columns(2)
        with v_col1_r3:
            nut_resp_rate = st.number_input("Respiration Rate (breaths per minute)", value=16, key="nut_rr")
        with v_col2_r3:
            nut_blood_oxygen = st.number_input("Blood Oxygen Saturation (%)", value=98, key="nut_bo")

        v_col1_r4, v_col2_r4 = st.columns(2)
        with v_col1_r4:
            nut_blood_ketones = st.number_input("Blood Ketones (mmol/L)", value=0.3, format="%.2f", key="nut_bk")
        with v_col2_r4:
            nut_body_temp = st.number_input("Body Temperature (°C)", value=36.7, format="%.1f", key="nut_bt")

        
        v_col1_r5, v_col2_r5 = st.columns(2)
        with v_col1_r5:
            nut_body_fat_percentage = st.number_input("Body Fat %", value=20.0, format="%.1f", min_value=0.0, max_value=100.0, key="nut_bfp")
        with v_col2_r5:
            pass

        nut_submitted = st.form_submit_button("Design My 7-Day Plan ✦", type="primary", use_container_width=True, help="Creates or updates your Nutrition Profile and generates a personalized seven-day meal plan.", key="nut_submit_btn")

        # Optional profile sections are intentionally disabled for this page.
        fit_submitted = False
        ws_submitted = False
        booking_submitted = False
        _disabled_optional_profile_sections = """

        st.markdown("---")
        st.subheader("🏋️ Personalized Fitness Profile")
        
        col_f1, col_f2, col_f3 = st.columns(3)
        with col_f1:
            fit_name = st.text_input("Fitness Name *", placeholder="Your name", key="fit_name_input")
        with col_f2:
            fit_age = st.number_input("Fitness Age *", min_value=13, max_value=100, value=30, key="fit_age_input")
        with col_f3:
            fit_gender = st.selectbox("Fitness Gender *", ["Male", "Female", "Other"], key="fit_gender_input")
        
        col_f4, col_f5, col_f6 = st.columns(3)
        with col_f4:
            fit_unit_system = st.radio("Units *", ["Metric (kg, cm)", "Imperial (lbs, in)"], key="fit_unit_input")
            
            fit_weight_kg = 0.0
            fit_height_cm = 0.0

        with col_f5:
            if fit_unit_system == "Metric (kg, cm)":
                fit_weight_metric = st.number_input("Fitness Weight (kg) *", min_value=30.0, max_value=300.0, value=70.0, key="fit_weight_kg_input")
                fit_weight_kg = fit_weight_metric
            else:
                fit_weight_lbs = st.number_input("Fitness Weight (lbs) *", min_value=66.0, max_value=660.0, value=154.3, key="fit_weight_lbs_input")
                fit_weight_kg = fit_weight_lbs * 0.453592

        with col_f6:
            if fit_unit_system == "Metric (kg, cm)":
                fit_height_metric = st.number_input("Fitness Height (cm) *", min_value=100.0, max_value=250.0, value=170.0, key="fit_height_cm_input")
                fit_height_cm = fit_height_metric
            else:
                fit_height_in = st.number_input("Fitness Height (in) *", min_value=39.0, max_value=98.0, value=66.9, key="fit_height_in_input")
                fit_height_cm = fit_height_in * 2.54

        st.markdown("##### Fitness Goals & Level")
        
        col_f7, col_f8 = st.columns(2)
        with col_f7:
            fit_primary_goal = st.selectbox(
                "Primary Goal *",
                FITNESS_GOAL_OPTIONS, key="fit_primary_goal_input"
            )
        
        with col_f8:
            fit_secondary_goal = st.selectbox(
                "Secondary Goal (Optional)",
                ["None"] + FITNESS_GOAL_OPTIONS, key="fit_secondary_goal_input"
            )
        
        fit_fitness_level = st.selectbox(
            "Fitness Level *",
            FITNESS_LEVEL_OPTIONS, key="fit_fitness_level_input"
        )
        
        st.markdown("##### Health Screening")
        
        fit_medical_conditions = st.multiselect(
            "Fitness Medical Conditions *",
            FITNESS_MEDICAL_CONDITIONS_OPTIONS, 
            default=["None"], 
            key="fit_medical_conditions_input"
        )
        
        st.warning("⚠️ **Physical Limitations** - Describe ANY injuries, pain, or movement restrictions")
        fit_physical_limitation = st.text_area( 
            "Physical Limitations (Important for Safety) *",
            placeholder="E.g., 'Previous right knee surgery - avoid deep squats'",
            height=100, key="fit_physical_limitation_input"
        )
        
        # --- NEW: Specific Exercise Avoidance ---
        st.warning("⚠️ **Specific Exercise Restrictions**")
        fit_specific_avoidance_input = st.text_area(
            "Have you been advised to avoid any specific exercises? (List them below if known to avoid in plan generation):",
            placeholder="E.g., 'Heavy deadlifts, overhead pressing due to shoulder issue, any exercise that causes sharp pain in the elbow.'",
            height=100,
            key="specific_avoidance_text_input"
        )
        # --- END NEW: Specific Exercise Avoidance ---
        
        st.markdown("##### Training Schedule & Equipment")
        
        col_f9, col_f10 = st.columns(2)
        with col_f9:
            fit_days_per_week = st.multiselect(
                "Training Days *",
                FITNESS_DAYS_OPTIONS,
                default=["Monday", "Wednesday", "Friday"], key="fit_days_per_week_input"
            )
        
        with col_f10:
            fit_session_duration = st.selectbox(
                "Session Duration *",
                FITNESS_DURATION_OPTIONS, key="fit_session_duration_input"
            )

        fit_workout_location = st.selectbox(
            "Where will you primarily work out? *",
            FITNESS_LOCATION_OPTIONS, key="fit_location_input"
        )
        
        fit_equipment = st.multiselect("Select all available equipment: *", FITNESS_EQUIPMENT_OPTIONS, default=["Bodyweight Only"], key="fit_equipment_input")

        fit_submitted = st.form_submit_button("🚀 Generate Fitness Plan", type="secondary", use_container_width=True, help="Creates or updates your Fitness Profile and generates a personalized Workout Plan.", key="fit_submit_btn")


        st.markdown("---")
        st.subheader("💧 Water & Step Goal Profile")
        st.caption("These fields are used to calculate personalized daily step and water targets.")
        
        current_constraints = st.session_state.current_constraints
        default_age_ws = current_constraints.get('age', 30)
        default_weight_ws = current_constraints.get('weight_kg', 70.0)
        default_gender_ws = current_constraints.get('gender', 'Male')
        default_height_ws = current_constraints.get('height_cm', 170.0)
        default_conditions_ws = current_constraints.get('medical_conditions', [])
        default_goal_ws = current_constraints.get('goal', {}).get('type', 'Maintenance')
        
        st.markdown("##### Core Profile Details (Manual Entry for Calculation)")
        
        col_ws_1, col_ws_2 = st.columns(2)
        
        with col_ws_1:
            ws_age = st.number_input("WS Age *", min_value=13, max_value=100, value=default_age_ws, step=1, key="ws_age_input")
            ws_weight = st.number_input("WS Weight (kg) *", min_value=30.0, max_value=300.0, value=default_weight_ws, step=0.1, format="%.1f", key="ws_weight_input")
            ws_gender = st.selectbox("WS Gender *", options=["Male", "Female", "Other"], index=(0 if default_gender_ws == 'Male' else 1) if default_gender_ws in ['Male', 'Female'] else 2, key="ws_gender_input")
            
        with col_ws_2:
            ws_height = st.number_input("WS Height (cm) *", min_value=100.0, max_value=250.0, value=default_height_ws, step=0.1, format="%.1f", key="ws_height_input")
            
            ws_bmi = ws_weight / ((ws_height / 100) ** 2) if ws_weight > 0 and ws_height > 0 else 0.0
            st.markdown(f"Calculated BMI: **{ws_bmi:.2f}**")
            
            ws_medical_conditions = st.multiselect("WS Medical Conditions (for calculation)", 
                                                 options=["None", "Diabetes Type 2", "Hypertension", "Pregnancy", "Lactation"], 
                                                 default=[c for c in default_conditions_ws if c in ["Diabetes Type 2", "Hypertension", "Pregnancy", "Lactation"]], 
                                                 key="ws_med_cond_input")

        st.markdown("##### 👣 Step Goal Parameters")
        
        step_occupations = ["Sedentary (office)", "Moderate (teacher, retail, etc.)", "Active job (labor, delivery, etc.)"]
        ws_occupation = st.selectbox("Occupation (Affects Base Activity)", step_occupations, key="ws_occupation")
        
        step_fitness_lvls = ["Beginner", "Intermediate", "Advanced"]
        ws_activity = st.selectbox("Current Fitness Level (WS)", step_fitness_lvls, key="ws_activity")
        
        step_goals = ["Weight loss", "Maintenance", "Weight Gain", "Health management (e.g., diabetes)"]
        ws_goal = st.selectbox("Primary Health Goal (WS)", step_goals, index=step_goals.index(default_goal_ws) if default_goal_ws in step_goals else 1, key="ws_goal")
        
        st.markdown("##### 💧 Water Goal Parameters")
        
        ws_climates = ["Moderate / Temperate", "Hot / Humid (>30°C)", "High Altitude (>1500m)", "Dry / Air-Conditioned"]
        ws_climate = st.selectbox("Typical Climate/Environment", ws_climates, key="ws_climate")
        
        ws_diets = ["Standard", "High Protein (>1.6g/kg)", "High Fiber", "Salty / Spicy", "Low-carb/Keto"]
        ws_diet = st.selectbox("Dietary Intake Pattern (WS)", ws_diets, key="ws_diet")

        ws_submitted = st.form_submit_button("Save and Calculate Goals", type="secondary", use_container_width=True, help="Calculates and saves your Step and Water intake goals to your profile.", key="ws_submit_btn")

        
        st.markdown("---")
        st.subheader("📅 Expert Session Booking")
        st.caption("Fill in your preferred session details to automatically trigger the booking conversation.")

        default_user_name = st.session_state.current_constraints.get('name', 'User')
        
        expert_name = st.selectbox("Select Expert", EXPERT_NAMES, key="booking_expert")
        session_type = st.selectbox("Session Focus", SESSION_TYPES, key="booking_focus")
        appointment_type = st.selectbox("Appointment Type", APPOINTMENT_TYPES, key="booking_type")
        
        date_col, time_col = st.columns(2)
        with date_col:
            session_date = st.date_input("Preferred Date", key="booking_date")
        with time_col:
            session_time = st.time_input("Preferred Time", value=datetime.now().time(), key="booking_time")

        booking_submitted = st.form_submit_button("Book Session Now", type="secondary", use_container_width=True, help="Initiates a conversational flow to book the specified session.", key="booking_submit_btn")
        """


    if nut_submitted:
        with st.spinner("Designing your personalised week..."):

            blood_report_path = None
            if nut_uploaded_blood_report:
                blood_report_path = save_uploaded_file_temp(nut_uploaded_blood_report)
                st.session_state.blood_report_temp_path = blood_report_path

            nut_height_cm = (nut_height_feet * 12 + nut_height_inches) * 2.54

            profile_data = {
                "name": nut_user_name_input,
                "age": nut_age,
                "gender": nut_gender,
                "weight_kg": nut_weight,
                "height_cm": nut_height_cm,
                "waist_circumference_cm": nut_waist_circumference * 2.54, 
                "activity_level": nut_activity_level.lower(),
                "dietary_preference": [nut_dietary_preference], 
                "restrictions": nut_dietary_restrictions,
                "medical_conditions": nut_medical_conditions,
                "allergies": nut_food_allergies,
                "cuisine": nut_cuisine,
                "digestive_issues": nut_digestive_issues,
                "symptom_aggravating_foods": nut_symptom_aggravating_foods,
                "vitals_numeric": {
                    "Heart Rate": nut_hr, 
                    "Blood Pressure": f"{nut_bp_sys}/{nut_bp_dia}", 
                    "Blood Glucose": nut_glucose,
                    "Respiration Rate": nut_resp_rate,
                    "Blood Oxygen Saturation": nut_blood_oxygen,
                    "Blood Ketones": nut_blood_ketones,
                    "Body Temperature": nut_body_temp,
                    "Body Fat %": nut_body_fat_percentage
                }
            }

            profile_data = {
                key: value
                for key, value in profile_data.items()
                if key in PROFILE_CONSTRAINT_FIELDS
            }

            st.session_state.current_constraints = profile_data.copy()
            plan_start_date = datetime.now().date()
            plan_end_date = plan_start_date + timedelta(days=6)
            plan_start = plan_start_date.isoformat()
            plan_end = plan_end_date.isoformat()
            initial_prompt = "Create a complete 7-day meal plan for me based on my profile."

            # Send the frontend-collected profile directly to the backend so it is
            # used as the meal-plan constraints, bypassing the DB profile lookup/auth.
            response_data = call_meal_plan_api(
                profile=profile_data,
                query=initial_prompt,
                token=st.session_state.auth_token,
                force_tool_type="weekly_meal_plan_generator",
                start_date=plan_start,
                end_date=plan_end
            )

            if response_data and "data" in response_data:
                ai_content = response_data["data"]
                st.session_state.weekly_meal_plan = ai_content.get("answer", "")
                st.session_state.messages.append({
                    "role": "assistant", 
                    "content": ai_content.get("answer"), 
                    "audio": ai_content.get("audio"),
                    "context": ai_content.get("context")
                })
                
                st.session_state.current_constraints = ai_content.get("updated_constraints", st.session_state.current_constraints)
                st.session_state.last_agent_context = ai_content.get("updated_last_agent_context", {})
                
                st.session_state.profile_complete = True
                st.rerun()

    _disabled_optional_profile_handlers = """
    elif fit_submitted:
        if not fit_name.strip():
            st.error("❌ Please enter your name.")
        elif not fit_days_per_week:
            st.error("❌ Please select at least one training day.")
        elif (fit_weight_kg <= 0 and fit_unit_system == "Metric (kg, cm)") or (fit_height_cm <= 0 and fit_unit_system == "Metric (kg, cm)"):
            st.error("❌ Please ensure valid weight and height inputs.")
        else:
            sanitized_med_conditions = [mc for mc in fit_medical_conditions if mc != "None"]
            if not sanitized_med_conditions:
                sanitized_med_conditions = ["None"]
                
            sanitized_equipment = fit_equipment if fit_equipment else ["Bodyweight Only"]

            fit_weight_kg_final = fit_weight_kg
            fit_height_cm_final = fit_height_cm
            if fit_unit_system == "Imperial (lbs, in)":
                fit_weight_kg_final = st.session_state["fit_weight_lbs_input"] * 0.453592
                fit_height_cm_final = st.session_state["fit_height_in_input"] * 2.54

            fit_profile = {
                "name": fit_name.strip(),
                "age": fit_age,
                "gender": fit_gender,
                "weight_kg": fit_weight_kg_final,
                "height_cm": fit_height_cm_final,
                "bmi": fit_weight_kg_final / ((fit_height_cm_final / 100) ** 2) if fit_weight_kg_final > 0 and fit_height_cm_final > 0 else 0,
                "primary_goal": fit_primary_goal,
                "secondary_goal": fit_secondary_goal if fit_secondary_goal != "None" else None,
                "fitness_level": fit_fitness_level,
                "medical_conditions": sanitized_med_conditions,
                "physical_limitation": fit_physical_limitation.strip(),
                "days_per_week": fit_days_per_week,
                "session_duration": fit_session_duration,
                "available_equipment": sanitized_equipment,
                "workout_location": fit_workout_location,
                "specific_avoidance": fit_specific_avoidance_input.strip() or "None" # CRITICAL: Include new avoidance field
            }
            
            st.session_state.fitness_profile_data = fit_profile
            
            st.session_state.current_constraints.update({
                "name": fit_name.strip(),
                "age": fit_age,
                "gender": fit_gender,
                "weight_kg": fit_weight_kg_final,
                "height_cm": fit_height_cm_final,
                "primary_goal": fit_primary_goal,
                "fitness_level": fit_fitness_level,
                "days_per_week": fit_days_per_week,
                "available_equipment": sanitized_equipment,
                "physical_limitation": fit_physical_limitation.strip(), 
                "specific_avoidance": fit_specific_avoidance_input.strip() or "None", # CRITICAL: Include new avoidance field in constraints
                "medical_conditions": sanitized_med_conditions,
                "fitness_profile_data": fit_profile
            })
            
            initial_prompt = f"Generate a detailed weekly workout plan based on my profile, starting now."
            
            st.session_state.messages = [] 
            st.session_state.messages.append({"role": "user", "content": initial_prompt})
            
            with st.spinner("Generating your comprehensive fitness plan..."):
                api_payload = {
                    "query": initial_prompt,
                    "chat_history": st.session_state.messages, 
                    "current_constraints": st.session_state.current_constraints,
                    "last_agent_context": st.session_state.last_agent_context,
                    "force_tool_type": "fitness_plan_generator" 
                }
                
                response_data = call_backend_api(api_payload, st.session_state.auth_token)

                if response_data:
                    st.session_state.messages.append({
                        "role": "assistant", 
                        "content": response_data.get("answer"), 
                        "audio": response_data.get("audio"), 
                        "context": response_data.get("context")
                    })
                    st.session_state.current_constraints = response_data.get("updated_constraints", {})
                    st.session_state.last_agent_context = response_data.get("updated_last_agent_context", {})
                
                st.session_state.profile_complete = True
                st.rerun()

    elif ws_submitted:
        if ws_weight <= 30.0 or ws_height <= 100.0:
            st.error("❌ Please provide realistic weight and height inputs for Water/Step calculation.")
            return

        ws_update_payload = {
            "age": ws_age, 
            "gender": ws_gender, 
            "weight_kg": ws_weight, 
            "height_cm": ws_height, 
            "medical_conditions": [c for c in ws_medical_conditions if c != "None"], 
            "occupation": ws_occupation,
            "activity_level": ws_activity, 
            "goal": {"type": ws_goal},
            "climate": ws_climate,
            "diet_type": ws_diet
        }
        
        st.session_state.current_constraints.update(ws_update_payload)
        
        initial_prompt = "What are my daily step and water goals?"
        
        st.session_state.messages.append({"role": "user", "content": initial_prompt})
        
        send_message_to_backend(initial_prompt, force_tool="water_step_advisor", images=None)
        
        st.session_state.messages.append({"role": "assistant", "content": "Water and Step calculation parameters saved. You can now ask: **'What is my daily step goal?'** or **'How much water should I drink?'**"})
        st.rerun()
        
    elif booking_submitted:
        user_query = (
            f"I want to book a {session_type} session with {expert_name} on "
            f"{session_date.strftime('%B %d, %Y')} at {session_time.strftime('%I:%M %p')} for a {appointment_type}. "
            f"My name is {default_user_name}."
        )

        st.session_state.messages.append({"role": "user", "content": user_query})
        
        send_message_to_backend(user_query, force_tool="session_booking")
        
        st.session_state.profile_complete = True
        st.rerun()

    """



if not st.session_state.profile_complete:
    render_unified_profile_form()
else:
    render_premium_meal_plan_dashboard()

    st.sidebar.title("MG AI Controls")
    st.sidebar.title("Developer Diagnostics")

    last_assistant_message = None
    for msg in reversed(st.session_state.messages):
        if msg.get("role") == "assistant":
            last_assistant_message = msg
            break

    if last_assistant_message:
        original_text = last_assistant_message.get("content", "N/A")
        summarized_text = last_assistant_message.get("audio")

        st.sidebar.subheader("Latest AI response")
        st.sidebar.text_area("Original", value=original_text, height=200, key="debug_original", disabled=True)

        # Updated Debugging for the Base64 Audio data
        st.sidebar.subheader("Audio data (developer view)")
        if summarized_text:
            st.sidebar.text_area("Base64 String", value=summarized_text[:100] + "..." if len(summarized_text) > 100 else summarized_text, height=100, key="debug_summary", disabled=True)
            if len(summarized_text) < 50:
                 st.sidebar.warning("Note: The 'audio' field is short. This might be a mistake if Base64 audio was expected.")
        else:
            st.sidebar.info("No audio data was sent for the last message.")
    else:
        st.sidebar.info("No assistant message yet to display.")

    st.sidebar.divider()
    
    if st.sidebar.button("Start Over"):
        token = st.session_state.get("auth_token", "")
        if st.session_state.blood_report_temp_path and os.path.exists(st.session_state.blood_report_temp_path):
            try:
                os.remove(st.session_state.blood_report_temp_path)
                print(f"Cleaned up temporary blood report file: {st.session_state.blood_report_temp_path}")
            except Exception as e:
                print(f"Error cleaning up temporary blood report file: {e}")
        
        for key in list(st.session_state.keys()):
            if key not in ["auth_token"]:
                del st.session_state[key]
        st.session_state.profile_complete = False
        st.rerun()
    
    if st.sidebar.button("Save Chat Logs"):
        doc = Document()
        profile = st.session_state.get('current_constraints', {})
        doc.add_heading(f"MG AI Chat History for {profile.get('name', 'User')}", level=1)
        if profile:
            doc.add_heading("User Profile", level=2)
            profile_text_parts = []
            for k, v in profile.items():
                if k not in ['vitals_numeric', 'vitals_text']:
                    key_title = k.replace('_', ' ').title()
                    value_str = ", ".join(v) if isinstance(v, list) else str(v)
                    profile_text_parts.append(f"{key_title}: {value_str}")
            
            if 'vitals_numeric' in profile and profile['vitals_numeric']:
                profile_text_parts.append("\nVitals:")
                for vital_key, vital_value in profile['vitals_numeric'].items():
                        profile_text_parts.append(f"- {vital_key.replace('_', ' ').title()}: {vital_value}")

            doc.add_paragraph("\n".join(profile_text_parts))
        
        doc.add_heading("Conversation", level=2)
        for msg in st.session_state.messages:
            content = msg.get('content', '')
            if isinstance(content, dict):
                content = json.dumps(content, indent=2)

            clean_content = re.sub(r'<.*?>', '', content)
            clean_content = clean_content.replace("**", "")
            doc.add_paragraph(f"{msg['role'].title()}: {clean_content}", style='Body Text')
        
        doc_io = BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)
        
        st.sidebar.download_button(
            label="Download Chat Logs as .docx",
            data=doc_io,
            file_name=f"mg_ai_chat_{profile.get('name', 'user').lower()}_history.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )