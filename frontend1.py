import streamlit as st
import requests
import json
from docx import Document
from io import BytesIO
import re
from typing import Optional, Dict
import os
import tempfile
import base64
import binascii
from datetime import datetime, timedelta
from audio_recorder_streamlit import audio_recorder
try:
    from helpers.meal_plan_component import MealPlanRenderer
except Exception:
    MealPlanRenderer = None

try:
    from helpers.premium_meal_plan_parser import parse_weekly_plan
except Exception:
    parse_weekly_plan = None

try:
    from helpers import premium_ui
except Exception:
    premium_ui = None

# Backend base URL is configurable so the deployed frontend (e.g. Streamlit Cloud)
# can point at a publicly reachable backend instead of localhost.
try:
    _backend_base_url = st.secrets.get("BACKEND_BASE_URL", os.getenv("BACKEND_BASE_URL", "http://127.0.0.1:7000"))
except Exception:
    _backend_base_url = os.getenv("BACKEND_BASE_URL", "http://127.0.0.1:7000")
_backend_base_url = _backend_base_url.rstrip("/")

BACKEND_API_URL = f"{_backend_base_url}/chat"
# Generates a meal plan directly from the frontend-collected profile, bypassing DB profile lookup/auth.
MEAL_PLAN_FROM_PROFILE_API_URL = f"{_backend_base_url}/generate-meal-plan-from-profile"

FITNESS_GOAL_OPTIONS = ["Weight Loss", "Muscle Gain", "Increase Overall Strength", "Improve Cardiovascular Fitness", "Improve Flexibility & Mobility", "Rehabilitation & Injury Prevention", "Improve Posture and Balance", "General Fitness", "Weight Maintenance"]
FITNESS_LEVEL_OPTIONS = ["Beginner (0–6 months)", "Intermediate (6 months–2 years)", "Advanced (2+ years)"]
FITNESS_MEDICAL_CONDITIONS_OPTIONS = ["None", "Hypertension (High Blood Pressure)", "Type 2 Diabetes", "Osteoarthritis", "Chronic Lower Back Pain", "Other"]
FITNESS_DAYS_OPTIONS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
FITNESS_DURATION_OPTIONS = ["15-20 minutes", "20-30 minutes", "30-45 minutes", "45-60 minutes"]
FITNESS_LOCATION_OPTIONS = ["Home", "Gym", "Outdoor", "Any"]
FITNESS_EQUIPMENT_OPTIONS = ["Bodyweight Only", "Dumbbells", "Resistance Bands", "Kettlebells", "Barbell", "Bench", "Pull-up Bar", "Yoga Mat", "Machines"]

DIETARY_PREFERENCE_OPTIONS = [
    "Vegetarian", "Vegan", "Non Vegetarian", "No Pork", "No Beef", "No Red Meat"
]
DIETARY_RESTRICTION_OPTIONS = ["Dairy-Free", "Lactose Intolerance", "Gluten-Free"]
MEDICAL_CONDITIONS_NUTRITION_OPTIONS = [
    "Alzheimer’s Disease", "Anemia", "Cardiovascular Disease", "Constipation",
    "Coronary Artery Disease", "Crohn’s Disease", "Depression", "Heart Failure",
    "High Cholesterol", "Hypertension", "Hyperthyroidism", "Hypothyroidism",
    "Liver Disease", "Metabolic Syndrome", "Non-Alcoholic Fatty Liver Disease (NAFLD)",
    "Obesity", "Polycystic Ovary Syndrome (PCOS)", "Type 1 Diabetes", "Type 2 Diabetes"
]
DIGESTIVE_ISSUES_OPTIONS = [
    "Occasional bloating", "Constipation", "Diarrhea", "Stomach pain or cramping",
    "Nausea or vomiting", "Excessive gas"
]
SYMPTOM_AGGRAVATING_FOODS_OPTIONS = [
    "Fried or greasy foods", "Spicy foods", "Caffeinated beverages",
    "Carbonated drinks", "Alcohol", "High-fiber foods", "Raw vegetables",
    "Chocolate", "Garlic and onions", "Processed or packaged foods"
]
CUISINE_OPTIONS = [
    "North American", "African", "Central American", "Chinese", "European", "Indian",
    "Indian South", "International", "Italian", "Japanese", "Korean", "Latin American",
    "Middle Eastern", "North Indian", "South American", "Spanish", "Thai", "Turkish",
    "Vietnamese"
]
FOOD_ALLERGIES_OPTIONS = [
    "Tree Nuts (e.g. Almonds, Walnuts, Cashews)", "Shellfish", "Fish", "Eggs",
    "Milk (Dairy Allergy)", "Wheat", "Soy", "Sesame", "Corn",
    "Legumes (e.g. Lentils, Chickpeas)", "Peanuts"
]

PROFILE_CONSTRAINT_FIELDS = {
    "name", "age", "gender", "height_cm", "weight_kg", "activity_level",
    "waist_circumference_cm", "cuisine", "symptom_aggravating_foods",
    "dietary_preference", "restrictions", "allergies", "medical_conditions",
    "digestive_issues", "vitals_numeric"
}

EXPERT_NAMES = ["Dr. Anjali Singh", "Dr. Amit Patel"]
SESSION_TYPES = ["Yoga", "Stress Relief", "Nutritional Advice"]
APPOINTMENT_TYPES = ["Video Call", "Audio Call"]

st.set_page_config(page_title="Personalised Meal Plan Generator", layout="wide")

st.markdown("""
<style>
    .meal-summary {
        font-size: 30px !important;
        font-weight: bold !important;
        color: #FFFFFF;
    }
    .stButton>button {
        width: 100%;
        margin-bottom: 5px;
        background-color: #4CAF50;
        color: white;
        padding: 10px 15px;
        border: none;
        border-radius: 8px;
        cursor: pointer;
        font-size: 16px;
        transition: background-color 0.3s ease;
        box-shadow: 2px 2px 5px rgba(0,0,0,0.2);
    }
    .stButton>button:hover {
        background-color: #45a049;
    }
    .stButton>button:active {
        background-color: #3e8e41;
        box_shadow: 0 0 0 rgba(0,0,0,0.2);
        transform: translateY(1px);
    }
    .unified-button-nutrition > button {
        background-color: #4CAF50 !important;
    }
    .unified-button-fitness > button {
        background-color: #764ba2 !important;
    }
    .unified-button-waterstep > button {
        background-color: #3e8e41 !important;
    }
    .unified-button-booking > button {
        background-color: #007bff !important;
    }
    /* Custom style for the combined input area */
    /* st.chat_input uses a complex structure. Targeting the elements within the column is safer. */
    .stAudioRecorder button {
        /* Ensure the microphone button is visually distinct and prominent */
        background-color: #dc3545 !important; 
        border-radius: 50%;
        width: 40px;
        height: 40px;
        padding: 0;
        display: flex;
        align-items: center;
        justify-content: center;
    }
    .stAudioRecorder button:hover {
        background-color: #c82333 !important;
    }
</style>
""", unsafe_allow_html=True)


if premium_ui is not None:
    premium_ui.inject_premium_css()
    premium_ui.render_top_nav()
else:
    st.set_page_config(page_title="Personalised Meal Plan Generator", layout="wide")
    st.title("Personalised Meal Plan Generator")

if "messages" not in st.session_state:
    st.session_state.messages = []
if "profile_complete" not in st.session_state:
    st.session_state.profile_complete = False
if "current_constraints" not in st.session_state:
    st.session_state.current_constraints = {}
if "last_agent_context" not in st.session_state:
    st.session_state.last_agent_context = {}
if "pending_images" not in st.session_state:
    st.session_state.pending_images = []
if "auth_token" not in st.session_state:
    # Use the token from your provided code
    st.session_state.auth_token = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiIrOTE5NzE3MjY1NjExIiwianRpIjoiNDAwMjI5YjAtNzg0OS00NjYwLTg1Y2QtMjkzM2RjMmU3OWRhIiwiVXNlcklkIjoiZWI0NmVkYTItODUzMy00MDNkLWE4ZjktMWQ3NjlkMjk4Nzc1IiwiSWQiOiJlYjQ2ZWRhMi04NTMzLTQwM2QtYThmOS0xZDc2OWQyOTg3NzUiLCJFbWFpbCI6ImF5dXNoLnJhanB1dEBub3VyaXEuYWkiLCJUZW5hbnRJZCI6IjQiLCJEYXRhYmFzZU5hbWUiOiJGcmlza2FBaUNDTV9IRldMIiwiRGF0YUJhc2VOYW1lIjoiRnJpc2thQWlDQ01fSEZXTCIsIlJvbGVOYW1lIjoiUGF0aWVudCIsIldlbGxuZXNzU3RhdHVzIjoiRW5yb2xsZWQiLCJXZWxsbmVzc1N0YXR1c0NvZGUiOiJFTlJPTExFRCIsIlBhdGllbnRGaXJzdE5hbWUiOiJBeXVzaCIsIklzSW50ZXJuYWxVc2VyIjoiRmFsc2UiLCJQcm9qZWN0VHlwZSI6IkNDTVRlc3RpbmciLCJleHAiOjE3Nzk0NzMyMjcsImlzcyI6IkhGV0wiLCJhdWQiOiJIRldMIn0.9BM8MP2EUv1dhs7kaPH4EkdzePlRufWd4xUSue65d_c"
if "blood_report_temp_path" not in st.session_state:
    st.session_state.blood_report_temp_path = None
if "fitness_profile_data" not in st.session_state:
    st.session_state.fitness_profile_data = {}
    
# --- CRITICAL FIX 1: DYNAMIC MIC KEY INITIALIZATION ---
if "mic_key" not in st.session_state:
    st.session_state.mic_key = "initial_mic_key"


def call_backend_api(payload: dict, token: str, files: list = None):
    try:
        headers = {
            "Authorization": f"Bearer {token}"
            # Do NOT manually set Content-Type here; 'requests' does it for multipart
        }
        
        # We need to send 'query' and 'last_agent_context' as Form fields
        form_data = {
            "query": payload.get("query", ""),
            "last_agent_context": json.dumps(payload.get("last_agent_context", {}))
        }

        # Handle Image files for the 'image' parameter in your backend
        file_payload = None
        if files and len(files) > 0:
            # Taking the first image as your backend 'image' parameter expects one file
            file_payload = {"image": (files[0].name, files[0].getvalue(), files[0].type)}

        # Use 'data=' for form-data, 'files=' for the image
        response = requests.post(
            BACKEND_API_URL, 
            data=form_data, 
            files=file_payload, 
            headers=headers, 
            timeout=1200
        )
        
        response.raise_for_status()
        return response.json()
    except requests.exceptions.RequestException as e:
        st.error(f"Backend Error: {e}")
        return None

def call_meal_plan_api(profile: dict, query: str, token: str, force_tool_type: str = "meal_plan_generator", start_date: Optional[str] = None, end_date: Optional[str] = None):
    """Sends the frontend-collected profile directly to the backend so it can be
    used as the meal-plan constraints, bypassing the DB profile lookup/auth."""
    try:
        headers = {"Authorization": f"Bearer {token}"}
        body = {
            "profile": profile,
            "query": query,
            "last_agent_context": st.session_state.last_agent_context,
            "force_tool_type": force_tool_type,
            "start_date": start_date,
            "end_date": end_date
        }
        response = requests.post(
            MEAL_PLAN_FROM_PROFILE_API_URL,
            json=body,
            headers=headers,
            timeout=1200
        )
        response.raise_for_status()
        return response.json()
    except requests.exceptions.RequestException as e:
        st.error(f"Backend Error: {e}")
        return None

def convert_images_to_base64_urls(files: list):
    try:
        base64_urls = []
        for uploaded_file in files:
            image_bytes = uploaded_file.getvalue()
            base64_encoded = base64.b64encode(image_bytes).decode('utf-8')
            
            mime_type = uploaded_file.type if uploaded_file.type else "image/jpeg"
            data_url = f"data:{mime_type};base64,{base64_encoded}"
            base64_urls.append(data_url)
        
        return base64_urls
    except Exception as e:
        st.error(f"Failed to process images: {e}")
        return None

# def send_message_to_backend(user_input: str, force_tool: Optional[str] = None, images: list = None):
#     with st.spinner("Friska is thinking..."):
#         api_payload = {
#             "query": user_input,
#             "chat_history": st.session_state.messages, 
#             "current_constraints": st.session_state.current_constraints,
#             "last_agent_context": st.session_state.last_agent_context,
#             "force_tool_type": force_tool 
#         }
        
#         response_data = call_backend_api(api_payload, st.session_state.auth_token, files = images) 

#         if response_data:
#             st.session_state.messages.append({
#                 "role": "assistant", 
#                 "content": response_data.get("answer"), 
#                 "audio": response_data.get("audio"), 
#                 "context": response_data.get("context")
#             })
                
#             st.session_state.current_constraints = response_data.get("updated_constraints", {})
#             st.session_state.last_agent_context = response_data.get("updated_last_agent_context", {})
    
#             st.rerun()
            
def send_message_to_backend(user_input: str, force_tool: Optional[str] = None, images: list = None):
    with st.spinner("Friska is thinking..."):
        api_payload = {
            "query": user_input,
            "constraints": st.session_state.current_constraints,
            "last_agent_context": st.session_state.last_agent_context,
            "force_tool_type": force_tool 
        }
        
        full_response = call_backend_api(api_payload, st.session_state.auth_token, files=images) 

        if full_response and "data" in full_response:
            # Extract 'ai_data' from the 'data' key returned by your FastAPI chat endpoint
            ai_data = full_response["data"]
            
            st.session_state.messages.append({
                "role": "assistant", 
                "content": ai_data.get("answer"), # Backend uses 'answer'
                "audio": ai_data.get("audio_base64"), # Your backend uses 'audio_base64'
                "context": ai_data.get("context")
            })
                
            # Update history and context
            st.session_state.last_agent_context = ai_data.get("updated_last_agent_context", {})
    
            st.rerun()

def save_uploaded_file_temp(uploaded_file):
    try:
        temp_dir = "temp_uploads"
        os.makedirs(temp_dir, exist_ok=True)
        
        temp_file_path = os.path.join(temp_dir, uploaded_file.name)
        counter = 1
        original_temp_file_path = temp_file_path
        while os.path.exists(temp_file_path):
            name, ext = os.path.splitext(original_temp_file_path)
            temp_file_path = f"{name}_{counter}{ext}"
            counter += 1

        with open(temp_file_path, "wb") as f:
            f.write(uploaded_file.getvalue())
        return temp_file_path
    except Exception as e:
        st.error(f"Error saving file: {e}")
        return None

def render_unified_profile_form():
    st.header("1. Complete Your Health Profile")
    st.info("Fill out the relevant sections below and click the corresponding button to generate your plan or set your goals.")
    
    with st.form(key="unified_profile_form"):

        st.markdown("---")
        st.subheader("🥝 Nutrition Assistant Profile")
        
        col1, col2 = st.columns(2)

        with col1:
            st.markdown("##### Personal Details")
            nut_user_name_input = st.text_input("First Name", value="Jansen", key="nut_name")
            nut_age = st.number_input("Age", min_value=18, max_value=100, value=30, step=1, key="nut_age")
            nut_weight = st.number_input("Weight (kg)", min_value=30.0, max_value=300.0, value=75.0, step=0.1, format="%.1f", key="nut_weight")
            
            h_col1, h_col2 = st.columns(2)
            with h_col1:
                nut_height_feet = st.number_input("Height (feet)", min_value=4, max_value=8, value=5, step=1, key="nut_height_ft")
            with h_col2:
                nut_height_inches = st.number_input("Height (inches)", min_value=0, max_value=11, value=9, step=1, key="nut_height_in")
            
            nut_waist_circumference = st.number_input("Waist Circumference (inches)", min_value=20.0, max_value=80.0, value=35.0, step=0.1, format="%.1f", key="nut_waist")

        with col2:
            st.markdown("##### Dietary & Lifestyle")
            nut_gender = st.selectbox("Gender", options=["Male", "Female"], index=0, key="nut_gender")
            nut_activity_level = st.selectbox("Activity Level", options=["Sedentary", "Lightly active", "Moderately active", "Very active", "Extra active"], index=2, key="nut_activity")
            nut_dietary_preference = st.selectbox("Dietary Preference", options=DIETARY_PREFERENCE_OPTIONS, key="nut_diet_pref")
            nut_dietary_restrictions = st.multiselect("Dietary Restrictions", options=DIETARY_RESTRICTION_OPTIONS, key="nut_diet_rest")
            
            nut_medical_conditions = st.multiselect("Medical Conditions", options=MEDICAL_CONDITIONS_NUTRITION_OPTIONS, key="nut_med_cond")
            nut_digestive_issues = st.multiselect("Digestive Issues", options=DIGESTIVE_ISSUES_OPTIONS, key="nut_dig_issues")
            nut_symptom_aggravating_foods = st.multiselect("Symptom-Aggravating Foods", options=SYMPTOM_AGGRAVATING_FOODS_OPTIONS, key="nut_symptom_foods")
            nut_cuisine = st.multiselect("Preferred Cuisine", options=CUISINE_OPTIONS, key="nut_cuisine")
            nut_food_allergies = st.multiselect("Food Allergies", options=FOOD_ALLERGIES_OPTIONS, key="nut_allergies")

        st.markdown("##### Upload Blood Report (Optional)")
        nut_uploaded_blood_report = st.file_uploader("Upload a blood report (PDF, JPG, JPEG, PNG, DOCX, CSV, TXT)", type=["pdf", "jpg", "jpeg", "png", "docx", "csv", "txt"], key="nut_uploader")

        st.markdown("##### Vital Signs (Optional, numeric values only)")
        
        v_col1_r1, v_col2_r1 = st.columns(2)
        with v_col1_r1:
            nut_hr = st.number_input("Heart Rate (beats per minute)", value=80, key="nut_hr")
        with v_col2_r1:
            nut_bp_sys = st.number_input("Blood Pressure - Systolic (mmHg)", value=120, key="nut_bps")

        v_col1_r2, v_col2_r2 = st.columns(2)
        with v_col1_r2:
            nut_glucose = st.number_input("Blood Glucose (mg/dL)", value=90, key="nut_glucose")
        with v_col2_r2:
            nut_bp_dia = st.number_input("Blood Pressure - Diastolic (mmHg)", value=80, key="nut_bpd")

        v_col1_r3, v_col2_r3 = st.columns(2)
        with v_col1_r3:
            nut_resp_rate = st.number_input("Respiration Rate (breaths per minute)", value=16, key="nut_rr")
        with v_col2_r3:
            nut_blood_oxygen = st.number_input("Blood Oxygen Saturation (%)", value=98, key="nut_bo")

        v_col1_r4, v_col2_r4 = st.columns(2)
        with v_col1_r4:
            nut_blood_ketones = st.number_input("Blood Ketones (mmol/L)", value=0.3, format="%.2f", key="nut_bk")
        with v_col2_r4:
            nut_body_temp = st.number_input("Body Temperature (°C)", value=36.7, format="%.1f", key="nut_bt")

        
        v_col1_r5, v_col2_r5 = st.columns(2)
        with v_col1_r5:
            nut_body_fat_percentage = st.number_input("Body Fat %", value=20.0, format="%.1f", min_value=0.0, max_value=100.0, key="nut_bfp")
        with v_col2_r5:
            pass

        nut_submitted = st.form_submit_button("Generate My 7-Day Meal Plan", type="primary", use_container_width=True, help="Creates or updates your Nutrition Profile and generates a personalized seven-day meal plan.", key="nut_submit_btn")

        # Optional profile sections are intentionally disabled for this page.
        fit_submitted = False
        ws_submitted = False
        booking_submitted = False
        _disabled_optional_profile_sections = """

        st.markdown("---")
        st.subheader("🏋️ Personalized Fitness Profile")
        
        col_f1, col_f2, col_f3 = st.columns(3)
        with col_f1:
            fit_name = st.text_input("Fitness Name *", placeholder="Your name", key="fit_name_input")
        with col_f2:
            fit_age = st.number_input("Fitness Age *", min_value=13, max_value=100, value=30, key="fit_age_input")
        with col_f3:
            fit_gender = st.selectbox("Fitness Gender *", ["Male", "Female", "Other"], key="fit_gender_input")
        
        col_f4, col_f5, col_f6 = st.columns(3)
        with col_f4:
            fit_unit_system = st.radio("Units *", ["Metric (kg, cm)", "Imperial (lbs, in)"], key="fit_unit_input")
            
            fit_weight_kg = 0.0
            fit_height_cm = 0.0

        with col_f5:
            if fit_unit_system == "Metric (kg, cm)":
                fit_weight_metric = st.number_input("Fitness Weight (kg) *", min_value=30.0, max_value=300.0, value=70.0, key="fit_weight_kg_input")
                fit_weight_kg = fit_weight_metric
            else:
                fit_weight_lbs = st.number_input("Fitness Weight (lbs) *", min_value=66.0, max_value=660.0, value=154.3, key="fit_weight_lbs_input")
                fit_weight_kg = fit_weight_lbs * 0.453592

        with col_f6:
            if fit_unit_system == "Metric (kg, cm)":
                fit_height_metric = st.number_input("Fitness Height (cm) *", min_value=100.0, max_value=250.0, value=170.0, key="fit_height_cm_input")
                fit_height_cm = fit_height_metric
            else:
                fit_height_in = st.number_input("Fitness Height (in) *", min_value=39.0, max_value=98.0, value=66.9, key="fit_height_in_input")
                fit_height_cm = fit_height_in * 2.54

        st.markdown("##### Fitness Goals & Level")
        
        col_f7, col_f8 = st.columns(2)
        with col_f7:
            fit_primary_goal = st.selectbox(
                "Primary Goal *",
                FITNESS_GOAL_OPTIONS, key="fit_primary_goal_input"
            )
        
        with col_f8:
            fit_secondary_goal = st.selectbox(
                "Secondary Goal (Optional)",
                ["None"] + FITNESS_GOAL_OPTIONS, key="fit_secondary_goal_input"
            )
        
        fit_fitness_level = st.selectbox(
            "Fitness Level *",
            FITNESS_LEVEL_OPTIONS, key="fit_fitness_level_input"
        )
        
        st.markdown("##### Health Screening")
        
        fit_medical_conditions = st.multiselect(
            "Fitness Medical Conditions *",
            FITNESS_MEDICAL_CONDITIONS_OPTIONS, 
            default=["None"], 
            key="fit_medical_conditions_input"
        )
        
        st.warning("⚠️ **Physical Limitations** - Describe ANY injuries, pain, or movement restrictions")
        fit_physical_limitation = st.text_area( 
            "Physical Limitations (Important for Safety) *",
            placeholder="E.g., 'Previous right knee surgery - avoid deep squats'",
            height=100, key="fit_physical_limitation_input"
        )
        
        # --- NEW: Specific Exercise Avoidance ---
        st.warning("⚠️ **Specific Exercise Restrictions**")
        fit_specific_avoidance_input = st.text_area(
            "Have you been advised to avoid any specific exercises? (List them below if known to avoid in plan generation):",
            placeholder="E.g., 'Heavy deadlifts, overhead pressing due to shoulder issue, any exercise that causes sharp pain in the elbow.'",
            height=100,
            key="specific_avoidance_text_input"
        )
        # --- END NEW: Specific Exercise Avoidance ---
        
        st.markdown("##### Training Schedule & Equipment")
        
        col_f9, col_f10 = st.columns(2)
        with col_f9:
            fit_days_per_week = st.multiselect(
                "Training Days *",
                FITNESS_DAYS_OPTIONS,
                default=["Monday", "Wednesday", "Friday"], key="fit_days_per_week_input"
            )
        
        with col_f10:
            fit_session_duration = st.selectbox(
                "Session Duration *",
                FITNESS_DURATION_OPTIONS, key="fit_session_duration_input"
            )

        fit_workout_location = st.selectbox(
            "Where will you primarily work out? *",
            FITNESS_LOCATION_OPTIONS, key="fit_location_input"
        )
        
        fit_equipment = st.multiselect("Select all available equipment: *", FITNESS_EQUIPMENT_OPTIONS, default=["Bodyweight Only"], key="fit_equipment_input")

        fit_submitted = st.form_submit_button("🚀 Generate Fitness Plan", type="secondary", use_container_width=True, help="Creates or updates your Fitness Profile and generates a personalized Workout Plan.", key="fit_submit_btn")


        st.markdown("---")
        st.subheader("💧 Water & Step Goal Profile")
        st.caption("These fields are used to calculate personalized daily step and water targets.")
        
        current_constraints = st.session_state.current_constraints
        default_age_ws = current_constraints.get('age', 30)
        default_weight_ws = current_constraints.get('weight_kg', 70.0)
        default_gender_ws = current_constraints.get('gender', 'Male')
        default_height_ws = current_constraints.get('height_cm', 170.0)
        default_conditions_ws = current_constraints.get('medical_conditions', [])
        default_goal_ws = current_constraints.get('goal', {}).get('type', 'Maintenance')
        
        st.markdown("##### Core Profile Details (Manual Entry for Calculation)")
        
        col_ws_1, col_ws_2 = st.columns(2)
        
        with col_ws_1:
            ws_age = st.number_input("WS Age *", min_value=13, max_value=100, value=default_age_ws, step=1, key="ws_age_input")
            ws_weight = st.number_input("WS Weight (kg) *", min_value=30.0, max_value=300.0, value=default_weight_ws, step=0.1, format="%.1f", key="ws_weight_input")
            ws_gender = st.selectbox("WS Gender *", options=["Male", "Female", "Other"], index=(0 if default_gender_ws == 'Male' else 1) if default_gender_ws in ['Male', 'Female'] else 2, key="ws_gender_input")
            
        with col_ws_2:
            ws_height = st.number_input("WS Height (cm) *", min_value=100.0, max_value=250.0, value=default_height_ws, step=0.1, format="%.1f", key="ws_height_input")
            
            ws_bmi = ws_weight / ((ws_height / 100) ** 2) if ws_weight > 0 and ws_height > 0 else 0.0
            st.markdown(f"Calculated BMI: **{ws_bmi:.2f}**")
            
            ws_medical_conditions = st.multiselect("WS Medical Conditions (for calculation)", 
                                                 options=["None", "Diabetes Type 2", "Hypertension", "Pregnancy", "Lactation"], 
                                                 default=[c for c in default_conditions_ws if c in ["Diabetes Type 2", "Hypertension", "Pregnancy", "Lactation"]], 
                                                 key="ws_med_cond_input")

        st.markdown("##### 👣 Step Goal Parameters")
        
        step_occupations = ["Sedentary (office)", "Moderate (teacher, retail, etc.)", "Active job (labor, delivery, etc.)"]
        ws_occupation = st.selectbox("Occupation (Affects Base Activity)", step_occupations, key="ws_occupation")
        
        step_fitness_lvls = ["Beginner", "Intermediate", "Advanced"]
        ws_activity = st.selectbox("Current Fitness Level (WS)", step_fitness_lvls, key="ws_activity")
        
        step_goals = ["Weight loss", "Maintenance", "Weight Gain", "Health management (e.g., diabetes)"]
        ws_goal = st.selectbox("Primary Health Goal (WS)", step_goals, index=step_goals.index(default_goal_ws) if default_goal_ws in step_goals else 1, key="ws_goal")
        
        st.markdown("##### 💧 Water Goal Parameters")
        
        ws_climates = ["Moderate / Temperate", "Hot / Humid (>30°C)", "High Altitude (>1500m)", "Dry / Air-Conditioned"]
        ws_climate = st.selectbox("Typical Climate/Environment", ws_climates, key="ws_climate")
        
        ws_diets = ["Standard", "High Protein (>1.6g/kg)", "High Fiber", "Salty / Spicy", "Low-carb/Keto"]
        ws_diet = st.selectbox("Dietary Intake Pattern (WS)", ws_diets, key="ws_diet")

        ws_submitted = st.form_submit_button("Save and Calculate Goals", type="secondary", use_container_width=True, help="Calculates and saves your Step and Water intake goals to your profile.", key="ws_submit_btn")

        
        st.markdown("---")
        st.subheader("📅 Expert Session Booking")
        st.caption("Fill in your preferred session details to automatically trigger the booking conversation.")

        default_user_name = st.session_state.current_constraints.get('name', 'User')
        
        expert_name = st.selectbox("Select Expert", EXPERT_NAMES, key="booking_expert")
        session_type = st.selectbox("Session Focus", SESSION_TYPES, key="booking_focus")
        appointment_type = st.selectbox("Appointment Type", APPOINTMENT_TYPES, key="booking_type")
        
        date_col, time_col = st.columns(2)
        with date_col:
            session_date = st.date_input("Preferred Date", key="booking_date")
        with time_col:
            session_time = st.time_input("Preferred Time", value=datetime.now().time(), key="booking_time")

        booking_submitted = st.form_submit_button("Book Session Now", type="secondary", use_container_width=True, help="Initiates a conversational flow to book the specified session.", key="booking_submit_btn")
        """


    if nut_submitted:
        with st.spinner("Creating your profile and generating a personalized meal plan..."):

            blood_report_path = None
            if nut_uploaded_blood_report:
                blood_report_path = save_uploaded_file_temp(nut_uploaded_blood_report)
                st.session_state.blood_report_temp_path = blood_report_path

            nut_height_cm = (nut_height_feet * 12 + nut_height_inches) * 2.54

            profile_data = {
                "name": nut_user_name_input,
                "age": nut_age,
                "gender": nut_gender,
                "weight_kg": nut_weight,
                "height_cm": nut_height_cm,
                "waist_circumference_cm": nut_waist_circumference * 2.54, 
                "activity_level": nut_activity_level.lower(),
                "dietary_preference": [nut_dietary_preference], 
                "restrictions": nut_dietary_restrictions,
                "medical_conditions": nut_medical_conditions,
                "allergies": nut_food_allergies,
                "cuisine": nut_cuisine,
                "digestive_issues": nut_digestive_issues,
                "symptom_aggravating_foods": nut_symptom_aggravating_foods,
                "vitals_numeric": {
                    "Heart Rate": nut_hr, 
                    "Blood Pressure": f"{nut_bp_sys}/{nut_bp_dia}", 
                    "Blood Glucose": nut_glucose,
                    "Respiration Rate": nut_resp_rate,
                    "Blood Oxygen Saturation": nut_blood_oxygen,
                    "Blood Ketones": nut_blood_ketones,
                    "Body Temperature": nut_body_temp,
                    "Body Fat %": nut_body_fat_percentage
                }
            }

            profile_data = {
                key: value
                for key, value in profile_data.items()
                if key in PROFILE_CONSTRAINT_FIELDS
            }

            st.session_state.current_constraints = profile_data.copy()
            plan_start_date = datetime.now().date()
            plan_end_date = plan_start_date + timedelta(days=6)
            plan_start = plan_start_date.isoformat()
            plan_end = plan_end_date.isoformat()
            initial_prompt = "Create a complete 7-day meal plan for me based on my profile."

            # Send the frontend-collected profile directly to the backend so it is
            # used as the meal-plan constraints, bypassing the DB profile lookup/auth.
            response_data = call_meal_plan_api(
                profile=profile_data,
                query=initial_prompt,
                token=st.session_state.auth_token,
                force_tool_type="weekly_meal_plan_generator",
                start_date=plan_start,
                end_date=plan_end
            )

            if response_data and "data" in response_data:
                ai_content = response_data["data"]
                st.session_state.weekly_meal_plan = ai_content.get("answer", "")
                st.session_state.messages.append({
                    "role": "assistant", 
                    "content": ai_content.get("answer"), 
                    "audio": ai_content.get("audio"),
                    "context": ai_content.get("context")
                })
                
                st.session_state.current_constraints = ai_content.get("updated_constraints", st.session_state.current_constraints)
                st.session_state.last_agent_context = ai_content.get("updated_last_agent_context", {})
                
                st.session_state.profile_complete = True
                st.rerun()

    _disabled_optional_profile_handlers = """
    elif fit_submitted:
        if not fit_name.strip():
            st.error("❌ Please enter your name.")
        elif not fit_days_per_week:
            st.error("❌ Please select at least one training day.")
        elif (fit_weight_kg <= 0 and fit_unit_system == "Metric (kg, cm)") or (fit_height_cm <= 0 and fit_unit_system == "Metric (kg, cm)"):
            st.error("❌ Please ensure valid weight and height inputs.")
        else:
            sanitized_med_conditions = [mc for mc in fit_medical_conditions if mc != "None"]
            if not sanitized_med_conditions:
                sanitized_med_conditions = ["None"]
                
            sanitized_equipment = fit_equipment if fit_equipment else ["Bodyweight Only"]

            fit_weight_kg_final = fit_weight_kg
            fit_height_cm_final = fit_height_cm
            if fit_unit_system == "Imperial (lbs, in)":
                fit_weight_kg_final = st.session_state["fit_weight_lbs_input"] * 0.453592
                fit_height_cm_final = st.session_state["fit_height_in_input"] * 2.54

            fit_profile = {
                "name": fit_name.strip(),
                "age": fit_age,
                "gender": fit_gender,
                "weight_kg": fit_weight_kg_final,
                "height_cm": fit_height_cm_final,
                "bmi": fit_weight_kg_final / ((fit_height_cm_final / 100) ** 2) if fit_weight_kg_final > 0 and fit_height_cm_final > 0 else 0,
                "primary_goal": fit_primary_goal,
                "secondary_goal": fit_secondary_goal if fit_secondary_goal != "None" else None,
                "fitness_level": fit_fitness_level,
                "medical_conditions": sanitized_med_conditions,
                "physical_limitation": fit_physical_limitation.strip(),
                "days_per_week": fit_days_per_week,
                "session_duration": fit_session_duration,
                "available_equipment": sanitized_equipment,
                "workout_location": fit_workout_location,
                "specific_avoidance": fit_specific_avoidance_input.strip() or "None" # CRITICAL: Include new avoidance field
            }
            
            st.session_state.fitness_profile_data = fit_profile
            
            st.session_state.current_constraints.update({
                "name": fit_name.strip(),
                "age": fit_age,
                "gender": fit_gender,
                "weight_kg": fit_weight_kg_final,
                "height_cm": fit_height_cm_final,
                "primary_goal": fit_primary_goal,
                "fitness_level": fit_fitness_level,
                "days_per_week": fit_days_per_week,
                "available_equipment": sanitized_equipment,
                "physical_limitation": fit_physical_limitation.strip(), 
                "specific_avoidance": fit_specific_avoidance_input.strip() or "None", # CRITICAL: Include new avoidance field in constraints
                "medical_conditions": sanitized_med_conditions,
                "fitness_profile_data": fit_profile
            })
            
            initial_prompt = f"Generate a detailed weekly workout plan based on my profile, starting now."
            
            st.session_state.messages = [] 
            st.session_state.messages.append({"role": "user", "content": initial_prompt})
            
            with st.spinner("Generating your comprehensive fitness plan..."):
                api_payload = {
                    "query": initial_prompt,
                    "chat_history": st.session_state.messages, 
                    "current_constraints": st.session_state.current_constraints,
                    "last_agent_context": st.session_state.last_agent_context,
                    "force_tool_type": "fitness_plan_generator" 
                }
                
                response_data = call_backend_api(api_payload, st.session_state.auth_token)

                if response_data:
                    st.session_state.messages.append({
                        "role": "assistant", 
                        "content": response_data.get("answer"), 
                        "audio": response_data.get("audio"), 
                        "context": response_data.get("context")
                    })
                    st.session_state.current_constraints = response_data.get("updated_constraints", {})
                    st.session_state.last_agent_context = response_data.get("updated_last_agent_context", {})
                
                st.session_state.profile_complete = True
                st.rerun()

    elif ws_submitted:
        if ws_weight <= 30.0 or ws_height <= 100.0:
            st.error("❌ Please provide realistic weight and height inputs for Water/Step calculation.")
            return

        ws_update_payload = {
            "age": ws_age, 
            "gender": ws_gender, 
            "weight_kg": ws_weight, 
            "height_cm": ws_height, 
            "medical_conditions": [c for c in ws_medical_conditions if c != "None"], 
            "occupation": ws_occupation,
            "activity_level": ws_activity, 
            "goal": {"type": ws_goal},
            "climate": ws_climate,
            "diet_type": ws_diet
        }
        
        st.session_state.current_constraints.update(ws_update_payload)
        
        initial_prompt = "What are my daily step and water goals?"
        
        st.session_state.messages.append({"role": "user", "content": initial_prompt})
        
        send_message_to_backend(initial_prompt, force_tool="water_step_advisor", images=None)
        
        st.session_state.messages.append({"role": "assistant", "content": "Water and Step calculation parameters saved. You can now ask: **'What is my daily step goal?'** or **'How much water should I drink?'**"})
        st.rerun()
        
    elif booking_submitted:
        user_query = (
            f"I want to book a {session_type} session with {expert_name} on "
            f"{session_date.strftime('%B %d, %Y')} at {session_time.strftime('%I:%M %p')} for a {appointment_type}. "
            f"My name is {default_user_name}."
        )

        st.session_state.messages.append({"role": "user", "content": user_query})
        
        send_message_to_backend(user_query, force_tool="session_booking")
        
        st.session_state.profile_complete = True
        st.rerun()

    """


if not st.session_state.profile_complete:
    render_unified_profile_form()

else:
    _mg_days_preview = parse_weekly_plan(st.session_state.get("weekly_meal_plan", "")) if callable(parse_weekly_plan) else []
    if premium_ui is not None:
        premium_ui.render_hero(st.session_state.current_constraints, _mg_days_preview)
        premium_ui.render_profile_chips(st.session_state.current_constraints)
    else:
        st.title("Personalised Meal Plan Generator")
        st.write("Premium UI is unavailable in this environment.")

    with st.expander("✏️ Edit Profile and Regenerate Plan", expanded=False):
        constraints = st.session_state.current_constraints
        existing_diet = constraints.get("dietary_preference", [])
        if isinstance(existing_diet, str):
            existing_diet = [existing_diet]

        with st.form("edit_meal_plan_profile_form"):
            edit_col1, edit_col2 = st.columns(2)
            with edit_col1:
                edit_name = st.text_input("First Name", value=constraints.get("name", ""), key="edit_plan_name")
                edit_age = st.number_input("Age", min_value=18, max_value=100, value=int(constraints.get("age", 30)), key="edit_plan_age")
                edit_weight = st.number_input("Weight (kg)", min_value=30.0, max_value=300.0, value=float(constraints.get("weight_kg", 75.0)), key="edit_plan_weight")
                edit_gender = st.selectbox("Gender", ["Male", "Female"], index=0 if constraints.get("gender", "Male") == "Male" else 1, key="edit_plan_gender")
                edit_activity = st.selectbox("Activity Level", ["Sedentary", "Lightly active", "Moderately active", "Very active", "Extra active"], index=2, key="edit_plan_activity")
            with edit_col2:
                edit_diet = st.multiselect("Dietary Preference", DIETARY_PREFERENCE_OPTIONS, default=[x for x in existing_diet if x in DIETARY_PREFERENCE_OPTIONS], key="edit_plan_diet")
                edit_restrictions = st.multiselect("Dietary Restrictions", DIETARY_RESTRICTION_OPTIONS, default=[x for x in constraints.get("restrictions", []) if x in DIETARY_RESTRICTION_OPTIONS], key="edit_plan_restrictions")
                edit_conditions = st.multiselect("Medical Conditions", MEDICAL_CONDITIONS_NUTRITION_OPTIONS, default=[x for x in constraints.get("medical_conditions", []) if x in MEDICAL_CONDITIONS_NUTRITION_OPTIONS], key="edit_plan_conditions")
                edit_allergies = st.multiselect("Food Allergies", FOOD_ALLERGIES_OPTIONS, default=[x for x in constraints.get("allergies", []) if x in FOOD_ALLERGIES_OPTIONS], key="edit_plan_allergies")
                edit_cuisine = st.multiselect("Preferred Cuisine", CUISINE_OPTIONS, default=[x for x in constraints.get("cuisine", []) if x in CUISINE_OPTIONS], key="edit_plan_cuisine")
                edit_digestive = st.multiselect("Digestive Issues", DIGESTIVE_ISSUES_OPTIONS, default=[x for x in constraints.get("digestive_issues", []) if x in DIGESTIVE_ISSUES_OPTIONS], key="edit_plan_digestive")
                edit_symptoms = st.multiselect("Symptom-Aggravating Foods", SYMPTOM_AGGRAVATING_FOODS_OPTIONS, default=[x for x in constraints.get("symptom_aggravating_foods", []) if x in SYMPTOM_AGGRAVATING_FOODS_OPTIONS], key="edit_plan_symptoms")

            edit_profile_submitted = st.form_submit_button("Save Profile and Regenerate 7-Day Plan", type="primary", use_container_width=True)

        if edit_profile_submitted:
            updated_profile = {
                **constraints,
                "name": edit_name,
                "age": edit_age,
                "weight_kg": edit_weight,
                "gender": edit_gender,
                "activity_level": edit_activity.lower(),
                "dietary_preference": edit_diet,
                "restrictions": edit_restrictions,
                "medical_conditions": edit_conditions,
                "allergies": edit_allergies,
                "cuisine": edit_cuisine,
                "digestive_issues": edit_digestive,
                "symptom_aggravating_foods": edit_symptoms,
            }
            edit_start_date = datetime.now().date()
            edit_end_date = edit_start_date + timedelta(days=6)
            with st.spinner("Regenerating your meal plan from the updated profile..."):
                edit_response = call_meal_plan_api(
                    profile=updated_profile,
                    query="Generate a complete 7-day meal plan based on my updated profile.",
                    token=st.session_state.auth_token,
                    force_tool_type="weekly_meal_plan_generator",
                    start_date=edit_start_date.isoformat(),
                    end_date=edit_end_date.isoformat()
                )
            if edit_response and isinstance(edit_response.get("data"), dict):
                edit_data = edit_response["data"]
                st.session_state.current_constraints = edit_data.get("updated_constraints", updated_profile)
                st.session_state.weekly_meal_plan = edit_data.get("answer", "")
                st.session_state.last_agent_context = edit_data.get("updated_last_agent_context", st.session_state.last_agent_context)
                st.rerun()
            else:
                st.error("The updated seven-day meal plan could not be generated.")

    weekly_start_date = datetime.now().date()
    weekly_end_date = weekly_start_date + timedelta(days=6)
    weekly_start = weekly_start_date.isoformat()
    weekly_end = weekly_end_date.isoformat()

    st.caption(f"Generate a complete meal plan from {weekly_start} to {weekly_end}.")

    if st.button("Explore My Week \u2192", type="primary", key="generate_7_day_meal_plan"):
        weekly_query = "Generate a complete 7-day meal plan based on my profile."
        with st.spinner("Designing your week..."):
            weekly_response = call_meal_plan_api(
                profile=st.session_state.current_constraints,
                query=weekly_query,
                token=st.session_state.auth_token,
                force_tool_type="weekly_meal_plan_generator",
                start_date=weekly_start,
                end_date=weekly_end
            )

        if weekly_response and isinstance(weekly_response.get("data"), dict):
            weekly_data = weekly_response["data"]
            st.session_state.weekly_meal_plan = weekly_data.get("answer", "")
            st.session_state.current_constraints = weekly_data.get(
                "updated_constraints", st.session_state.current_constraints
            )
            st.session_state.last_agent_context = weekly_data.get(
                "updated_last_agent_context", st.session_state.last_agent_context
            )
        else:
            st.error("We couldn't prepare your plan. Something interrupted the generation. Your existing plan is safe.")

    weekly_meal_plan = st.session_state.get("weekly_meal_plan", "")
    mg_days = parse_weekly_plan(weekly_meal_plan)

    if not weekly_meal_plan:
        premium_ui.render_empty_plan_state()
    elif mg_days and premium_ui is not None:
        selected_idx = premium_ui.render_day_selector(mg_days)
        selected_day = mg_days[selected_idx]

        premium_ui.render_daily_overview(selected_day)

        for meal_i, meal in enumerate(selected_day.meals):
            premium_ui.render_meal_card(meal, selected_day.day_number, meal_i)

        premium_ui.render_nutrition_intelligence(selected_day)
        premium_ui.render_ai_insight(selected_day, st.session_state.current_constraints)
    elif mg_days:
        st.subheader("Weekly Meal Plan")
        for day in mg_days:
            st.markdown(f"### Day {day.day_number}: {day.day_name}, {day.date_label}")
            for meal in day.meals:
                st.markdown(f"**{meal.name}**")
                for item in meal.items:
                    st.write(f"- {item.name}")
    else:
        # Fall back to the original renderer if the text doesn't match the
        # structured "### Day N: ..." format (e.g. older/edited plans).
        if MealPlanRenderer is not None:
            MealPlanRenderer.render_meal_plan_simple(weekly_meal_plan)
        else:
            st.text(weekly_meal_plan or "No meal plan available yet.")

    _legacy_chat_and_save_ui = """
    chat_col, buttons_col = st.columns([3, 1])

    with chat_col:
        for message in st.session_state.messages:
            with st.chat_message(message["role"]):
                content_to_display = message.get("content", "")
                
                if not isinstance(content_to_display, str):
                    content_to_display = json.dumps(content_to_display, indent=2)

                context = message.get("context", {})
                is_meal_plan = context.get("is_meal_plan", False) if isinstance(context, dict) else False
                
                # 🎯 NEW: Use interactive meal plan component for meal plans
                if message["role"] == "assistant" and is_meal_plan:
                    # Use the new meal plan renderer component
                    # This will automatically strip metadata and make food items clickable
                    display_meal_plan(content_to_display, interactive=True, style="simple")
                else:
                    # Regular message rendering
                    st.markdown(content_to_display, unsafe_allow_html=True)
                
                if message["role"] == "assistant" and context and context.get("play_video"):
                    video_file_name = context["play_video"]
                    video_path = video_file_name 
                    
                    if os.path.exists(video_path):
                        st.video(video_path)
                    else:
                        st.warning(f"Video file '{video_file_name}' not found. Please ensure it is placed in the 'videos' directory.")
      
                
                if message["role"] == "assistant" and message.get("audio"):
                    audio_data = message["audio"]
    # Check if the data is a string and long enough to be a real Base64 audio file
                    if audio_data and isinstance(audio_data, str) and len(audio_data) > 100: 
                        try:
                            audio_bytes = base64.b64decode(audio_data, validate=True)
                            st.audio(audio_bytes, format="audio/mp3")
                        except (binascii.Error, TypeError, ValueError) as e:
            # Error now caught, providing a debug message without crashing the frontend
                            print(f"DEBUG: Could not decode audio for a message. Error: {e}")
                            st.info("TTS Audio playback skipped due to an audio decoding error from the backend.")


        if st.session_state.last_agent_context.get("expecting_meal_image_upload"):
            uploaded_meal_image = st.file_uploader("Upload a photo of what you ate", type=["jpg", "jpeg", "png"], key="meal_image_upload")

            if uploaded_meal_image:
                st.session_state.last_agent_context["expecting_meal_image_upload"] = False

                with st.spinner("Analyzing your meal image..."):
                    try:
                        response = requests.post(
                            "http://127.0.0.1:8502/analyze-meal-image",
                            files={"file": uploaded_meal_image.getvalue()}
                        )
                        result = response.json()

                        if result.get("success"):
                            analysis_text = result["analysis_text"]
                            st.session_state.messages.append({"role": "user", "content": analysis_text})
                            send_message_to_backend(analysis_text)
                        else:
                            st.error(f"Analysis failed: {result.get('error')}")
                    except Exception as e:
                        st.error(f"Error contacting backend for image analysis: {e}")

        # --- START MODIFIED MIC/INPUT SECTION (Optimized for Streamlit 1.52.1) ---
        
        # Streamlit's chat_input is complex. Use the old-style columns for better control over mic placement
        col_mic, col_upload, col_input = st.columns([0.5, 0.5, 4])
        
        audio_bytes = None
        mic_available = False

        # Audio Recording Section
        with col_mic:
            st.markdown("<div style='height: 30px;'></div>", unsafe_allow_html=True) # Vertical alignment spacer
            try:
                # --- CRITICAL FIX 2: Use the dynamic key for the audio recorder ---
                audio_bytes = audio_recorder(
                    text="",
                    key=st.session_state.mic_key 
                )
                mic_available = True
            except Exception as e:
                # Catch-all for any failure, allowing the app to continue
                st.error(f"🎙️ Mic component failed to load. Error: {e}", icon="⚠️")
                mic_available = False

        # Image Upload
        with col_upload:
             uploaded_images = st.file_uploader(
                "📎",
                type=["jpg", "jpeg", "png"],
                accept_multiple_files=True,
                key="chat_image_uploader",
                label_visibility="collapsed"
            )
        
        # Text Input & Logic
        with col_input:
            
            # --- CRITICAL FIX 4: CHECK AND RESET PENDING VOICE STATE ---
            # Use a session state flag to prevent re-execution of the audio logic 
            # after the first successful transcription.
            if "voice_input_ready" not in st.session_state:
                st.session_state.voice_input_ready = None
                
            transcribed_text_to_process = None
            if st.session_state.voice_input_ready:
                transcribed_text_to_process = st.session_state.voice_input_ready
                st.session_state.voice_input_ready = None # Consume the state
                st.session_state.pending_images = []
                
            if transcribed_text_to_process:
                user_input = transcribed_text_to_process
                st.session_state.messages.append({"role": "user", "content": user_input})
                
                # Display the transcribed text immediately in the chat column
                with chat_col:
                    with st.chat_message("user"):
                        st.markdown(f"*(Voice)*: {user_input}") 
                
                # Send the transcribed text to your main conversational API
                send_message_to_backend(user_input, images=None) # Images are now handled in the previous turn if needed
                
                # We do NOT call st.rerun here; send_message_to_backend will call it after the API response is received.
            
            # Check if audio was recorded successfully
            # This is the original logic that *prepares* the transcribed text.
            elif audio_bytes is not None and mic_available:
                st.info("Processing recorded voice command...")
                with st.spinner("Transcribing voice command..."):
                    try:
                        # Call your backend's STT endpoint to transcribe the audio
                        base_url = BACKEND_API_URL.rsplit('/', 1)[0]
                        
                        stt_res = requests.post(
                            f"{base_url}/transcribe/", 
                            files={'file': ('audio.wav', audio_bytes, 'audio/wav')},
                            timeout=100
                        )
                        stt_res.raise_for_status()
                        transcribed_text = stt_res.json().get('transcribed_text', '')

                        if transcribed_text and "Could not understand" not in transcribed_text:
                            # CRITICAL: Store the transcribed text and force a rerun
                            st.session_state.voice_input_ready = transcribed_text
                            
                            # Force key change to reset the microphone widget
                            st.session_state.mic_key = str(datetime.now()) 
                            
                            # CRITICAL: Rerun to process the `voice_input_ready` state
                            st.rerun() 

                        else:
                            st.warning("Sorry, I didn't catch that. Could you please type your command instead?")
                            # Clear the key even on failure to reset the mic button.
                            st.session_state.mic_key = str(datetime.now()) 
                            st.rerun() # Rerun to reset mic state
                            
                    except Exception as e:
                        st.error(f"Error during transcription: {e}")
                        # Clear the key even on critical failure
                        st.session_state.mic_key = str(datetime.now()) 
                        st.rerun() # Rerun to reset mic state

            elif user_input := st.chat_input("Ask me about your plan or anything else..."):
                
                images_to_send = st.session_state.pending_images.copy() if st.session_state.pending_images else None
                
                st.session_state.messages.append({"role": "user", "content": user_input})
                with chat_col:
                    with st.chat_message("user"):
                        st.markdown(user_input)
                        if images_to_send:
                            for img in images_to_send:
                                st.image(img, width=200)
                
                st.session_state.pending_images = []
                
                send_message_to_backend(user_input, images=images_to_send)

        # Handle pending images display
        if uploaded_images:
            st.session_state.pending_images = uploaded_images
            
        # Display image upload count clearly *outside* the chat_input if images are pending
        if st.session_state.pending_images:
            st.caption(f"{len(st.session_state.pending_images)} image(s) selected for next message.")

        # --- END MODIFIED MIC/INPUT SECTION ---


    with buttons_col:
        st.subheader("Meal Check-in")
        
        meal_types = ["Breakfast", "Morning Snack", "Lunch", "Evening Snack", "Dinner"]

        plan_context = st.session_state.last_agent_context
        current_plan_text = plan_context.get("pending_meal_plan_for_save") or plan_context.get("active_meal_plan")
        
        is_daily_plan = current_plan_text and plan_context.get("plan_type") == "DAILY"
        is_weekly_plan = current_plan_text and plan_context.get("plan_type") == "WEEKLY"
        plan_exists = bool(current_plan_text)

        for meal_type in meal_types:
            if st.button(f"Check-in for {meal_type}", disabled=not plan_exists, help="Please generate a meal plan first to enable check-in." if not plan_exists else None, key=f"checkin_{meal_type}"):
                user_query = f"Hi I just got my {meal_type} reminder. Can you help me log what I ate"
                
                st.session_state.messages.append({"role": "user", "content": user_query})
                
                send_message_to_backend(user_query, force_tool="meal_check_in", images=None)

        st.subheader("Grocery List")
        
        if st.button("Generate Daily Grocery List", 
                      disabled=not is_daily_plan, 
                      help="Generate a daily meal plan first." if not is_daily_plan else "Get ingredients for your daily plan.", 
                      key="daily_grocery_list"):
            user_query = "Generate a grocery list for my daily meal plan."
            st.session_state.messages.append({"role": "user", "content": user_query})
            send_message_to_backend(user_query, force_tool="general_query", images=None)
            
        if st.button("Generate Weekly Grocery List", 
                      disabled=not is_weekly_plan, 
                      help="Generate a weekly meal plan first." if not is_weekly_plan else "Get ingredients for your weekly plan.", 
                      key="weekly_grocery_list"):
            user_query = "Generate a grocery list for my weekly meal plan."
            st.session_state.messages.append({"role": "user", "content": user_query})
            send_message_to_backend(user_query, force_tool="general_query", images=None)

            """

    st.sidebar.title("Actions")
    st.sidebar.title("🧪 TTS Debugging")

    last_assistant_message = None
    for msg in reversed(st.session_state.messages):
        if msg.get("role") == "assistant":
            last_assistant_message = msg
            break

    if last_assistant_message:
        original_text = last_assistant_message.get("content", "N/A")
        summarized_text = last_assistant_message.get("audio")

        st.sidebar.subheader("Original Text (Input to Summary)")
        st.sidebar.text_area("Original", value=original_text, height=200, key="debug_original", disabled=True)

        # Updated Debugging for the Base64 Audio data
        st.sidebar.subheader("Base64 Audio Data (from backend)")
        if summarized_text:
            st.sidebar.text_area("Base64 String", value=summarized_text[:100] + "..." if len(summarized_text) > 100 else summarized_text, height=100, key="debug_summary", disabled=True)
            if len(summarized_text) < 50:
                 st.sidebar.warning("Note: The 'audio' field is short. This might be a mistake if Base64 audio was expected.")
        else:
            st.sidebar.info("No audio data was sent for the last message.")
    else:
        st.sidebar.info("No assistant message yet to display.")

    st.sidebar.divider()
    
    if st.sidebar.button("Start Over"):
        token = st.session_state.get("auth_token", "")
        if st.session_state.blood_report_temp_path and os.path.exists(st.session_state.blood_report_temp_path):
            try:
                os.remove(st.session_state.blood_report_temp_path)
                print(f"Cleaned up temporary blood report file: {st.session_state.blood_report_temp_path}")
            except Exception as e:
                print(f"Error cleaning up temporary blood report file: {e}")
        
        for key in list(st.session_state.keys()):
            if key not in ["auth_token"]:
                del st.session_state[key]
        st.session_state.profile_complete = False
        st.rerun()
    
    if st.sidebar.button("Save Chat Logs"):
        doc = Document()
        profile = st.session_state.get('current_constraints', {})
        doc.add_heading(f"Friska Chat History for {profile.get('name', 'User')}", level=1)
        if profile:
            doc.add_heading("User Profile", level=2)
            profile_text_parts = []
            for k, v in profile.items():
                if k not in ['vitals_numeric', 'vitals_text']:
                    key_title = k.replace('_', ' ').title()
                    value_str = ", ".join(v) if isinstance(v, list) else str(v)
                    profile_text_parts.append(f"{key_title}: {value_str}")
            
            if 'vitals_numeric' in profile and profile['vitals_numeric']:
                profile_text_parts.append("\nVitals:")
                for vital_key, vital_value in profile['vitals_numeric'].items():
                        profile_text_parts.append(f"- {vital_key.replace('_', ' ').title()}: {vital_value}")

            doc.add_paragraph("\n".join(profile_text_parts))
        
        doc.add_heading("Conversation", level=2)
        for msg in st.session_state.messages:
            content = msg.get('content', '')
            if isinstance(content, dict):
                content = json.dumps(content, indent=2)

            clean_content = re.sub(r'<.*?>', '', content)
            clean_content = clean_content.replace("**", "")
            doc.add_paragraph(f"{msg['role'].title()}: {clean_content}", style='Body Text')
        
        doc_io = BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)
        
        st.sidebar.download_button(
            label="Download Chat Logs as .docx",
            data=doc_io,
            file_name=f"friska_chat_{profile.get('name', 'user').lower()}_history.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    premium_ui.render_footer()